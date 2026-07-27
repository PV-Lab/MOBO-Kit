"""Read-only structural audits for historical and updated D2D workbooks."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import DefaultDict, Sequence

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .candidate_pool import physical_rows_to_grid_indices
from .design import InputSpec, build_design


_V2_INPUT_HEADERS = (
    "speed_1",
    "time_1",
    "speed_2",
    "time_2",
    "precur_conc",
    "precur_vol (uL)",
    "anneal_temp",
    "anneal_time",
    "anti_vol",
    "anti_time",
)
_V2_CANONICAL_INPUTS = (
    "speed_1",
    "time_1",
    "speed_2",
    "time_2",
    "precur_conc",
    "precur_vol",
    "anneal_temp",
    "anneal_time",
    "anti_vol",
    "anti_time",
)
_V2_HEADERS = (
    "Sample number",
    *_V2_INPUT_HEADERS,
    "Coverage",
    "Uniformity",
    "1 - Uniformity",
    "Phase purity",
    "PL - Implied Voc (Max)",
    "PL - Implied Voc (Max) Normalized",
    "Photoconductance (Max)",
    "Photoconductance (Max) Normalized",
    "Uniformity score",
    "Optoelectronic score",
    "Thickness (avg)",
    "Normalized thickness (sigma = 250)",
    "Total combination - addition",
    None,
    "Uniformity score absolute difference",
    "Optoelectronic score absolute difference",
    "Thickness absolute difference",
    "Total score absolute difference",
)
_V3_HEADERS = (
    "Sample number",
    *_V2_INPUT_HEADERS,
    "Coverage",
    "Uniformity",
    "1 - Uniformity",
    "Phase purity",
    "PL - Implied Voc (Max)",
    "Photoconductance (Max)",
    "Log10 (Photoconductance (Max) x PL - Implied Voc (Max))",
    "T1",
    "T2",
    "T3",
    "T4",
    "T anom",
    "Thickness (avg)",
    "Normalized thickness (sigma = 250)",
    "Uniformity score",
    "Optoelectronic score",
    "Thickness score",
    "Stability score?",
    "Total combination - addition",
    "Total combination - multiplied",
    "Uniformity score absolute difference",
    "Optoelectronic score absolute difference",
    "Thickness absolute difference",
    "Total score absolute difference",
)
_V3_OBJECTIVE_HEADERS = (
    "Uniformity score",
    "Optoelectronic score",
    "Thickness score",
)
_V3_CANONICAL_OBJECTIVES = (
    "uniformity_score",
    "optoelectronic_score",
    "thickness_score",
)
_V3_IGNORED_MODEL_HEADERS = (
    "Stability score?",
    "Total combination - addition",
    "Total combination - multiplied",
    "Uniformity score absolute difference",
    "Optoelectronic score absolute difference",
    "Thickness absolute difference",
    "Total score absolute difference",
)


@dataclass(frozen=True)
class WorkbookInputException:
    """A versioned, observed-only input exception accepted by one profile."""

    sample_id: int
    row: int
    input_name: str
    observed_value: float
    reason: str


@dataclass(frozen=True)
class WorkbookInputExceptionRule:
    """Explicit runtime rule for one observed-only off-grid input value."""

    sample_id: int
    input_name: str
    observed_value: float
    reason: str


@dataclass(frozen=True)
class WorkbookNote:
    cell: str
    row: int
    column: int
    value: str


@dataclass(frozen=True)
class WorkbookAudit:
    """Raw schema, recognized profile, and non-mutating validation findings."""

    workbook_path: Path
    sheet_names: list[str]
    active_sheet: str
    used_range: str
    raw_headers: list[str | None]
    duplicate_headers: dict[str, list[int]]
    blank_header_columns: list[int]
    sample_row_count: int
    warnings: list[str]
    profile: str = "unknown"
    header_positions: dict[str, list[int]] | None = None
    canonical_input_mapping: dict[str, str] | None = None
    canonical_input_positions: dict[str, int] | None = None
    sample_rows: list[int] | None = None
    notes: list[WorkbookNote] | None = None
    input_rows_valid: bool | None = None
    input_validation_errors: list[str] | None = None
    objective_mapping_approved: bool = False
    formula_cell_count: int = 0
    canonical_objective_mapping: dict[str, str] | None = None
    canonical_objective_positions: dict[str, int] | None = None
    ignored_model_positions: dict[str, int] | None = None
    objective_mapping_resolved_for_debug: bool = False
    input_exceptions: list[WorkbookInputException] | None = None

    @property
    def column_count(self) -> int:
        return len(self.raw_headers)


def _is_nonblank(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.replace("\u00a0", " ").strip())
    return True


def _header_text(value: object) -> str | None:
    if not _is_nonblank(value):
        return None
    return value if isinstance(value, str) else str(value)


def _normalized_header(value: str) -> str:
    return "".join(character for character in value.casefold() if character.isalnum())


def _sample_identifier_columns(header_cells: dict[int, object]) -> list[int]:
    aliases = {"sample", "sampleid", "samplenumber"}
    return [
        column
        for column, value in header_cells.items()
        if (header := _header_text(value)) is not None
        and _normalized_header(header) in aliases
    ]


def _numeric_sample_identifier(value: object) -> int | None:
    if isinstance(value, (bool, np.bool_)) or not isinstance(
        value, (int, float, np.integer, np.floating)
    ):
        return None
    numeric_value = float(value)
    if not np.isfinite(numeric_value) or not numeric_value.is_integer():
        return None
    return int(numeric_value)


def _numeric_sample_rows(
    row_values: dict[int, dict[int, object]], *, sample_column: int
) -> list[int]:
    return [
        row
        for row, values in sorted(row_values.items())
        if _numeric_sample_identifier(values.get(sample_column)) is not None
    ]


def _recognize_profile(
    raw_headers: list[str | None],
    *,
    active_sheet: str,
    header_row: int,
    min_column: int,
) -> str:
    if (
        active_sheet == "Sheet1"
        and header_row == 1
        and min_column == 1
        and tuple(raw_headers) == _V3_HEADERS
    ):
        return "d2d_summary_v3_scores"
    if (
        active_sheet == "Sheet1"
        and header_row == 1
        and min_column == 1
        and tuple(raw_headers) == _V2_HEADERS
    ):
        return "d2d_summary_v2"
    if (
        len(raw_headers) == 29
        and "Anneal Temp" in raw_headers
        and raw_headers.count("Uniformity score") == 2
    ):
        return "d2d_summary_step1_historical"
    return "unknown"


def _d2d_design():
    return build_design(
        [
            InputSpec("speed_1", 1000, 6000, 500),
            InputSpec("time_1", 5, 50, 5),
            InputSpec("speed_2", 0, 5000, 500),
            InputSpec("time_2", 10, 60, 5),
            InputSpec("precur_conc", 1.0, 2.0, 0.05),
            InputSpec("precur_vol", 40, 200, 10),
            InputSpec("anneal_temp", 100, 185, 5),
            InputSpec("anneal_time", 10, 60, 5),
            InputSpec("anti_vol", 100, 200, 5),
            InputSpec("anti_time", 9, 25, 2),
        ]
    )


def _validate_v2_inputs(
    row_values: dict[int, dict[int, object]], sample_rows: list[int]
) -> tuple[bool, list[str]]:
    errors: list[str] = []
    expected_rows = list(range(2, 17))
    if sample_rows != expected_rows:
        errors.append(
            "The v2 profile requires sample rows exactly at Excel rows 2-16; "
            f"found {sample_rows}."
        )
    if row_values.get(17):
        errors.append("The v2 profile requires Excel row 17 to be blank.")

    observed_identifiers: list[int | None] = []
    for row in expected_rows:
        raw_identifier = row_values.get(row, {}).get(1)
        if isinstance(raw_identifier, (bool, np.bool_)) or not isinstance(
            raw_identifier, (int, float, np.integer, np.floating)
        ):
            observed_identifiers.append(None)
            continue
        numeric_identifier = float(raw_identifier)
        if not np.isfinite(numeric_identifier) or not numeric_identifier.is_integer():
            observed_identifiers.append(None)
            continue
        observed_identifiers.append(int(numeric_identifier))
    numeric_identifiers = [
        identifier for identifier in observed_identifiers if identifier is not None
    ]
    if len(numeric_identifiers) != len(expected_rows) or len(
        set(numeric_identifiers)
    ) != len(numeric_identifiers):
        errors.append(
            "The v2 profile requires 15 unique numeric sample identifiers in Excel "
            f"rows 2-16; found {observed_identifiers}."
        )

    physical_rows: list[tuple[int, list[float]]] = []
    for row in sample_rows:
        values: list[float] = []
        for column, name in zip(range(2, 12), _V2_CANONICAL_INPUTS):
            raw = row_values.get(row, {}).get(column)
            if isinstance(raw, (bool, np.bool_)):
                errors.append(f"Row {row} input {name!r} is boolean, not numeric.")
                continue
            try:
                number = float(raw)
            except (TypeError, ValueError):
                errors.append(f"Row {row} input {name!r} is missing or nonnumeric.")
                continue
            if not np.isfinite(number):
                errors.append(f"Row {row} input {name!r} is non-finite.")
                continue
            values.append(number)
        if len(values) == 10:
            physical_rows.append((row, values))
    valid_indices: list[np.ndarray] = []
    valid_excel_rows: list[int] = []
    for excel_row, physical_row in physical_rows:
        try:
            valid_indices.append(
                physical_rows_to_grid_indices(
                    np.asarray(physical_row, dtype=float)[None, :], _d2d_design()
                )[0]
            )
            valid_excel_rows.append(excel_row)
        except ValueError as exc:
            errors.append(f"Excel row {excel_row}: {exc}")
    if valid_indices:
        index_array = np.asarray(valid_indices)
        unique_indices, inverse, counts = np.unique(
            index_array, axis=0, return_inverse=True, return_counts=True
        )
        del unique_indices
        duplicate_groups = [
            [
                valid_excel_rows[position]
                for position in np.flatnonzero(inverse == group)
            ]
            for group, count in enumerate(counts)
            if count > 1
        ]
        if duplicate_groups:
            errors.append(
                "The v2 recipe rows contain duplicate input grid tuples at Excel "
                f"row group(s) {duplicate_groups}."
            )
    return not errors, errors


def _validate_v3_inputs(
    row_values: dict[int, dict[int, object]],
    sample_rows: list[int],
    allowed_input_exceptions: Sequence[WorkbookInputExceptionRule],
) -> tuple[bool, list[str], list[WorkbookInputException]]:
    """Validate v3 observations against explicit runtime exception rules."""

    errors: list[str] = []
    exceptions: list[WorkbookInputException] = []
    expected_rows = list(range(2, 17))
    if sample_rows != expected_rows:
        errors.append(
            "The v3 profile requires numeric sample rows exactly at Excel rows "
            f"2-16; found {sample_rows}."
        )

    observed_identifiers = [
        _numeric_sample_identifier(row_values.get(row, {}).get(1))
        for row in expected_rows
    ]
    numeric_identifiers = [
        identifier for identifier in observed_identifiers if identifier is not None
    ]
    if len(numeric_identifiers) != len(expected_rows) or len(
        set(numeric_identifiers)
    ) != len(numeric_identifiers):
        errors.append(
            "The v3 profile requires 15 unique numeric sample identifiers in Excel "
            f"rows 2-16; found {observed_identifiers}."
        )

    design = _d2d_design()
    rules_by_key: dict[tuple[int, str], WorkbookInputExceptionRule] = {}
    for rule in allowed_input_exceptions:
        if not isinstance(rule, WorkbookInputExceptionRule):
            raise TypeError(
                "allowed_input_exceptions must contain WorkbookInputExceptionRule values."
            )
        key = (rule.sample_id, rule.input_name)
        if key in rules_by_key:
            raise ValueError(f"Duplicate workbook input exception rule for {key!r}.")
        if rule.input_name not in design.names:
            raise ValueError(
                f"Unknown workbook input exception field {rule.input_name!r}."
            )
        dimension = design.names.index(rule.input_name)
        observed = float(rule.observed_value)
        if not np.isfinite(observed):
            raise ValueError("Workbook input exception values must be finite.")
        if observed < design.lowers[dimension] or observed > design.uppers[dimension]:
            raise ValueError("Workbook input exception values must remain in bounds.")
        if np.any(
            np.isclose(
                observed,
                design.var_array[dimension],
                rtol=0.0,
                atol=1e-9,
            )
        ):
            raise ValueError("Workbook input exception values must be off-grid.")
        if not str(rule.reason).strip():
            raise ValueError("Workbook input exception rules require a reason.")
        rules_by_key[key] = rule

    valid_grid_indices: list[np.ndarray] = []
    valid_excel_rows: list[int] = []
    physical_rows: list[tuple[int, int | None, list[float]]] = []
    encountered_rule_keys: set[tuple[int, str]] = set()
    for row in sample_rows:
        sample_id = _numeric_sample_identifier(row_values.get(row, {}).get(1))
        values: list[float] = []
        for column, name in zip(range(2, 12), _V2_CANONICAL_INPUTS):
            raw = row_values.get(row, {}).get(column)
            if isinstance(raw, (bool, np.bool_)):
                errors.append(f"Row {row} input {name!r} is boolean, not numeric.")
                continue
            try:
                number = float(raw)
            except (TypeError, ValueError):
                errors.append(f"Row {row} input {name!r} is missing or nonnumeric.")
                continue
            if not np.isfinite(number):
                errors.append(f"Row {row} input {name!r} is non-finite.")
                continue
            values.append(number)
        if len(values) != len(_V2_CANONICAL_INPUTS):
            continue

        physical_rows.append((row, sample_id, values))
        physical_row = np.asarray(values, dtype=float)
        out_of_bounds = np.flatnonzero(
            (physical_row < design.lowers) | (physical_row > design.uppers)
        )
        if out_of_bounds.size:
            details = ", ".join(
                f"{design.names[index]}={physical_row[index]:g} outside "
                f"[{design.lowers[index]:g}, {design.uppers[index]:g}]"
                for index in out_of_bounds
            )
            errors.append(f"Excel row {row} has out-of-bounds input(s): {details}.")
            continue

        grid_probe = physical_row.copy()
        row_exceptions: list[WorkbookInputException] = []
        for dimension, input_name in enumerate(design.names):
            observed = float(physical_row[dimension])
            if np.any(
                np.isclose(
                    observed,
                    design.var_array[dimension],
                    rtol=0.0,
                    atol=1e-9,
                )
            ):
                continue
            rule = (
                None
                if sample_id is None
                else rules_by_key.get((int(sample_id), input_name))
            )
            if rule is None or not np.isclose(
                observed, rule.observed_value, rtol=0.0, atol=1e-12
            ):
                continue
            grid_probe[dimension] = float(design.var_array[dimension][0])
            encountered_rule_keys.add((rule.sample_id, rule.input_name))
            row_exceptions.append(
                WorkbookInputException(
                    sample_id=rule.sample_id,
                    row=row,
                    input_name=rule.input_name,
                    observed_value=rule.observed_value,
                    reason=rule.reason,
                )
            )
        try:
            grid_index = physical_rows_to_grid_indices(grid_probe[None, :], design)[0]
        except ValueError as exc:
            errors.append(f"Excel row {row}: {exc}")
            continue

        if row_exceptions:
            exceptions.extend(row_exceptions)
        else:
            valid_grid_indices.append(grid_index)
            valid_excel_rows.append(row)

    missing_rules = sorted(set(rules_by_key) - encountered_rule_keys)
    if missing_rules:
        errors.append(
            "Configured workbook input exception rule(s) were not found exactly: "
            f"{missing_rules}."
        )

    if valid_grid_indices:
        index_array = np.asarray(valid_grid_indices)
        _, inverse, counts = np.unique(
            index_array, axis=0, return_inverse=True, return_counts=True
        )
        duplicate_groups = [
            [
                valid_excel_rows[position]
                for position in np.flatnonzero(inverse == group)
            ]
            for group, count in enumerate(counts)
            if count > 1
        ]
        if duplicate_groups:
            errors.append(
                "The v3 recipe rows contain duplicate on-grid input tuples at "
                f"Excel row group(s) {duplicate_groups}."
            )

    if physical_rows:
        tuples_to_rows: DefaultDict[tuple[float, ...], list[int]] = defaultdict(list)
        for row, _, values in physical_rows:
            tuples_to_rows[tuple(values)].append(row)
        duplicate_observed_groups = [
            rows for rows in tuples_to_rows.values() if len(rows) > 1
        ]
        if duplicate_observed_groups and not any(
            "duplicate on-grid input tuples" in error for error in errors
        ):
            errors.append(
                "The v3 recipe rows contain duplicate observed input tuples at "
                f"Excel row group(s) {duplicate_observed_groups}."
            )

    return not errors, errors, exceptions


def audit_campaign_workbook(
    path: str | Path,
    *,
    allowed_input_exceptions: Sequence[WorkbookInputExceptionRule] = (),
) -> WorkbookAudit:
    """Inspect the active worksheet in read-only mode without saving it."""
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=False,
    )
    try:
        sheet_names = list(workbook.sheetnames)
        worksheet = workbook.active
        active_sheet = worksheet.title
        reset_dimensions = getattr(worksheet, "reset_dimensions", None)
        if callable(reset_dimensions):
            reset_dimensions()

        min_row: int | None = None
        min_column: int | None = None
        max_row: int | None = None
        max_column: int | None = None
        header_row: int | None = None
        header_cells: dict[int, object] = {}
        sample_columns: list[int] = []
        sample_rows: list[int] = []
        row_values: dict[int, dict[int, object]] = {}
        formula_count = 0

        for row in worksheet.iter_rows():
            nonblank_cells: list[tuple[int, int, object]] = []
            for cell in row:
                value = cell.value
                if not _is_nonblank(value):
                    continue
                cell_row = int(cell.row)
                cell_column = int(cell.column)
                nonblank_cells.append((cell_row, cell_column, value))
                row_values.setdefault(cell_row, {})[cell_column] = value
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                min_row = cell_row if min_row is None else min(min_row, cell_row)
                min_column = (
                    cell_column if min_column is None else min(min_column, cell_column)
                )
                max_row = cell_row if max_row is None else max(max_row, cell_row)
                max_column = (
                    cell_column if max_column is None else max(max_column, cell_column)
                )
            if not nonblank_cells:
                continue
            row_number = nonblank_cells[0][0]
            if header_row is None:
                header_row = row_number
                header_cells = {column: value for _, column, value in nonblank_cells}
                sample_columns = _sample_identifier_columns(header_cells)
                continue
            if sample_columns and any(
                column == sample_columns[0] for _, column, _ in nonblank_cells
            ):
                sample_rows.append(row_number)

        if (
            min_row is None
            or min_column is None
            or max_row is None
            or max_column is None
            or header_row is None
        ):
            raise ValueError(
                f"Active worksheet {active_sheet!r} contains no nonblank cell values."
            )

        raw_headers = [
            _header_text(header_cells.get(column))
            for column in range(min_column, max_column + 1)
        ]
        positions_by_header: DefaultDict[str, list[int]] = defaultdict(list)
        for column, header in zip(range(min_column, max_column + 1), raw_headers):
            if header is not None:
                positions_by_header[header].append(column)
        header_positions = dict(positions_by_header)
        duplicate_headers = {
            header: positions
            for header, positions in header_positions.items()
            if len(positions) > 1
        }
        blank_header_columns = [
            column
            for column, header in zip(range(min_column, max_column + 1), raw_headers)
            if header is None
        ]

        warnings: list[str] = []
        for header, positions in duplicate_headers.items():
            warnings.append(
                f"Duplicate header {header!r} appears at 1-based columns "
                f"{positions}; columns remain separate."
            )
        for column in blank_header_columns:
            previous_header = (
                raw_headers[column - min_column - 1] if column > min_column else None
            )
            after_text = (
                f" after {previous_header!r}" if previous_header is not None else ""
            )
            warnings.append(
                f"Blank header at 1-based column {column} "
                f"({get_column_letter(column)}){after_text} within the used range."
            )
        normalized_groups: DefaultDict[str, list[tuple[str, int]]] = defaultdict(list)
        for column, header in zip(range(min_column, max_column + 1), raw_headers):
            if header is not None:
                normalized_groups[_normalized_header(header)].append((header, column))
        for normalized, entries in normalized_groups.items():
            raw_names = {header for header, _ in entries}
            if len(raw_names) > 1:
                details = ", ".join(
                    f"{header!r} (column {column})" for header, column in entries
                )
                warnings.append(
                    f"Ambiguous related headers normalize to {normalized!r}: {details}."
                )
        if not sample_columns:
            warnings.append(
                "No sample identifier header was found; sample_row_count is 0."
            )
        elif len(sample_columns) > 1:
            warnings.append(
                "Multiple sample identifier headers were found at 1-based columns "
                f"{sample_columns}; sample rows were counted from column "
                f"{sample_columns[0]}."
            )

        used_range = (
            f"{get_column_letter(min_column)}{min_row}:"
            f"{get_column_letter(max_column)}{max_row}"
        )
        profile = _recognize_profile(
            raw_headers,
            active_sheet=active_sheet,
            header_row=header_row,
            min_column=min_column,
        )
        if profile == "d2d_summary_v3_scores":
            # Rows 17-20 contain explanatory notes in the v3 workbook. Only
            # integer-valued Sample number cells identify observations.
            sample_rows = _numeric_sample_rows(row_values, sample_column=1)

        canonical_mapping: dict[str, str] = {}
        canonical_positions: dict[str, int] = {}
        canonical_objective_mapping: dict[str, str] = {}
        canonical_objective_positions: dict[str, int] = {}
        ignored_model_positions: dict[str, int] = {}
        objective_mapping_resolved_for_debug = False
        input_exceptions: list[WorkbookInputException] = []
        input_valid: bool | None = None
        input_errors: list[str] = []
        if profile == "d2d_summary_v2":
            canonical_mapping = dict(zip(_V2_INPUT_HEADERS, _V2_CANONICAL_INPUTS))
            canonical_positions = {
                canonical: column
                for column, canonical in zip(range(2, 12), _V2_CANONICAL_INPUTS)
            }
            input_valid, input_errors = _validate_v2_inputs(row_values, sample_rows)
            if used_range != "A1:AC18":
                input_errors.append(
                    "The v2 profile requires content range A1:AC18; "
                    f"found {used_range}."
                )
                input_valid = False
            warnings.append(
                "Objective mapping is provisional and not approved for production."
            )
            if (
                row_values.get(18, {}).get(17)
                and row_values.get(18, {}).get(19)
                and not row_values.get(18, {}).get(16)
                and not row_values.get(18, {}).get(18)
            ):
                warnings.append(
                    "Supplied workbook notes are at Q18/S18 (normalized columns); "
                    "the Step 2A pack's P18/R18 note locations do not match the file."
                )
        elif profile == "d2d_summary_v3_scores":
            canonical_mapping = dict(zip(_V2_INPUT_HEADERS, _V2_CANONICAL_INPUTS))
            canonical_positions = {
                canonical: column
                for column, canonical in zip(range(2, 12), _V2_CANONICAL_INPUTS)
            }
            canonical_objective_mapping = dict(
                zip(_V3_OBJECTIVE_HEADERS, _V3_CANONICAL_OBJECTIVES)
            )
            canonical_objective_positions = {
                canonical: column
                for column, canonical in zip(range(26, 29), _V3_CANONICAL_OBJECTIVES)
            }
            ignored_model_positions = {
                header: column
                for column, header in zip(range(29, 36), _V3_IGNORED_MODEL_HEADERS)
            }
            objective_mapping_resolved_for_debug = True
            input_valid, input_errors, input_exceptions = _validate_v3_inputs(
                row_values, sample_rows, allowed_input_exceptions
            )
            if used_range != "A1:AI20":
                input_errors.append(
                    "The v3 profile requires content range A1:AI20; "
                    f"found {used_range}."
                )
                input_valid = False
            warnings.append(
                "Objective mapping Z/AA/AB is resolved for debug only and is "
                "not approved for production."
            )
            for exception in input_exceptions:
                warnings.append(
                    f"Observed-only input exception: Sample {exception.sample_id} "
                    f"at Excel row {exception.row} retains "
                    f"{exception.input_name}={exception.observed_value:g}; "
                    "new candidates must remain on the approved grid."
                )

        sample_row_set = set(sample_rows)
        notes = [
            WorkbookNote(
                cell=f"{get_column_letter(column)}{row}",
                row=row,
                column=column,
                value=str(value),
            )
            for row, values in sorted(row_values.items())
            if row != header_row and row not in sample_row_set
            for column, value in sorted(values.items())
            if _is_nonblank(value)
        ]
        return WorkbookAudit(
            workbook_path=workbook_path,
            sheet_names=sheet_names,
            active_sheet=active_sheet,
            used_range=used_range,
            raw_headers=raw_headers,
            duplicate_headers=duplicate_headers,
            blank_header_columns=blank_header_columns,
            sample_row_count=len(sample_rows),
            warnings=warnings,
            profile=profile,
            header_positions=header_positions,
            canonical_input_mapping=canonical_mapping,
            canonical_input_positions=canonical_positions,
            sample_rows=sample_rows,
            notes=notes,
            input_rows_valid=input_valid,
            input_validation_errors=input_errors,
            objective_mapping_approved=False,
            formula_cell_count=formula_count,
            canonical_objective_mapping=canonical_objective_mapping,
            canonical_objective_positions=canonical_objective_positions,
            ignored_model_positions=ignored_model_positions,
            objective_mapping_resolved_for_debug=(objective_mapping_resolved_for_debug),
            input_exceptions=input_exceptions,
        )
    finally:
        workbook.close()


__all__ = [
    "WorkbookAudit",
    "WorkbookInputException",
    "WorkbookInputExceptionRule",
    "WorkbookNote",
    "audit_campaign_workbook",
]
