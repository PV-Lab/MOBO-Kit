"""What a proposed batch actually says, before anyone fabricates it.

Fifteen films is a real cost, and until 2026-07-30 nothing in this repo showed a
human what a batch meant -- only that it passed its validity checks.  This module
builds one artifact: a table of the proposed conditions in physical units, what the
model predicts for each and how sure it is, how far each sits from anything already
measured, and which coordinates are pinned at a range edge.  It is written as a
``Review`` sheet beside the worklist and echoed into the launcher window, so it can
be forwarded to the experimental group on its own.

Three things it is careful about.

**The predictions come from the same path the acquisition used.**
:func:`campaign.fit_campaign_models` reproduces the round's model bit for bit from
the same data and seed, and utility moments come from
``ucb_hvi.posterior_utility_moments`` -- the function the round itself called.  A
review that computed utilities its own way could disagree with the batch it is
reviewing, which would be worse than no review.

**It reports physical values as well as utilities.** A utility of 0.87 means
nothing to the person running the coater; "predicted 612 nm, 68% interval
480-780" does.  For a log-link objective the decoded value is the posterior
*median*, because ``exp`` of a mean of logs is not a mean.

**Probes are declared in config, not hardcoded here.** A probe asks a
counterfactual: hold everything else, force one input to a value worth
interrogating, and report what the model thinks there.  What is worth
interrogating is campaign knowledge, so it lives in the campaign YAML under
``review.probes`` -- as do the standing notes under ``review.notes``.

The artifact ends where it should: nothing here is approved.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
import pandas as pd
import torch

from .campaign import (
    build_objective_transform,
    fit_campaign_models,
    normalise_inputs,
    objective_names,
)
from .design import build_design_from_config
from .scores import ScoreFinding, ScoreSeverity
from .ucb_hvi import posterior_utility_moments

__all__ = [
    "BatchReview",
    "NOT_APPROVED",
    "ProbeSpec",
    "SD_MATERIALITY_RATIO",
    "build_batch_review",
    "classify_probe_objective",
    "probe_specs_from_config",
    "review_notes_from_config",
    "write_review_sheet",
]

#: How much larger a probed region's posterior sd must be before it counts as
#: "the model finds this region uncertain" rather than "about the same".
#:
#: A bare ``>`` comparison is useless here: on the campaign's R0 fit the probed
#: sd came out 1-8% above the selected batch's on all three objectives, which a
#: strict inequality reads as "more uncertain" even while predicted thickness
#: utility falls from 0.79 to 0.22. A few percent of sd is not a reason to skip a
#: region; a 3.5x drop in predicted utility is. The ratio is printed either way, so
#: a reader who prefers a different line can draw it.
SD_MATERIALITY_RATIO = 1.25

def classify_probe_objective(
    twin_utility: float, batch_utility: float, sd_ratio: float
) -> str | None:
    """How to read one objective at a probed value.

    ``"known_and_bad"`` -- scores worse with no materially greater uncertainty.
    Under UCB that is the interesting case: uncertainty is what UCB pays for, so a
    region skipped despite equal uncertainty is being skipped on its predicted
    value, which means the model believes it knows.

    ``"uncertain_tradeoff"`` -- scores worse but genuinely more uncertain, so the
    skip is a trade-off against the other objectives. The benign reading.

    ``None`` -- does not score worse, so the model has no objection to the region
    and its absence is about batch spacing, not merit.
    """
    if not (twin_utility < batch_utility):
        return None
    return (
        "uncertain_tradeoff" if sd_ratio > SD_MATERIALITY_RATIO else "known_and_bad"
    )


NOT_APPROVED = (
    "Nothing here is approved. These conditions were proposed by an optimiser and "
    "have not been reviewed by anyone. Read them, decide, and record the decision "
    "outside this file."
)


# --------------------------------------------------------------------------- #
# configuration
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class ProbeSpec:
    """A counterfactual to report beside the batch.

    ``column`` forced to ``value``, everything else held at each proposed
    candidate's own coordinates.  ``note`` is the campaign's reason for asking.
    """

    name: str
    column: str
    value: float
    note: str = ""

    def __post_init__(self) -> None:
        if not isinstance(self.name, str) or not self.name.strip():
            raise ValueError("A probe needs a non-empty name.")
        if not isinstance(self.column, str) or not self.column.strip():
            raise ValueError(f"Probe {self.name!r} needs a column.")
        object.__setattr__(self, "name", self.name.strip())
        object.__setattr__(self, "column", self.column.strip())
        value = float(self.value)
        if not np.isfinite(value):
            raise ValueError(f"Probe {self.name!r} needs a finite value.")
        object.__setattr__(self, "value", value)


def probe_specs_from_config(config: Mapping[str, Any]) -> tuple[ProbeSpec, ...]:
    """Read ``review.probes``; absent means no probes, which is fine."""
    review = config.get("review") or {}
    if not isinstance(review, Mapping):
        raise ValueError("config['review'] must be a mapping.")
    raw = review.get("probes") or ()
    if isinstance(raw, Mapping):
        raw = [raw]
    specs = []
    for entry in raw:
        if not isinstance(entry, Mapping):
            raise ValueError("Each review probe must be a mapping.")
        specs.append(
            ProbeSpec(
                name=str(entry.get("name", entry.get("column", "probe"))),
                column=str(entry["column"]),
                value=float(entry["value"]),
                note=str(entry.get("note", "")).strip(),
            )
        )
    return tuple(specs)


def review_notes_from_config(config: Mapping[str, Any]) -> tuple[str, ...]:
    """Standing notes to print with every review of this campaign."""
    review = config.get("review") or {}
    raw = review.get("notes") or ()
    if isinstance(raw, str):
        raw = [raw]
    return tuple(str(note).strip() for note in raw if str(note).strip())


# --------------------------------------------------------------------------- #
# the review
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class BatchReview:
    """Everything a human needs to judge one proposed batch."""

    round_name: str
    candidates: pd.DataFrame
    """One row per condition: inputs, predicted utility and sd per objective,
    decoded physical prediction, distance to nearest observed point, boundary
    count and which coordinates are pinned."""
    probes: pd.DataFrame
    """Counterfactual rows, one block per declared probe. Empty if none."""
    probe_verdicts: tuple[str, ...] = ()
    notes: tuple[str, ...] = ()
    findings: tuple[ScoreFinding, ...] = ()
    model_warnings: tuple[str, ...] = ()
    """Fits that succeeded but deserve distrust. Printed before anything else,
    because they change how every number below should be read."""
    context: dict[str, Any] = field(default_factory=dict)

    def to_text(self) -> str:
        """The whole artifact as text, for the launcher pane and the console."""
        width = 78
        lines: list[str] = [
            f"BATCH REVIEW - {self.round_name}",
            "=" * width,
            "",
        ]
        for key, value in self.context.items():
            lines.append(f"{key:<22} {value}")
        if self.model_warnings:
            # first, not last: these change how every number below reads
            lines += ["", "!! READ THIS BEFORE THE NUMBERS", "-" * width]
            for warning in self.model_warnings:
                lines += _wrap(warning, width) + [""]
        lines += ["", "PROPOSED CONDITIONS", "-" * width]
        lines.append(
            self.candidates.to_string(
                index=False, float_format=lambda value: f"{value:g}"
            )
        )
        if not self.probes.empty:
            lines += ["", "PROBES", "-" * width]
            lines.append(
                self.probes.to_string(
                    index=False, float_format=lambda value: f"{value:g}"
                )
            )
        if self.probe_verdicts:
            lines += [""]
            for verdict in self.probe_verdicts:
                lines += _wrap(verdict, width) + [""]
        if self.notes:
            lines += ["NOTES", "-" * width]
            for note in self.notes:
                lines += _wrap(note, width) + [""]
        if self.findings:
            lines += ["CARRIED FROM THE MEASURED DATA", "-" * width]
            for finding in _ordered(self.findings):
                lines += _wrap(str(finding), width, hang=2)
            lines += [""]
        lines += ["=" * width] + _wrap(NOT_APPROVED, width)
        return "\n".join(lines)


def _wrap(text: str, width: int, *, hang: int = 0) -> list[str]:
    import textwrap

    out: list[str] = []
    for paragraph in str(text).split("\n"):
        wrapped = textwrap.wrap(paragraph.strip(), width=width) or [""]
        out.extend(wrapped[:1] + [" " * hang + line for line in wrapped[1:]])
    return out


def _ordered(findings: Sequence[ScoreFinding]) -> list[ScoreFinding]:
    rank = {ScoreSeverity.ERROR: 0, ScoreSeverity.WARNING: 1, ScoreSeverity.NOTE: 2}
    return sorted(findings, key=lambda f: (rank[f.severity], f.row_position))


def _utility_moments(
    config: Mapping[str, Any],
    model: Any,
    X_phys: np.ndarray,
    *,
    round_name: str,
    seed: int,
) -> tuple[np.ndarray, np.ndarray]:
    """Utility mean and sd through the acquisition's own posterior-sample path."""
    transform = build_objective_transform(config)
    settings = (config.get("rounds") or {}).get(round_name.lower()) or {}
    samples = int(settings.get("posterior_samples", settings.get("mc_samples", 256)))
    X_norm = torch.tensor(normalise_inputs(config, X_phys), dtype=torch.double)
    moments = posterior_utility_moments(
        model, X_norm, transform, mc_samples=samples, seed=seed
    )
    return moments.utility_mean, moments.utility_std


def _physical_predictions(
    config: Mapping[str, Any], model: Any, X_phys: np.ndarray
) -> dict[str, np.ndarray]:
    """Decoded model output per objective, in the measurement's own units.

    For a log-link objective this is ``exp(mu)``: the posterior median, not the
    mean.  Labelling it a median is the honest option -- the mean of a lognormal
    is ``exp(mu + v/2)``, and quietly reporting one as the other is the kind of
    small lie that gets quoted back later.
    """
    transform = build_objective_transform(config)
    X_norm = torch.tensor(normalise_inputs(config, X_phys), dtype=torch.double)
    model.eval()
    with torch.no_grad():
        posterior = model.posterior(X_norm)
        mean = posterior.mean.detach().cpu().numpy()
        sd = posterior.variance.clamp_min(0.0).sqrt().detach().cpu().numpy()
    out: dict[str, np.ndarray] = {}
    for index, spec in enumerate(transform.specs):
        mu, sigma = mean[:, index], sd[:, index]
        if spec.model_link == "log":
            out[spec.name] = np.column_stack(
                [np.exp(mu), np.exp(mu - sigma), np.exp(mu + sigma)]
            )
        else:
            out[spec.name] = np.column_stack([mu, mu - sigma, mu + sigma])
    return out


def build_batch_review(
    config: Mapping[str, Any],
    observed_X_phys: np.ndarray,
    observed_Y_raw: np.ndarray,
    conditions: pd.DataFrame,
    *,
    round_name: str,
    seed: int | None = None,
    findings: Sequence[ScoreFinding] = (),
    context: Mapping[str, Any] | None = None,
    observed_Yvar: np.ndarray | None = None,
) -> BatchReview:
    """Assemble the review of ``conditions`` against the model that proposed them.

    ``observed_Yvar`` must be whatever the round's runner was given. Under
    ``replicate_pooled`` the proposing GPs carry the measured replicate noise; a
    review refitted without it fits its own noise and describes a different model
    from the one that chose the batch -- on the first R2 that put the uniformity
    utilities up to 0.027 high, on a sheet that said the noise was measured.
    """
    design = build_design_from_config(dict(config))
    input_names = list(design.names)
    names = list(objective_names(config))
    resolved_seed = (
        int(config.get("reproducibility", {}).get("seed", 0)) if seed is None else seed
    )

    observed = np.asarray(observed_X_phys, dtype=float)
    proposed = conditions[input_names].to_numpy(dtype=float)

    # validate the probes BEFORE fitting: a typo in a config column name should
    # cost nothing, not three GP fits
    probes = probe_specs_from_config(config)
    for probe in probes:
        if probe.column not in input_names:
            raise ValueError(
                f"Probe {probe.name!r} names {probe.column!r}, which is not a "
                f"declared input. Declared inputs: {input_names}."
            )

    model, model_warnings = fit_campaign_models(
        config, observed, observed_Y_raw, seed=resolved_seed, Yvar=observed_Yvar
    )

    utility_mean, utility_sd = _utility_moments(
        config, model, proposed, round_name=round_name, seed=resolved_seed
    )
    physical = _physical_predictions(config, model, proposed)

    observed_norm = normalise_inputs(config, observed)
    proposed_norm = normalise_inputs(config, proposed)
    gaps = np.linalg.norm(
        proposed_norm[:, None, :] - observed_norm[None, :, :], axis=-1
    )
    nearest = gaps.min(axis=1)
    nearest_index = gaps.argmin(axis=1)

    at_lower = np.isclose(proposed_norm, 0.0, atol=1e-9)
    at_upper = np.isclose(proposed_norm, 1.0, atol=1e-9)
    pinned = at_lower | at_upper

    rows: dict[str, Any] = {
        "candidate": [f"{round_name.upper()}_C{i:02d}" for i in range(1, len(conditions) + 1)]
    }
    for column in input_names:
        rows[column] = conditions[column].to_numpy(dtype=float)
    for index, name in enumerate(names):
        rows[f"{name}_utility"] = utility_mean[:, index]
        rows[f"{name}_sd"] = utility_sd[:, index]
        # numeric, not a formatted range: this lands in a spreadsheet, where a
        # string reads as text and cannot be sorted, plotted or compared
        rows[f"{name}_predicted"] = physical[name][:, 0]
        rows[f"{name}_lo68"] = physical[name][:, 1]
        rows[f"{name}_hi68"] = physical[name][:, 2]
    rows["distance_to_nearest"] = nearest
    rows["nearest_observed_row"] = nearest_index + 1
    rows["n_at_range_edge"] = pinned.sum(axis=1)
    rows["which_at_range_edge"] = [
        ", ".join(
            f"{input_names[j]}={'min' if at_lower[i, j] else 'max'}"
            for j in range(len(input_names))
            if pinned[i, j]
        )
        or "-"
        for i in range(len(conditions))
    ]
    candidates = pd.DataFrame(rows)

    probe_frames: list[pd.DataFrame] = []
    verdicts: list[str] = []
    for probe in probes:
        frame, verdict = _run_probe(
            config,
            model,
            probe,
            proposed=proposed,
            observed=observed,
            observed_Y_raw=np.asarray(observed_Y_raw, dtype=float),
            names=names,
            input_names=input_names,
            batch_sd=utility_sd,
            batch_utility=utility_mean,
            round_name=round_name,
            seed=resolved_seed,
        )
        probe_frames.append(frame)
        verdicts.append(verdict)

    return BatchReview(
        round_name=round_name.upper(),
        candidates=candidates,
        probes=(
            pd.concat(probe_frames, ignore_index=True) if probe_frames else pd.DataFrame()
        ),
        probe_verdicts=tuple(verdicts),
        notes=review_notes_from_config(config),
        findings=tuple(findings),
        model_warnings=tuple(model_warnings),
        context=dict(context or {}),
    )


def _run_probe(
    config: Mapping[str, Any],
    model: Any,
    probe: ProbeSpec,
    *,
    proposed: np.ndarray,
    observed: np.ndarray,
    observed_Y_raw: np.ndarray,
    names: list[str],
    input_names: list[str],
    batch_sd: np.ndarray,
    batch_utility: np.ndarray,
    round_name: str,
    seed: int,
) -> tuple[pd.DataFrame, str]:
    """Evaluate one counterfactual and say what its numbers mean.

    The comparison that matters is the *sd*, not the mean.  UCB rewards
    uncertainty, so a region the batch avoids while the model still calls it
    uncertain is simply losing a trade-off.  A region the batch avoids while the
    model calls it *certain* is a different thing: the model has resolved it, and
    if the data there is two observations that contradict each other, what it has
    resolved is an average rather than a fact.
    """
    if probe.column not in input_names:
        raise ValueError(
            f"Probe {probe.name!r} names {probe.column!r}, which is not a declared "
            f"input. Declared inputs: {input_names}."
        )
    position = input_names.index(probe.column)

    twins = proposed.copy()
    twins[:, position] = probe.value
    twin_mean, twin_sd = _utility_moments(
        config, model, twins, round_name=round_name, seed=seed
    )

    here = np.isclose(observed[:, position], probe.value, rtol=0.0, atol=1e-9)
    rows: list[dict[str, Any]] = []
    for i in range(len(twins)):
        row: dict[str, Any] = {
            "probe": probe.name,
            "kind": f"{round_name.upper()}_C{i + 1:02d} moved to {probe.column}={probe.value:g}",
        }
        for index, name in enumerate(names):
            row[f"{name}_utility"] = twin_mean[i, index]
            row[f"{name}_sd"] = twin_sd[i, index]
            row[f"{name}_utility_selected"] = batch_utility[i, index]
            row[f"{name}_sd_selected"] = batch_sd[i, index]
        rows.append(row)

    if here.any():
        observed_mean, observed_sd = _utility_moments(
            config, model, observed[here], round_name=round_name, seed=seed
        )
        observed_indices = np.flatnonzero(here)
        for slot, original_row in enumerate(observed_indices):
            row = {
                "probe": probe.name,
                "kind": f"observed row {original_row + 1} (already at {probe.column}={probe.value:g})",
            }
            for index, name in enumerate(names):
                row[f"{name}_utility"] = observed_mean[slot, index]
                row[f"{name}_sd"] = observed_sd[slot, index]
                row[f"{name}_utility_selected"] = np.nan
                row[f"{name}_sd_selected"] = np.nan
            row["measured"] = ", ".join(
                f"{name}={observed_Y_raw[original_row, index]:g}"
                for index, name in enumerate(names)
            )
            rows.append(row)

    frame = pd.DataFrame(rows)
    observed_count = int(here.sum())

    # Per objective, not pooled: the three utilities have sd on different scales
    # here (optoelectronic near 0.06 against thickness near 0.23), so a median over
    # the whole matrix can hide an objective whose uncertainty genuinely rises.
    per_objective: list[str] = []
    quieter: list[str] = []
    noisier: list[str] = []
    for index, name in enumerate(names):
        twin_u = float(np.median(twin_mean[:, index]))
        base_u = float(np.median(batch_utility[:, index]))
        twin_s = float(np.median(twin_sd[:, index]))
        base_s = float(np.median(batch_sd[:, index]))
        ratio = twin_s / base_s if base_s > 0 else float("inf")
        per_objective.append(
            f"{name}: utility {twin_u:.3f} vs {base_u:.3f} "
            f"({twin_u - base_u:+.3f}), sd {twin_s:.3f} vs {base_s:.3f} "
            f"(x{ratio:.2f})"
        )
        verdict_kind = classify_probe_objective(twin_u, base_u, ratio)
        if verdict_kind == "known_and_bad":
            quieter.append(name)
        elif verdict_kind == "uncertain_tradeoff":
            noisier.append(name)

    verdict = [
        f"PROBE '{probe.name}' ({probe.column} = {probe.value:g}), probed value "
        f"against the selected batch, medians -- {'; '.join(per_objective)}. "
        f"{observed_count} observation(s) already sit there. An sd ratio above "
        f"{SD_MATERIALITY_RATIO:g}x counts as materially more uncertain; anything "
        "below that is the same uncertainty at a worse predicted value."
    ]

    if quieter:
        verdict.append(
            f"For {', '.join(quieter)} the probed region scores WORSE with no "
            "materially greater uncertainty than the batch that was selected. Since "
            "UCB rewards uncertainty, that region is not being skipped because it "
            "looks unexplored -- it is being skipped because it looks KNOWN AND BAD."
        )
    if noisier:
        verdict.append(
            f"For {', '.join(noisier)} the region scores worse and does read as "
            "materially more uncertain, so there its absence is a trade-off against "
            "the other objectives rather than absorbed confidence."
        )
    if not quieter and not noisier:
        verdict.append(
            "The probed region does not score worse than the selected batch on any "
            "objective, so its absence from the batch is about the batch-spacing "
            "penalty rather than about the model's opinion of the region."
        )

    trend_driven = _objectives_with_mean_feature(config, probe.column)
    if trend_driven and quieter:
        verdict.append(
            f"Where that confidence comes from matters: {', '.join(trend_driven)} "
            f"carries {probe.column} in its mean function, so the prediction here is "
            "a fitted global trend evaluated at the edge of its range, not a local "
            "average of the nearby observations. The trend can be confident at an "
            "edge that holds almost no data, and it will be confidently wrong if the "
            "few points there are unreliable. That makes this a question about the "
            "measurements at the edge, not a settled model conclusion."
        )
    if probe.note:
        verdict.append(probe.note)
    return frame, " ".join(verdict)


def _objectives_with_mean_feature(
    config: Mapping[str, Any], column: str
) -> tuple[str, ...]:
    """Objectives whose structured mean uses ``column`` as a feature.

    A probe on such a column is asking the fitted trend to extrapolate, which is a
    different kind of claim from a GP interpolating between nearby points -- and
    worth naming, because a monotone trend is confident at a range edge by
    construction.
    """
    from .structured_mean import mean_spec_from_config

    names: list[str] = []
    for entry in config["objectives"]["specs"]:
        spec = mean_spec_from_config(entry)
        if spec is not None and any(f.column == column for f in spec.features):
            names.append(str(entry["name"]))
    return tuple(names)


# --------------------------------------------------------------------------- #
# writing it beside the worklist
# --------------------------------------------------------------------------- #


def write_review_sheet(
    candidate_workbook: str | Path, review: BatchReview, *, sheet_name: str = "Review"
) -> Path:
    """Add the review to the candidate workbook this round just wrote.

    Safe to open for writing, unlike the source workbook: this file was created by
    :func:`workbook_io.write_candidate_sheet` moments ago and contains no formulas,
    so openpyxl has no cached values to discard.  Never point this at
    ``Summary Table.xlsx``.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    path = Path(candidate_workbook)
    workbook = load_workbook(path)
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    sheet = workbook.create_sheet(sheet_name)

    bold = Font(bold=True)
    row = 1

    def heading(text: str) -> None:
        nonlocal row
        sheet.cell(row=row, column=1, value=text).font = bold
        row += 1

    def blank() -> None:
        nonlocal row
        row += 1

    def paragraph(text: str) -> None:
        nonlocal row
        cell = sheet.cell(row=row, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 14 * max(1, len(text) // 110 + 1)
        row += 1

    def table(frame: pd.DataFrame) -> None:
        nonlocal row
        for offset, column in enumerate(frame.columns, start=1):
            sheet.cell(row=row, column=offset, value=str(column)).font = bold
        row += 1
        for _, record in frame.iterrows():
            for offset, column in enumerate(frame.columns, start=1):
                value = record[column]
                if isinstance(value, (np.floating, np.integer)):
                    value = value.item()
                if isinstance(value, float) and not np.isfinite(value):
                    value = None
                sheet.cell(row=row, column=offset, value=value)
            row += 1

    heading(f"BATCH REVIEW - {review.round_name}")
    blank()
    for key, value in review.context.items():
        sheet.cell(row=row, column=1, value=key).font = bold
        sheet.cell(row=row, column=2, value=str(value))
        row += 1
    blank()

    if review.model_warnings:
        # above the table, for the same reason it is first in the text version
        heading("READ THIS BEFORE THE NUMBERS")
        for warning in review.model_warnings:
            paragraph(warning)
        blank()

    heading("PROPOSED CONDITIONS")
    table(review.candidates)
    blank()

    if not review.probes.empty:
        heading("PROBES")
        table(review.probes)
        blank()

    if review.probe_verdicts:
        heading("WHAT THE PROBES MEAN")
        for verdict in review.probe_verdicts:
            paragraph(verdict)
        blank()

    if review.notes:
        heading("NOTES")
        for note in review.notes:
            paragraph(note)
        blank()

    if review.findings:
        heading("CARRIED FROM THE MEASURED DATA")
        table(
            pd.DataFrame(
                {
                    "severity": [f.severity.value for f in _ordered(review.findings)],
                    "sample": [f.sample_id for f in _ordered(review.findings)],
                    "objective": [f.objective for f in _ordered(review.findings)],
                    "message": [f.message for f in _ordered(review.findings)],
                }
            )
        )
        blank()

    heading("APPROVAL")
    paragraph(NOT_APPROVED)

    sheet.column_dimensions["A"].width = 46
    for index in range(2, 40):
        sheet.column_dimensions[get_column_letter(index)].width = 16
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return path
