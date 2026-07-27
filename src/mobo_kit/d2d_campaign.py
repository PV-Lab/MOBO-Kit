"""Resolved D2D Step 2B debug campaign contract and data preparation.

This module deliberately separates historical observations from proposed grid
points.  The one approved off-grid control remains a continuous GP training
coordinate, while the strict Step 2A grid APIs continue to govern every new
candidate.
"""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
import json
from numbers import Real
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook

from .candidate_pool import physical_rows_to_grid_indices
from .data import x_normalizer_np
from .design import Design, build_design_from_config
from .objectives import ObjectiveSpec, ObjectiveTransform
from .workbook_schema import (
    WorkbookAudit,
    WorkbookInputExceptionRule,
    audit_campaign_workbook,
)


D2D_INPUT_COLUMNS = (
    "speed_1",
    "time_1",
    "speed_2",
    "time_2",
    "precur_conc",
    "precur_vol",
    "anneal_temp",
    "anneal_time",
    "anti_vol",
    "anti_time",
)
D2D_WORKBOOK_INPUT_COLUMNS = (
    "speed_1",
    "time_1",
    "speed_2",
    "time_2",
    "precur_conc",
    "precur_vol (uL)",
    "anneal_temp",
    "anneal_time",
    "anti_vol",
    "anti_time",
)
D2D_OBJECTIVE_COLUMNS = (
    "Uniformity score",
    "Optoelectronic score",
    "Thickness score",
)
D2D_BOUNDED_OBJECTIVE_COLUMNS = ("Uniformity score", "Thickness score")
D2D_OBJECTIVE_NAMES = (
    "uniformity_score",
    "optoelectronic_score",
    "thickness_score",
)
D2D_REFERENCE_POINT_UTILITY = np.asarray((-0.01, -10.0, -0.01), dtype=float)
D2D_DEBUG_WATERMARK = "DEBUG ONLY - NOT APPROVED FOR EXPERIMENT"
D2D_OBJECTIVE_CONTRACT_VERSION = "d2d-step2b-debug-objectives-v1"
D2D_APPROVED_INPUT_SPECS = (
    ("speed_1", "rpm", 1000.0, 6000.0, 500.0),
    ("time_1", "s", 5.0, 50.0, 5.0),
    ("speed_2", "rpm", 0.0, 5000.0, 500.0),
    ("time_2", "s", 10.0, 60.0, 5.0),
    ("precur_conc", "M", 1.0, 2.0, 0.05),
    ("precur_vol", "uL", 40.0, 200.0, 10.0),
    ("anneal_temp", "C", 100.0, 185.0, 5.0),
    ("anneal_time", "min", 10.0, 60.0, 5.0),
    ("anti_vol", "uL", 100.0, 200.0, 5.0),
    ("anti_time", "s", 9.0, 25.0, 2.0),
)


class D2DDebugConfigError(ValueError):
    """Raised when a debug configuration weakens the resolved contract."""


@dataclass(frozen=True)
class OffGridObservedException:
    sample_id: int
    input_name: str
    observed_value: float
    reason: str


@dataclass(frozen=True)
class ResolvedD2DDebugConfig:
    raw: dict[str, Any]
    design: Design
    config_hash: str
    workbook_profile: str
    workbook_sheet: str
    expected_content_range: str
    expected_workbook_sha256: str
    expected_sample_ids: tuple[int, ...]
    reference_point_utility: np.ndarray
    control_sample_ids: tuple[int, ...]
    include_control_in_debug_model: bool
    control_measurement_provenance_assumption: str
    off_grid_exceptions: tuple[OffGridObservedException, ...]
    seed: int
    r1_batch_size: int
    replicates_per_condition: int
    beta: float
    posterior_samples: int
    candidate_pool_size: int
    score_chunk_size: int
    local_radius: float
    min_batch_distance: float
    min_observed_distance: float
    dimension_weights: np.ndarray | None
    output_root: str
    debug_watermark: str


@dataclass(frozen=True)
class D2DTrainingData:
    campaign_id: str
    X_phys_all: np.ndarray
    X_norm_all: np.ndarray
    Y_objectives: np.ndarray
    sample_ids: np.ndarray
    row_roles: tuple[str, ...]
    include_in_model: np.ndarray
    on_grid_mask: np.ndarray
    off_grid_exceptions: tuple[OffGridObservedException, ...]
    on_grid_grid_indices: np.ndarray
    on_grid_sample_ids: np.ndarray
    measurement_provenance: str
    off_grid_exception_reasons: tuple[str, ...]
    warnings: tuple[str, ...]

    def manifest_frame(self) -> pd.DataFrame:
        """Return an auditable row-level training manifest."""
        frame = pd.DataFrame(
            {
                "campaign_id": self.campaign_id,
                "round": "R0",
                "sample_id": self.sample_ids,
                "candidate_id": [
                    f"R0-S{int(sample_id):02d}" for sample_id in self.sample_ids
                ],
                "row_role": self.row_roles,
                "replicate_group": [
                    f"R0-S{int(sample_id):02d}" for sample_id in self.sample_ids
                ],
                "replicate_number": 1,
                "candidate_status": "observed_debug_input",
                "include_in_model": self.include_in_model,
                "measurement_provenance": self.measurement_provenance,
                "on_grid": self.on_grid_mask,
                "off_grid_exception": ~self.on_grid_mask,
                "off_grid_exception_reason": self.off_grid_exception_reasons,
                "exclusion_reason": "",
            }
        )
        for index, name in enumerate(D2D_INPUT_COLUMNS):
            frame[name] = self.X_phys_all[:, index]
        for index, name in enumerate(D2D_OBJECTIVE_COLUMNS):
            frame[name] = self.Y_objectives[:, index]
        return frame


@dataclass(frozen=True)
class ReplicateAggregationResult:
    frame: pd.DataFrame
    warnings: tuple[str, ...]


def _canonical_json_hash(value: Mapping[str, Any]) -> str:
    payload = json.dumps(
        value, sort_keys=True, separators=(",", ":"), ensure_ascii=True
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _mapping(value: Any, *, field: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise D2DDebugConfigError(f"{field} must be a mapping.")
    return value


def _positive_int(value: Any, *, field: str) -> int:
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, (int, np.integer)):
        raise D2DDebugConfigError(f"{field} must be a positive integer.")
    result = int(value)
    if result <= 0:
        raise D2DDebugConfigError(f"{field} must be a positive integer.")
    return result


def _finite_number(value: Any, *, field: str, positive: bool = False) -> float:
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
        raise D2DDebugConfigError(f"{field} must be a real non-boolean number.")
    result = float(value)
    if not np.isfinite(result) or (positive and result <= 0):
        qualifier = "finite and strictly positive" if positive else "finite"
        raise D2DDebugConfigError(f"{field} must be {qualifier}.")
    return result


def load_d2d_debug_config(
    path: str | Path,
    *,
    allow_public_template: bool = False,
) -> ResolvedD2DDebugConfig:
    """Load and strictly validate a private config or explicit public fixture."""
    if not isinstance(allow_public_template, bool):
        raise TypeError("allow_public_template must be a boolean.")
    config_path = Path(path)
    if not config_path.is_file():
        raise FileNotFoundError(f"D2D debug config not found: {config_path}")
    with config_path.open("r", encoding="utf-8") as handle:
        raw = yaml.safe_load(handle)
    if not isinstance(raw, dict):
        raise D2DDebugConfigError("D2D debug config must contain a mapping.")
    template_only = raw.get("template_only", False)
    if not isinstance(template_only, bool):
        raise D2DDebugConfigError("template_only must be a boolean when supplied.")
    if template_only:
        if not allow_public_template:
            raise D2DDebugConfigError(
                "The tracked Step 2B config is a public template. Copy it into an "
                "ignored private location and supply runtime workbook identity fields."
            )
        if raw.get("campaign_id") != "public-d2d-template-debug":
            raise D2DDebugConfigError(
                "allow_public_template may resolve only the tracked public template."
            )
        workbook_template = _mapping(raw.get("workbook"), field="workbook")
        if workbook_template.get("expected_sha256") != (
            "REQUIRED_IN_IGNORED_PRIVATE_CONFIG"
        ):
            raise D2DDebugConfigError(
                "The public template workbook identity placeholder was modified."
            )
        workbook_template["expected_sha256"] = _canonical_json_hash(
            {
                "schema_version": raw.get("schema_version"),
                "campaign_id": raw.get("campaign_id"),
                "identity": "runtime-generated-public-synthetic-fixture",
            }
        )
        raw["template_only"] = False
    if raw.get("schema_version") != "d2d-step2b-debug-1":
        raise D2DDebugConfigError(
            "schema_version must be exactly 'd2d-step2b-debug-1'."
        )
    campaign_id = raw.get("campaign_id")
    if not isinstance(campaign_id, str) or not campaign_id.strip():
        raise D2DDebugConfigError("campaign_id must be a nonblank string.")

    required_flags = {
        "run_mode": "debug",
        "debug_run_authorized": True,
        "approved_for_production": False,
        "approved_for_experiment": False,
    }
    for field, expected in required_flags.items():
        value = raw.get(field)
        exact_boolean = not isinstance(expected, bool) or (
            isinstance(value, bool) and value is expected
        )
        if value != expected or not exact_boolean:
            raise D2DDebugConfigError(
                f"{field} must be exactly {expected!r} for a Step 2B debug run."
            )

    inputs = raw.get("inputs")
    if not isinstance(inputs, list) or len(inputs) != len(D2D_APPROVED_INPUT_SPECS):
        raise D2DDebugConfigError(
            "inputs must contain the exact ten approved Step 2A grid definitions."
        )
    for index, (item, approved) in enumerate(zip(inputs, D2D_APPROVED_INPUT_SPECS)):
        input_spec = _mapping(item, field=f"inputs[{index}]")
        name, unit, start, stop, step = approved
        if input_spec.get("name") != name or input_spec.get("unit") != unit:
            raise D2DDebugConfigError(
                f"inputs[{index}] must preserve approved name/unit {name!r}/{unit!r}."
            )
        for field, approved_value in (
            ("start", start),
            ("stop", stop),
            ("step", step),
        ):
            actual = _finite_number(
                input_spec.get(field), field=f"inputs[{index}].{field}"
            )
            if actual != approved_value:
                raise D2DDebugConfigError(
                    f"inputs[{index}].{field} must preserve the approved Step 2A "
                    f"grid value {approved_value}."
                )
    design = build_design_from_config(raw)

    workbook = _mapping(raw.get("workbook"), field="workbook")
    if workbook.get("profile") != "d2d_summary_v3_scores":
        raise D2DDebugConfigError("workbook.profile must be 'd2d_summary_v3_scores'.")
    if workbook.get("sheet") != "Sheet1":
        raise D2DDebugConfigError("workbook.sheet must be 'Sheet1'.")
    if workbook.get("expected_content_range") != "A1:AI20":
        raise D2DDebugConfigError(
            "workbook.expected_content_range must be exactly 'A1:AI20'."
        )
    if workbook.get("input_aliases") != {"precur_vol (uL)": "precur_vol"}:
        raise D2DDebugConfigError(
            "workbook.input_aliases must preserve the resolved precursor-volume alias."
        )
    if workbook.get("sample_id_column") != "Sample number":
        raise D2DDebugConfigError(
            "workbook.sample_id_column must be exactly 'Sample number'."
        )
    expected_hash = str(workbook.get("expected_sha256", "")).upper()
    if len(expected_hash) != 64 or any(
        c not in "0123456789ABCDEF" for c in expected_hash
    ):
        raise D2DDebugConfigError(
            "workbook.expected_sha256 must be a SHA-256 hex digest."
        )
    expected_ids_raw = workbook.get("expected_sample_ids")
    if not isinstance(expected_ids_raw, list) or not expected_ids_raw:
        raise D2DDebugConfigError(
            "workbook.expected_sample_ids must be a nonempty list of unique integers."
        )
    if any(
        isinstance(value, (bool, np.bool_)) or not isinstance(value, (int, np.integer))
        for value in expected_ids_raw
    ):
        raise D2DDebugConfigError(
            "workbook.expected_sample_ids must contain only integers."
        )
    expected_ids = tuple(int(value) for value in expected_ids_raw)
    if len(set(expected_ids)) != len(expected_ids):
        raise D2DDebugConfigError(
            "workbook.expected_sample_ids must contain unique integers."
        )

    expected_objectives = [
        {
            "name": "uniformity_score",
            "source_column": "Uniformity score",
            "excel_column": "Z",
            "direction": "maximize",
            "utility_transform": "identity",
            "support_formula": "Coverage * (1 - Uniformity) * Phase purity",
            "mismatch_policy": "warn_use_supplied",
        },
        {
            "name": "optoelectronic_score",
            "source_column": "Optoelectronic score",
            "excel_column": "AA",
            "direction": "maximize",
            "utility_transform": "identity",
            "support_formula": "log10((PL - Implied Voc (Max)) * Photoconductance (Max))",
            "mismatch_policy": "error",
        },
        {
            "name": "thickness_score",
            "source_column": "Thickness score",
            "excel_column": "AB",
            "direction": "maximize",
            "utility_transform": "identity",
            "support_formula": "exp(-((mean(valid T1:T4) - 650.0) / 250.0)^2)",
            "target_nm": 650.0,
            "scale_nm": 250.0,
            "exponent_factor": 1.0,
            "exclude_columns": ["T anom"],
            "mismatch_policy": "error",
        },
    ]
    objectives = raw.get("objectives")
    if not isinstance(objectives, list) or len(objectives) != 3:
        raise D2DDebugConfigError("objectives must contain exactly three mappings.")
    for index, (item, expected) in enumerate(zip(objectives, expected_objectives)):
        objective = _mapping(item, field=f"objectives[{index}]")
        for field, expected_value in expected.items():
            if objective.get(field) != expected_value:
                raise D2DDebugConfigError(
                    f"objectives[{index}].{field} must be {expected_value!r}."
                )

    expected_qc_policy = {
        "final_scores_required": True,
        "complete_case_rule": "require_all_three_final_scores",
        "failed_measurement_rule": "block_row",
        "outlier_rule": "report_do_not_auto_remove",
        "uniformity_known_mismatch": "warn_and_continue_debug",
        "optoelectronic_mismatch": "fail",
        "thickness_mismatch": "fail",
    }
    if raw.get("qc_policy") != expected_qc_policy:
        raise D2DDebugConfigError(
            "qc_policy must preserve the resolved Step 2B score-validation policy."
        )
    expected_ignored_columns = [
        "Stability score?",
        "Total combination - addition",
        "Total combination - multiplied",
        "Uniformity score absolute difference",
        "Optoelectronic score absolute difference",
        "Thickness absolute difference",
        "Total score absolute difference",
    ]
    if raw.get("ignored_model_columns") != expected_ignored_columns:
        raise D2DDebugConfigError(
            "ignored_model_columns must preserve the resolved AC:AI exclusion list."
        )

    reference = np.asarray(raw.get("reference_point_utility"), dtype=float)
    if reference.shape != (3,) or not np.array_equal(
        reference, D2D_REFERENCE_POINT_UTILITY
    ):
        raise D2DDebugConfigError(
            "reference_point_utility must be exactly [-0.01, -10.0, -0.01]."
        )
    if raw.get("constraints") != []:
        raise D2DDebugConfigError("constraints must be an explicit empty list.")

    r0 = _mapping(raw.get("r0"), field="r0")
    condition_count = _positive_int(
        r0.get("condition_count"), field="r0.condition_count"
    )
    if condition_count != len(expected_ids):
        raise D2DDebugConfigError(
            "r0.condition_count must match workbook.expected_sample_ids."
        )
    control_ids_raw = r0.get("control_sample_ids")
    if not isinstance(control_ids_raw, list) or not control_ids_raw:
        raise D2DDebugConfigError(
            "r0.control_sample_ids must be a nonempty list of sample identifiers."
        )
    if any(
        isinstance(value, (bool, np.bool_)) or not isinstance(value, (int, np.integer))
        for value in control_ids_raw
    ):
        raise D2DDebugConfigError("r0.control_sample_ids must contain only integers.")
    control_ids = tuple(int(value) for value in control_ids_raw)
    if len(set(control_ids)) != len(control_ids) or not set(control_ids).issubset(
        expected_ids
    ):
        raise D2DDebugConfigError(
            "r0.control_sample_ids must be unique members of expected_sample_ids."
        )
    if r0.get("include_control_in_debug_model") is not True:
        raise D2DDebugConfigError(
            "The control must be included in the primary debug model."
        )
    provenance = r0.get("control_measurement_provenance_assumption")
    if not isinstance(provenance, str) or not provenance.strip():
        raise D2DDebugConfigError(
            "The control measurement provenance assumption must be a nonblank string."
        )
    raw_exceptions = r0.get("off_grid_observed_exceptions")
    if not isinstance(raw_exceptions, list):
        raise D2DDebugConfigError("r0.off_grid_observed_exceptions must be a list.")
    exceptions: list[OffGridObservedException] = []
    exception_keys: set[tuple[int, str]] = set()
    for index, item in enumerate(raw_exceptions):
        exception_raw = _mapping(
            item, field=f"r0.off_grid_observed_exceptions[{index}]"
        )
        sample_id_raw = exception_raw.get("sample_id")
        if isinstance(sample_id_raw, (bool, np.bool_)) or not isinstance(
            sample_id_raw, (int, np.integer)
        ):
            raise D2DDebugConfigError(
                "Off-grid exception sample_id must be an integer."
            )
        sample_id = int(sample_id_raw)
        input_name = str(exception_raw.get("input_name", "")).strip()
        if sample_id not in expected_ids or input_name not in design.names:
            raise D2DDebugConfigError(
                "Off-grid exceptions must reference a configured sample and input."
            )
        if "observed_value" in exception_raw:
            observed_value = _finite_number(
                exception_raw.get("observed_value"), field="off-grid observed_value"
            )
        elif (
            exception_raw.get("observed_value_strategy")
            == "midpoint_between_first_two_grid_values"
        ):
            dimension = design.names.index(input_name)
            grid = np.asarray(design.var_array[dimension], dtype=float)
            if grid.size < 2:
                raise D2DDebugConfigError(
                    "Synthetic midpoint exceptions require at least two grid values."
                )
            observed_value = float((grid[0] + grid[1]) / 2.0)
        else:
            raise D2DDebugConfigError(
                "Off-grid exceptions require observed_value or the supported "
                "synthetic midpoint strategy."
            )
        dimension = design.names.index(input_name)
        grid = np.asarray(design.var_array[dimension], dtype=float)
        if (
            observed_value < design.lowers[dimension]
            or observed_value > design.uppers[dimension]
        ):
            raise D2DDebugConfigError(
                "Off-grid exception values must remain in bounds."
            )
        if np.any(np.isclose(observed_value, grid, rtol=0.0, atol=1e-9)):
            raise D2DDebugConfigError("Off-grid exception values must not be on-grid.")
        reason = str(exception_raw.get("reason", "")).strip()
        if not reason or exception_raw.get("retain_observed_value") is not True:
            raise D2DDebugConfigError(
                "Off-grid exceptions require a reason and retain_observed_value=true."
            )
        key = (sample_id, input_name)
        if key in exception_keys:
            raise D2DDebugConfigError(f"Duplicate off-grid exception for {key!r}.")
        exception_keys.add(key)
        exceptions.append(
            OffGridObservedException(
                sample_id=sample_id,
                input_name=input_name,
                observed_value=observed_value,
                reason=reason,
            )
        )

    r1 = _mapping(raw.get("r1"), field="r1")
    if r1.get("method") != "ucb_hvi":
        raise D2DDebugConfigError("r1.method must be 'ucb_hvi'.")
    expected_r1 = {
        "method": "ucb_hvi",
        "batch_size_unique_conditions": 5,
        "replicates_per_condition": 3,
        "beta": 4.0,
        "posterior_samples": 256,
        "candidate_pool_size": 10000,
        "score_chunk_size": 512,
    }
    if r1 != expected_r1:
        raise D2DDebugConfigError(
            "r1 must preserve the resolved Step 2B baseline acquisition settings."
        )
    expected_r2 = {
        "method": "qlognehvi",
        "batch_size_unique_conditions": 3,
        "replicates_per_condition": 3,
        "mc_samples": 128,
        "candidate_pool_size": 5000,
        "sequential_pending": True,
    }
    if raw.get("r2_test_only") != expected_r2:
        raise D2DDebugConfigError(
            "r2_test_only must preserve the resolved synthetic-only settings."
        )
    local = _mapping(raw.get("local_penalization"), field="local_penalization")
    if local.get("distance_metric") != "normalized_euclidean":
        raise D2DDebugConfigError(
            "local_penalization.distance_metric must be 'normalized_euclidean'."
        )
    if local.get("allow_hard_distance_relaxation") is not False:
        raise D2DDebugConfigError("Hard-distance relaxation must remain false.")
    weights_raw = local.get("dimension_weights")
    weights = None if weights_raw is None else np.asarray(weights_raw, dtype=float)
    if weights is not None and (
        weights.shape != (10,)
        or not np.all(np.isfinite(weights))
        or np.any(weights <= 0)
    ):
        raise D2DDebugConfigError(
            "local_penalization.dimension_weights must be null or ten positive values."
        )
    expected_local = {
        "distance_metric": "normalized_euclidean",
        "radius": 0.25,
        "min_batch_distance": 0.15,
        "min_observed_distance": 0.0,
        "dimension_weights": None,
        "allow_hard_distance_relaxation": False,
    }
    if local != expected_local:
        raise D2DDebugConfigError(
            "local_penalization must preserve the resolved Step 2B baseline settings."
        )

    outputs = _mapping(raw.get("outputs"), field="outputs")
    if outputs.get("root") != "local_outputs/d2d_step2b_debug":
        raise D2DDebugConfigError(
            "outputs.root must be exactly 'local_outputs/d2d_step2b_debug'."
        )
    if outputs.get("debug_watermark") != D2D_DEBUG_WATERMARK:
        raise D2DDebugConfigError(
            f"outputs.debug_watermark must be {D2D_DEBUG_WATERMARK!r}."
        )
    if outputs.get("write_source_workbook") is not False:
        raise D2DDebugConfigError("outputs.write_source_workbook must be false.")
    reproducibility = _mapping(raw.get("reproducibility"), field="reproducibility")
    expected_reproducibility = {
        "seed": 73,
        "record_git_commit": True,
        "record_environment_versions": True,
        "record_resolved_config_hash": True,
        "record_workbook_hash": True,
    }
    if reproducibility != expected_reproducibility:
        raise D2DDebugConfigError(
            "reproducibility must preserve seed 73 and all provenance records."
        )
    resolved = ResolvedD2DDebugConfig(
        raw=raw,
        design=design,
        config_hash=_canonical_json_hash(raw),
        workbook_profile="d2d_summary_v3_scores",
        workbook_sheet="Sheet1",
        expected_content_range=str(workbook.get("expected_content_range")),
        expected_workbook_sha256=expected_hash,
        expected_sample_ids=expected_ids,
        reference_point_utility=reference.copy(),
        control_sample_ids=tuple(int(value) for value in control_ids),
        include_control_in_debug_model=True,
        control_measurement_provenance_assumption=str(provenance),
        off_grid_exceptions=tuple(exceptions),
        seed=_positive_int(reproducibility.get("seed"), field="reproducibility.seed"),
        r1_batch_size=_positive_int(
            r1.get("batch_size_unique_conditions"),
            field="r1.batch_size_unique_conditions",
        ),
        replicates_per_condition=_positive_int(
            r1.get("replicates_per_condition"), field="r1.replicates_per_condition"
        ),
        beta=_finite_number(r1.get("beta"), field="r1.beta", positive=True),
        posterior_samples=_positive_int(
            r1.get("posterior_samples"), field="r1.posterior_samples"
        ),
        candidate_pool_size=_positive_int(
            r1.get("candidate_pool_size"), field="r1.candidate_pool_size"
        ),
        score_chunk_size=_positive_int(
            r1.get("score_chunk_size"), field="r1.score_chunk_size"
        ),
        local_radius=_finite_number(
            local.get("radius"), field="local_penalization.radius", positive=True
        ),
        min_batch_distance=_finite_number(
            local.get("min_batch_distance"),
            field="local_penalization.min_batch_distance",
        ),
        min_observed_distance=_finite_number(
            local.get("min_observed_distance"),
            field="local_penalization.min_observed_distance",
        ),
        dimension_weights=None if weights is None else weights.copy(),
        output_root=str(outputs.get("root")),
        debug_watermark=D2D_DEBUG_WATERMARK,
    )
    if tuple(resolved.design.names) != D2D_INPUT_COLUMNS:
        raise D2DDebugConfigError(
            f"inputs must use the exact canonical order {D2D_INPUT_COLUMNS}."
        )
    if resolved.r1_batch_size != 5 or resolved.replicates_per_condition != 3:
        raise D2DDebugConfigError(
            "R1 must select five conditions with three replicates each."
        )
    if resolved.min_batch_distance < 0 or resolved.min_observed_distance < 0:
        raise D2DDebugConfigError("Configured minimum distances must be non-negative.")
    return resolved


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _normalize_missing(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and not value.replace("\u00a0", " ").strip():
        return None
    return value


def _numeric_sample_id(value: Any) -> int | None:
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
        return None
    number = float(value)
    if not np.isfinite(number) or not number.is_integer():
        return None
    return int(number)


def load_d2d_workbook_frame(
    path: str | Path,
    *,
    expected_profile: str = "d2d_summary_v3_scores",
    expected_sample_ids: Sequence[int] | None = None,
    allowed_input_exceptions: Sequence[OffGridObservedException] = (),
) -> tuple[pd.DataFrame, WorkbookAudit]:
    """Read only the numeric v3 sample rows and cached values from the workbook."""
    workbook_path = Path(path)
    audit = audit_campaign_workbook(
        workbook_path,
        allowed_input_exceptions=tuple(
            WorkbookInputExceptionRule(
                sample_id=exception.sample_id,
                input_name=exception.input_name,
                observed_value=exception.observed_value,
                reason=exception.reason,
            )
            for exception in allowed_input_exceptions
        ),
    )
    if audit.profile != expected_profile:
        raise ValueError(
            f"Expected workbook profile {expected_profile!r}; found {audit.profile!r}."
        )
    if audit.active_sheet != "Sheet1" or audit.used_range != "A1:AI20":
        raise ValueError(
            "The v3 workbook must use Sheet1 with content range A1:AI20; "
            f"found {audit.active_sheet!r} and {audit.used_range!r}."
        )
    if audit.input_rows_valid is not True:
        raise ValueError(
            "The v3 workbook input-row audit failed: "
            f"{audit.input_validation_errors or ['unspecified input error']}."
        )
    headers = tuple(audit.raw_headers)
    if len(headers) != 35 or any(header is None for header in headers):
        raise ValueError("The v3 workbook must contain 35 nonblank unique headers.")
    if len(set(headers)) != 35:
        raise ValueError("The v3 workbook headers must be unique.")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Sheet1"]
        rows: list[list[Any]] = []
        for values in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=35,
            values_only=True,
        ):
            sample_id = _numeric_sample_id(values[0])
            if sample_id is None:
                continue
            normalized = [_normalize_missing(value) for value in values]
            normalized[0] = sample_id
            rows.append(normalized)
    finally:
        workbook.close()
    frame = pd.DataFrame(rows, columns=list(headers))
    observed_ids = tuple(int(value) for value in frame["Sample number"].tolist())
    if not observed_ids or len(set(observed_ids)) != len(observed_ids):
        raise ValueError(
            "The v3 workbook must contain unique numeric sample identifiers; "
            f"found {observed_ids}."
        )
    if expected_sample_ids is not None and observed_ids != tuple(expected_sample_ids):
        raise ValueError(
            f"Expected ordered sample identifiers {tuple(expected_sample_ids)}; "
            f"found {observed_ids}."
        )
    return frame, audit


def build_d2d_objective_transform() -> ObjectiveTransform:
    """Return the resolved ordered identity/maximize objective contract."""
    specs = [
        ObjectiveSpec(
            name=name,
            goal="maximize",
            transform="identity",
            source_column=source,
        )
        for name, source in zip(D2D_OBJECTIVE_NAMES, D2D_OBJECTIVE_COLUMNS)
    ]
    return ObjectiveTransform(specs, version=D2D_OBJECTIVE_CONTRACT_VERSION)


def _numeric_matrix(
    frame: pd.DataFrame, columns: Sequence[str], *, label: str
) -> np.ndarray:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise ValueError(f"Missing {label} column(s): {missing}.")
    numeric = frame.loc[:, list(columns)].apply(pd.to_numeric, errors="coerce")
    values = numeric.to_numpy(dtype=float)
    if not np.all(np.isfinite(values)):
        rows, cols = np.where(~np.isfinite(values))
        details = [(int(row), columns[int(col)]) for row, col in zip(rows, cols)]
        raise ValueError(
            f"{label} values must be complete and finite; invalid {details}."
        )
    return values


def prepare_d2d_training_data(
    frame: pd.DataFrame,
    config: ResolvedD2DDebugConfig,
    *,
    include_control: bool = True,
) -> D2DTrainingData:
    """Prepare control-aware, fixed-bound training arrays from workbook rows."""
    if not isinstance(frame, pd.DataFrame):
        raise TypeError("frame must be a pandas DataFrame.")
    sample_values = _numeric_matrix(
        frame, ("Sample number",), label="sample identifier"
    )[:, 0]
    if not np.all(sample_values == np.floor(sample_values)):
        raise ValueError("Sample identifiers must be integers.")
    sample_ids = sample_values.astype(int)
    if tuple(sample_ids.tolist()) != config.expected_sample_ids:
        raise ValueError(
            f"Expected ordered sample IDs {config.expected_sample_ids}; found {sample_ids.tolist()}."
        )
    X_phys = _numeric_matrix(frame, D2D_WORKBOOK_INPUT_COLUMNS, label="input")
    duplicate_inputs = pd.DataFrame(X_phys).duplicated(keep=False).to_numpy()
    if np.any(duplicate_inputs):
        duplicate_samples = sample_ids[duplicate_inputs].tolist()
        raise ValueError(
            "Observed training recipes must be unique; duplicate input rows were "
            f"found for samples {duplicate_samples}."
        )
    lower = np.asarray(config.design.lowers, dtype=float)
    upper = np.asarray(config.design.uppers, dtype=float)
    out_of_bounds = np.where((X_phys < lower) | (X_phys > upper))
    if out_of_bounds[0].size:
        details = [
            (int(sample_ids[row]), D2D_INPUT_COLUMNS[col], float(X_phys[row, col]))
            for row, col in zip(*out_of_bounds)
        ]
        raise ValueError(
            f"Observed inputs must remain within configured bounds: {details}."
        )

    exceptions_by_key = {
        (exception.sample_id, exception.input_name): exception
        for exception in config.off_grid_exceptions
    }
    on_grid_mask = np.ones(X_phys.shape[0], dtype=bool)
    encountered: list[OffGridObservedException] = []
    for row_index, (sample_id, row) in enumerate(zip(sample_ids, X_phys)):
        mismatches: list[tuple[int, str, float]] = []
        for column, (name, grid) in enumerate(
            zip(D2D_INPUT_COLUMNS, config.design.var_array)
        ):
            if not np.any(np.isclose(row[column], grid, rtol=0.0, atol=1e-9)):
                mismatches.append((column, name, float(row[column])))
        if not mismatches:
            continue
        on_grid_mask[row_index] = False
        for _, name, observed in mismatches:
            exception = exceptions_by_key.get((int(sample_id), name))
            if exception is None or not np.isclose(
                observed, exception.observed_value, rtol=0.0, atol=1e-12
            ):
                raise ValueError(
                    f"Sample {sample_id} has an unapproved off-grid value "
                    f"{name}={observed}."
                )
            encountered.append(exception)
    if set(encountered) != set(config.off_grid_exceptions):
        raise ValueError(
            "The configured off-grid exception was not found exactly in the data."
        )

    grid_indices = physical_rows_to_grid_indices(X_phys[on_grid_mask], config.design)
    X_norm = x_normalizer_np(X_phys, config.design)
    if np.any(X_norm < 0.0) or np.any(X_norm > 1.0):
        raise RuntimeError("Bounded observations failed fixed input normalization.")
    objectives = _numeric_matrix(frame, D2D_OBJECTIVE_COLUMNS, label="objective")
    for objective_name in D2D_BOUNDED_OBJECTIVE_COLUMNS:
        objective_index = D2D_OBJECTIVE_COLUMNS.index(objective_name)
        invalid_rows = np.flatnonzero(
            (objectives[:, objective_index] < 0.0)
            | (objectives[:, objective_index] > 1.0)
        )
        if invalid_rows.size:
            details = [
                (int(sample_ids[row]), float(objectives[row, objective_index]))
                for row in invalid_rows
            ]
            raise ValueError(
                f"Authoritative {objective_name} values must remain in [0, 1]; "
                f"violations: {details}."
            )
    if not np.all(objectives > config.reference_point_utility):
        failing = np.argwhere(objectives <= config.reference_point_utility)
        details = [
            (
                int(sample_ids[row]),
                D2D_OBJECTIVE_COLUMNS[col],
                float(objectives[row, col]),
            )
            for row, col in failing
        ]
        raise ValueError(
            "Every objective row must strictly dominate the fixed reference point; "
            f"violations: {details}."
        )
    roles = tuple(
        "control" if int(sample_id) in config.control_sample_ids else "r0_lhs"
        for sample_id in sample_ids
    )
    included = np.ones(sample_ids.shape[0], dtype=bool)
    if not include_control:
        included &= ~np.isin(sample_ids, config.control_sample_ids)
    warnings = (
        "Configured control observations are included under the declared "
        "measurement-provenance assumption.",
        *tuple(
            "A configured observed-only off-grid input exception is retained for "
            f"sample {exception.sample_id} and input {exception.input_name!r}; "
            "new candidates remain on the approved grid."
            for exception in config.off_grid_exceptions
        ),
    )
    return D2DTrainingData(
        campaign_id=str(config.raw.get("campaign_id")),
        X_phys_all=X_phys,
        X_norm_all=X_norm,
        Y_objectives=objectives,
        sample_ids=sample_ids,
        row_roles=roles,
        include_in_model=included,
        on_grid_mask=on_grid_mask,
        off_grid_exceptions=tuple(encountered),
        on_grid_grid_indices=grid_indices,
        on_grid_sample_ids=sample_ids[on_grid_mask],
        measurement_provenance=config.control_measurement_provenance_assumption,
        off_grid_exception_reasons=tuple(
            next(
                (
                    exception.reason
                    for exception in encountered
                    if exception.sample_id == int(sample_id)
                ),
                "",
            )
            for sample_id in sample_ids
        ),
        warnings=warnings,
    )


def expand_candidates_to_replicates(
    candidates: pd.DataFrame,
    *,
    replicates_per_condition: int = 3,
    round_name: str = "R1",
) -> pd.DataFrame:
    """Expand unique candidate conditions to explicitly numbered debug executions."""
    if not isinstance(candidates, pd.DataFrame):
        raise TypeError("candidates must be a pandas DataFrame.")
    replicate_count = _positive_int(
        replicates_per_condition, field="replicates_per_condition"
    )
    _numeric_matrix(candidates, D2D_INPUT_COLUMNS, label="candidate input")
    if candidates.shape[0] == 0:
        raise ValueError("At least one unique candidate is required.")
    if candidates.loc[:, list(D2D_INPUT_COLUMNS)].duplicated().any():
        raise ValueError(
            "Unique candidate conditions must not contain duplicate inputs."
        )
    if not isinstance(round_name, str) or not round_name.strip():
        raise ValueError("round_name must be a nonblank string.")
    if "candidate_id" in candidates.columns:
        candidate_ids = candidates["candidate_id"].map(str).map(str.strip)
        if candidate_ids.eq("").any() or candidates["candidate_id"].isna().any():
            raise ValueError("candidate_id values must be nonblank.")
        if candidate_ids.duplicated().any():
            raise ValueError("candidate_id values must be unique across conditions.")
    else:
        candidate_ids = pd.Series(
            [
                f"{round_name}-C{position:02d}"
                for position in range(1, len(candidates) + 1)
            ],
            index=candidates.index,
        )
    rows: list[dict[str, Any]] = []
    if "campaign_id" in candidates.columns:
        campaign_ids = candidates["campaign_id"].map(str).map(str.strip)
        if candidates["campaign_id"].isna().any() or campaign_ids.eq("").any():
            raise ValueError("campaign_id values must be nonblank when supplied.")
        if campaign_ids.nunique() != 1:
            raise ValueError("campaign_id must be consistent across candidates.")
        campaign_id = str(campaign_ids.iloc[0])
    else:
        campaign_id = "public-debug-campaign"
    for position, (_, candidate) in enumerate(candidates.iterrows(), start=1):
        candidate_id = str(candidate_ids.iloc[position - 1])
        for replicate in range(1, replicate_count + 1):
            row: dict[str, Any] = {
                "campaign_id": campaign_id,
                "round": round_name,
                "sample_id": pd.NA,
                "execution_id": f"{candidate_id}-R{replicate}",
                "candidate_id": candidate_id,
                "replicate_group": candidate_id,
                "replicate_number": replicate,
                "row_role": "candidate_replicate",
                "candidate_status": D2D_DEBUG_WATERMARK,
                "debug_only": True,
                "approved_for_experiment": False,
                "include_in_model": False,
                "measurement_provenance": "pending_measurement",
                "off_grid_exception": False,
                "exclusion_reason": "awaiting_measurement",
            }
            for name in D2D_INPUT_COLUMNS:
                row[name] = float(candidate[name])
            for objective in D2D_OBJECTIVE_COLUMNS:
                row[objective] = pd.NA
            rows.append(row)
    expanded = pd.DataFrame(rows)
    if not expanded["execution_id"].is_unique:
        raise RuntimeError("Expanded candidate execution IDs must be unique.")
    return expanded


def aggregate_replicate_objectives(
    records: pd.DataFrame,
    *,
    group_column: str = "replicate_group",
    objective_columns: Sequence[str] = D2D_OBJECTIVE_COLUMNS,
    input_columns: Sequence[str] = D2D_INPUT_COLUMNS,
    expected_replicates: int = 3,
    sample_id_column: str = "execution_id",
) -> ReplicateAggregationResult:
    """Aggregate physical repeats to one condition-level model observation."""
    if not isinstance(records, pd.DataFrame):
        raise TypeError("records must be a pandas DataFrame.")
    expected_count = _positive_int(expected_replicates, field="expected_replicates")
    required = [
        group_column,
        sample_id_column,
        "replicate_number",
        *input_columns,
        *objective_columns,
    ]
    missing = [column for column in required if column not in records.columns]
    if missing:
        raise ValueError(f"Missing replicate column(s): {missing}.")
    source_identifiers = records[sample_id_column]
    normalized_identifiers = source_identifiers.map(str).map(str.strip)
    if source_identifiers.isna().any() or normalized_identifiers.eq("").any():
        raise ValueError(f"{sample_id_column} values must be nonblank.")
    if normalized_identifiers.duplicated().any():
        raise ValueError(f"{sample_id_column} values must be unique.")
    output_rows: list[dict[str, Any]] = []
    warnings: list[str] = []
    for group, group_frame in records.groupby(group_column, sort=False, dropna=False):
        if pd.isna(group) or not str(group).strip():
            raise ValueError("replicate_group values must be nonblank.")
        input_values = _numeric_matrix(
            group_frame, input_columns, label="replicate input"
        )
        if not np.all(input_values == input_values[0]):
            raise ValueError(
                f"Replicate group {group!r} contains different requested input conditions."
            )
        row: dict[str, Any] = {group_column: group}
        for column, value in zip(input_columns, input_values[0]):
            row[column] = float(value)
        replicate_numbers = pd.to_numeric(
            group_frame["replicate_number"], errors="coerce"
        ).to_numpy(dtype=float)
        if (
            not np.all(np.isfinite(replicate_numbers))
            or not np.all(replicate_numbers == np.floor(replicate_numbers))
            or np.any(replicate_numbers < 1)
            or np.any(replicate_numbers > expected_count)
        ):
            raise ValueError(
                f"Replicate group {group!r} has invalid replicate_number values."
            )
        if np.unique(replicate_numbers).size != replicate_numbers.size:
            raise ValueError(
                f"Replicate group {group!r} has duplicate replicate_number values."
            )
        source_ids = group_frame[sample_id_column].astype(str).tolist()
        row["source_sample_ids"] = "|".join(source_ids)
        objective_counts: list[int] = []
        for objective in objective_columns:
            values = (
                pd.to_numeric(group_frame[objective], errors="coerce")
                .dropna()
                .to_numpy(dtype=float)
            )
            if not np.all(np.isfinite(values)):
                raise ValueError(
                    f"Replicate group {group!r} objective {objective!r} contains non-finite values."
                )
            if objective in D2D_BOUNDED_OBJECTIVE_COLUMNS and np.any(
                (values < 0.0) | (values > 1.0)
            ):
                raise ValueError(
                    f"Replicate group {group!r} objective {objective!r} must "
                    "remain in [0, 1]."
                )
            count = int(values.size)
            objective_counts.append(count)
            mean = float(np.mean(values)) if count else np.nan
            std = float(np.std(values, ddof=1)) if count > 1 else np.nan
            sem = float(std / np.sqrt(count)) if count > 1 else np.nan
            row[f"{objective}_mean"] = mean
            row[f"{objective}_sample_std"] = std
            row[f"{objective}_count"] = count
            row[f"{objective}_standard_error"] = sem
        complete = group_frame.shape[0] == expected_count and all(
            count == expected_count for count in objective_counts
        )
        row["complete_replicate_set"] = complete
        row["include_in_next_model"] = complete
        if not complete:
            warnings.append(
                f"Replicate group {group!r} is incomplete: rows={group_frame.shape[0]}, "
                f"objective_counts={objective_counts}, expected={expected_count}."
            )
        output_rows.append(row)
    return ReplicateAggregationResult(pd.DataFrame(output_rows), tuple(warnings))


def combine_r0_and_aggregated_r1(
    r0_training: D2DTrainingData,
    r1_aggregated: ReplicateAggregationResult,
) -> tuple[np.ndarray, np.ndarray]:
    """Build the future 20-condition R2 training matrices without triplicate inflation."""
    complete = r1_aggregated.frame[
        r1_aggregated.frame["include_in_next_model"].astype(bool)
    ]
    X_r1 = _numeric_matrix(complete, D2D_INPUT_COLUMNS, label="aggregated R1 input")
    mean_columns = tuple(f"{objective}_mean" for objective in D2D_OBJECTIVE_COLUMNS)
    Y_r1 = _numeric_matrix(complete, mean_columns, label="aggregated R1 objective")
    for objective_name in D2D_BOUNDED_OBJECTIVE_COLUMNS:
        objective_index = D2D_OBJECTIVE_COLUMNS.index(objective_name)
        if np.any((Y_r1[:, objective_index] < 0.0) | (Y_r1[:, objective_index] > 1.0)):
            raise ValueError(
                f"Aggregated {objective_name} values must remain in [0, 1]."
            )
    X = np.vstack([r0_training.X_phys_all[r0_training.include_in_model], X_r1])
    Y = np.vstack([r0_training.Y_objectives[r0_training.include_in_model], Y_r1])
    if X.shape[0] != np.unique(X, axis=0).shape[0]:
        raise ValueError(
            "Combined condition-level training data contain duplicate recipes."
        )
    return X, Y


__all__ = [
    "D2D_DEBUG_WATERMARK",
    "D2D_BOUNDED_OBJECTIVE_COLUMNS",
    "D2D_INPUT_COLUMNS",
    "D2D_OBJECTIVE_COLUMNS",
    "D2D_OBJECTIVE_NAMES",
    "D2D_REFERENCE_POINT_UTILITY",
    "D2DDebugConfigError",
    "D2DTrainingData",
    "OffGridObservedException",
    "ReplicateAggregationResult",
    "ResolvedD2DDebugConfig",
    "aggregate_replicate_objectives",
    "build_d2d_objective_transform",
    "combine_r0_and_aggregated_r1",
    "expand_candidates_to_replicates",
    "load_d2d_debug_config",
    "load_d2d_workbook_frame",
    "prepare_d2d_training_data",
    "sha256_file",
]
