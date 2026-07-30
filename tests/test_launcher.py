"""The launcher's decisions, tested without a display.

The tkinter window is a thin shell over `inspect_campaign`, `gather_observations`
and `generate_next_round`; those are what can go wrong, so those are what is
tested here. Importing `mobo_kit.launcher` must not require tkinter, and one test
asserts that.
"""

from __future__ import annotations

import shutil

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from mobo_kit.campaign import (
    load_campaign_config,
    measurement_entry_columns,
    objective_names,
)
from mobo_kit.launcher import (
    CampaignStatus,
    LauncherError,
    gather_observations,
    generate_next_round,
    inspect_campaign,
)
from mobo_kit.workbook_io import (
    CandidateSheetError,
    candidate_workbook_path,
    read_candidate_results,
    sheet_name_for_round,
    write_candidate_sheet,
)

CONFIG_PATH = "configs/campaign_d2d_perovskite.yaml"
SOURCE = "local_inputs/Summary Table.xlsx"

pytestmark = pytest.mark.skipif(
    not __import__("pathlib").Path(SOURCE).exists(),
    reason="requires the ignored private campaign workbook",
)


@pytest.fixture(scope="module")
def config() -> dict:
    return load_campaign_config(CONFIG_PATH)


@pytest.fixture
def workbook(tmp_path):
    destination = tmp_path / "Summary Table.xlsx"
    shutil.copy2(SOURCE, destination)
    return destination


def _conditions(config: dict, n: int = 5) -> pd.DataFrame:
    names = [item["name"] for item in config["inputs"]]
    rows = [
        [float(item["start"]) + i * float(item["step"]) for item in config["inputs"]]
        for i in range(n)
    ]
    return pd.DataFrame(rows, columns=names)


def _fill_candidate_sheet(
    path, config, *, thickness=(700.0, 720.0), rows=None, coverage=1.0
) -> None:
    """Enter plausible measurements into every film of an R1 sheet."""
    book = load_workbook(path)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    values = {
        "Coverage": coverage,
        "Uniformity": 0.3,
        "Phase purity": 0.95,
        "PL - Implied Voc (Max)": 0.05,
        "Photoconductance (Max)": 5e-07,
        "T1": thickness[0],
        "T2": thickness[1],
    }
    target_rows = rows or range(2, sheet.max_row + 1)
    for row in target_rows:
        for column, value in values.items():
            sheet.cell(row=row, column=headers.index(column) + 1).value = value
    book.save(path)


# --------------------------------------------------------------------------- #
# status
# --------------------------------------------------------------------------- #


def test_a_fresh_workbook_is_ready_for_r1(workbook, config) -> None:
    status = inspect_campaign(workbook, config)
    assert status.next_round == "R1"
    assert status.can_generate
    assert status.observed_conditions == 15
    assert "Ready to propose R1" in status.headline


def test_the_detail_text_surfaces_the_read_findings(workbook, config) -> None:
    """The experimentalist should see that samples 8, 12 and 15 hold thickness
    readings that disagree, without going looking for it."""
    detail = inspect_campaign(workbook, config).detail()
    assert "Worth a look" in detail
    assert "1600" in detail and "709" in detail
    assert "For the record" in detail  # the excluded T anom readings


def test_a_missing_workbook_is_a_plain_sentence(tmp_path, config) -> None:
    with pytest.raises(LauncherError, match="does not exist"):
        inspect_campaign(tmp_path / "nope.xlsx", config)


def test_an_unmeasured_r1_sheet_blocks_the_next_round(workbook, config) -> None:
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    status = inspect_campaign(workbook, config)
    assert not status.can_generate
    assert "no results have been entered" in status.reason
    assert "Coverage" in status.reason  # says what to fill in


# --------------------------------------------------------------------------- #
# observations
# --------------------------------------------------------------------------- #


def test_r1_trains_on_sheet1_alone(workbook, config) -> None:
    X, Y, Yvar, provenance = gather_observations(workbook, config, for_round="R1")
    assert X.shape == (15, 10)
    assert Y.shape == (15, 3)
    assert provenance == ["Sheet1: 15 conditions"]
    # no replicates exist yet, so the noise is still fitted rather than measured
    assert Yvar is None


def test_r2_trains_on_sheet1_plus_the_aggregated_r1_conditions(
    workbook, config
) -> None:
    """Three films are one design point, so R2 sees 15 + 5, not 15 + 15."""
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    _fill_candidate_sheet(out, config)
    X, Y, Yvar, provenance = gather_observations(workbook, config, for_round="R2")
    assert X.shape == (20, 10)
    assert Y.shape == (20, 3)
    assert "5 conditions from 15 films" in provenance[1]
    # the live config still fits the noise; measured variance is one key away
    assert Yvar is None


def test_measured_replicate_variance_switches_on_from_config(workbook, config) -> None:
    """The promise of wiring this before the data exists: when the triplicates
    land, enabling it is a config edit, not a code change."""
    import copy

    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    _fill_candidate_sheet(out, config, thickness=(700.0, 760.0))
    # the films of a condition must actually differ, or there is no variance to pool
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    for row in range(2, sheet.max_row + 1):
        offset = row % 3
        sheet.cell(row=row, column=headers.index("T1") + 1).value = 700.0 + 40.0 * offset
        sheet.cell(row=row, column=headers.index("Coverage") + 1).value = 0.9 + 0.02 * offset
        # every objective needs film-to-film variation, or its pooled variance is
        # zero -- which the pooling refuses, because identical replicates are a
        # transcription rather than a measurement
        sheet.cell(row=row, column=headers.index("Photoconductance (Max)") + 1).value = (
            5e-07 * (1.0 + 0.1 * offset)
        )
    book.save(out)

    enabled = copy.deepcopy(dict(config))
    enabled["model"] = dict(enabled["model"])
    enabled["model"]["observation_noise"] = "replicate_pooled"

    X, Y, Yvar, provenance = gather_observations(workbook, enabled, for_round="R2")
    assert Yvar is not None
    assert Yvar.shape == Y.shape
    assert np.all(Yvar > 0)
    # the 15 R0 rows carry the full between-film variance; the R1 conditions,
    # being means of three films, carry a third of it
    assert Yvar[0, 2] == pytest.approx(3.0 * Yvar[15, 2])
    assert any("pooled between-film variance" in item for item in provenance)


def test_gathering_refuses_a_half_measured_film(workbook, config) -> None:
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    _fill_candidate_sheet(out, config)
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    # blank every thickness reading of one whole condition
    for row in (2, 3, 4):
        for column in ("T1", "T2"):
            sheet.cell(row=row, column=headers.index(column) + 1).value = None
    book.save(out)

    with pytest.raises(LauncherError, match="cannot be turned into objective values"):
        gather_observations(workbook, config, for_round="R2")


# --------------------------------------------------------------------------- #
# replicate aggregation
# --------------------------------------------------------------------------- #


def test_replicates_aggregate_to_one_observation_per_condition(
    workbook, config
) -> None:
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    _fill_candidate_sheet(out, config)
    results = read_candidate_results(workbook, config, "R1")
    assert results.n_conditions == 5
    assert len(results.replicates) == 15
    assert list(results.model_values.columns) == list(objective_names(config))
    assert (results.films_used["thickness"] == 3).all()


def test_thickness_aggregates_as_a_geometric_mean(workbook, config) -> None:
    """`response: log` means the GP trains on log(T), so three films are averaged
    in that space. With identical films the two means agree, which is why the
    check uses films that differ."""
    out = write_candidate_sheet(workbook, config, _conditions(config, 1), round_name="R1")
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    values = {
        "Coverage": 1.0,
        "Uniformity": 0.3,
        "Phase purity": 0.95,
        "PL - Implied Voc (Max)": 0.05,
        "Photoconductance (Max)": 5e-07,
    }
    per_film = (400.0, 700.0, 1000.0)
    for offset, thickness in enumerate(per_film):
        row = 2 + offset
        for column, value in values.items():
            sheet.cell(row=row, column=headers.index(column) + 1).value = value
        sheet.cell(row=row, column=headers.index("T1") + 1).value = thickness
    book.save(out)

    results = read_candidate_results(workbook, config, "R1")
    observed = float(results.model_values["thickness"].iloc[0])
    assert observed == pytest.approx(float(np.exp(np.mean(np.log(per_film)))))
    assert observed == pytest.approx(654.2, abs=0.1)  # (400*700*1000) ** (1/3)
    # and not the arithmetic mean, which is 700
    assert abs(observed - 700.0) > 40.0


def test_the_spread_is_kept_in_the_aggregation_space(workbook, config) -> None:
    """What Phase 4 needs: thickness spread already in log space, matching the
    config's decision to pool train_Yvar there."""
    out = write_candidate_sheet(workbook, config, _conditions(config, 1), round_name="R1")
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    for offset, thickness in enumerate((400.0, 700.0, 1000.0)):
        row = 2 + offset
        for column, value in {
            "Coverage": 1.0,
            "Uniformity": 0.3,
            "Phase purity": 0.95,
            "PL - Implied Voc (Max)": 0.05,
            "Photoconductance (Max)": 5e-07,
        }.items():
            sheet.cell(row=row, column=headers.index(column) + 1).value = value
        sheet.cell(row=row, column=headers.index("T1") + 1).value = thickness
    book.save(out)

    results = read_candidate_results(workbook, config, "R1")
    expected = float(np.std(np.log([400.0, 700.0, 1000.0]), ddof=1))
    assert float(results.replicate_spread["thickness"].iloc[0]) == pytest.approx(expected)
    # uniformity is identical across the three films, so its spread is zero
    assert float(results.replicate_spread["uniformity"].iloc[0]) == pytest.approx(0.0)


def test_films_of_one_condition_must_share_a_recipe(workbook, config) -> None:
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    _fill_candidate_sheet(out, config)
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row=3, column=headers.index("speed_1") + 1).value = 4242.0
    book.save(out)

    with pytest.raises(CandidateSheetError, match="do not share the same speed_1"):
        read_candidate_results(workbook, config, "R1")


def test_reading_a_sheet_that_was_never_written_says_so(workbook, config) -> None:
    with pytest.raises(CandidateSheetError, match="does not exist"):
        read_candidate_results(workbook, config, "R1")


# --------------------------------------------------------------------------- #
# generating
# --------------------------------------------------------------------------- #


@pytest.mark.slow
def test_generating_r1_writes_a_sheet_and_leaves_the_source_alone(
    workbook, config
) -> None:
    import hashlib

    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    messages: list[str] = []
    generated = generate_next_round(workbook, config, progress=messages.append)

    assert generated.round_name == "R1"
    assert generated.sheet_path == candidate_workbook_path(workbook, "R1")
    assert generated.sheet_path.exists()
    assert generated.result.n_conditions == 5
    assert generated.n_films == 15
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before
    assert messages and "Done." in messages

    summary = generated.summary()
    assert "Nothing here is approved" in summary
    assert "Sheet1: 15 conditions" in summary
    # the sheet is immediately readable by the reader that will consume it
    required, _ = measurement_entry_columns(config)
    headers = [
        cell.value
        for cell in load_workbook(generated.sheet_path)[sheet_name_for_round("R1")][1]
    ]
    for column in required:
        assert column in headers


def test_generating_refuses_when_no_round_is_due(workbook, config) -> None:
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    with pytest.raises(LauncherError, match="no results have been entered"):
        generate_next_round(workbook, config)


def test_generating_never_overwrites_an_existing_sheet(workbook, config, monkeypatch) -> None:
    """The sheet may already hold measurements. Refusing is the only safe move,
    and it must happen before the ten seconds of model fitting, not after."""
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")

    def fail(*args, **kwargs):  # pragma: no cover - must never be reached
        raise AssertionError("the round was fitted despite an existing sheet")

    monkeypatch.setattr("mobo_kit.launcher.run_r1_ucb", fail)
    monkeypatch.setattr(
        "mobo_kit.launcher.inspect_campaign",
        lambda *a, **k: CampaignStatus(
            workbook=workbook,
            next_round="R1",
            reason="pretend R1 is due",
            scored_rows=0,
            total_rows=0,
            observed_conditions=15,
        ),
    )
    with pytest.raises(LauncherError, match="already exists"):
        generate_next_round(workbook, config)


# --------------------------------------------------------------------------- #
# the shell
# --------------------------------------------------------------------------- #


@pytest.fixture
def isolated_settings(monkeypatch):
    """No remembered workbook, and no writing to the user's home.

    Both matter. The launcher schedules a `check()` 200 ms after construction when
    it remembers a workbook, so a path left behind by an earlier test raced the
    explicit `check()` these tests perform and overwrote the pane with a different
    result -- an order-dependent failure that only appeared in a full-suite run.
    And a test suite has no business writing to ~/.mobo_kit either way.
    """
    monkeypatch.setattr("mobo_kit.launcher.remembered_workbook", lambda: None)
    monkeypatch.setattr("mobo_kit.launcher.remember_workbook", lambda path: None)


def _status_for(path) -> CampaignStatus:
    from pathlib import Path

    return CampaignStatus(
        workbook=Path(path).resolve(),
        next_round="R1",
        reason="pretend R1 is due",
        scored_rows=0,
        total_rows=0,
        observed_conditions=15,
    )


def _tk_available() -> bool:
    try:
        import tkinter

        root = tkinter.Tk()
    except Exception:
        return False
    root.destroy()
    return True


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_the_window_reports_status_through_its_worker_thread(
    workbook, config, isolated_settings
) -> None:
    """The UI does its work off the main thread and posts results through a queue.
    Nothing else covers that plumbing, and a deadlock there would look like a
    window that simply never responds."""
    import time

    from mobo_kit.launcher import LauncherWindow

    window = LauncherWindow(CONFIG_PATH)
    try:
        window.path_var.set(str(workbook))
        window.check()
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            window.root.update()
            if not window._busy and window._status is not None:
                break
            time.sleep(0.02)

        assert window._status is not None, "the window never reported a status"
        assert window._status.next_round == "R1"
        assert window.headline.cget("text") == "Ready to propose R1."
        assert window.generate_button.cget("text") == "Propose R1"
        assert str(window.generate_button.cget("state")) == "normal"
        body = window.text.get("1.0", "end")
        assert "15 conditions on Sheet1" in body
    finally:
        window.root.destroy()


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_the_window_shows_a_readable_error_rather_than_a_traceback(
    config, isolated_settings
) -> None:
    import time

    from mobo_kit.launcher import LauncherWindow

    window = LauncherWindow(CONFIG_PATH)
    try:
        window.path_var.set("nowhere/at/all.xlsx")
        window.check()
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            window.root.update()
            if not window._busy:
                break
            time.sleep(0.02)
        body = window.text.get("1.0", "end")
        assert "does not exist" in body
        assert "Traceback" not in body
        assert window.headline.cget("text") == "Cannot continue."
    finally:
        window.root.destroy()


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_a_result_for_a_workbook_the_user_left_is_discarded(
    workbook, config, isolated_settings
) -> None:
    """The race the test fixture hid, now closed at the source.

    Work runs off the main thread, so a check dispatched against one workbook can
    return after the user has selected another. Painting "Ready to propose R1" over
    a different workbook is worse than painting nothing.

    Driven through the queue rather than by racing two real threads. The first
    version of this test did race them, passed alone, and failed intermittently in
    a full-suite run -- a flaky test of a race-condition fix is worse than no test,
    because it teaches people to re-run until green.
    """
    from mobo_kit.launcher import LauncherWindow

    window = LauncherWindow(CONFIG_PATH)
    try:
        window.path_var.set(str(workbook))
        window._start("pretending to read")
        window._request_id = 1
        # the user navigates away before the reply lands
        window.path_var.set(str(workbook.parent / "somewhere else.xlsx"))
        window._queue.put((1, "status", _status_for(workbook)))
        window.drain_once()

        assert window._status is None, "the stale status must not be adopted"
        assert not window._busy, "a discarded reply must still clear the busy state"
        assert "Ready to propose" not in window.headline.cget("text")
        # buttons usable again rather than stuck disabled
        assert str(window.check_button.cget("state")) == "normal"
        assert str(window.generate_button.cget("state")) == "disabled"
    finally:
        window.root.destroy()


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_a_result_for_the_current_workbook_is_adopted(
    workbook, config, isolated_settings
) -> None:
    """The other half of the rule: it must not discard everything."""
    from mobo_kit.launcher import LauncherWindow

    window = LauncherWindow(CONFIG_PATH)
    try:
        window.path_var.set(str(workbook))
        window._start("pretending to read")
        window._request_id = 1
        window._queue.put((1, "status", _status_for(workbook)))
        window.drain_once()

        assert window._status is not None
        assert window.headline.cget("text") == "Ready to propose R1."
        assert str(window.generate_button.cget("state")) == "normal"
    finally:
        window.root.destroy()


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_a_superseded_reply_does_not_overwrite_a_newer_request(
    workbook, config, isolated_settings
) -> None:
    """Two presses: the earlier press's answer must not land after the later one."""
    from mobo_kit.launcher import LauncherWindow

    window = LauncherWindow(CONFIG_PATH)
    try:
        window.path_var.set(str(workbook))
        window._start("pretending to read")
        window._request_id = 2  # a second press is already in flight
        window._queue.put((1, "status", _status_for(workbook)))
        window.drain_once()
        assert window._status is None, "request 1's reply landed after request 2"
        assert not window._busy

        window._start("still pretending")
        window._queue.put((2, "status", _status_for(workbook)))
        window.drain_once()
        assert window._status is not None, "request 2's own reply must land"
    finally:
        window.root.destroy()


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_the_startup_auto_check_is_cancelled_when_the_user_acts(
    workbook, config, monkeypatch
) -> None:
    """The auto-check fires 200 ms after construction against the remembered
    workbook. If the user has already pressed something, that answer is about the
    wrong file."""
    from mobo_kit import launcher as launcher_module
    from mobo_kit.launcher import LauncherWindow

    monkeypatch.setattr(launcher_module, "remembered_workbook", lambda: workbook)
    monkeypatch.setattr(launcher_module, "remember_workbook", lambda path: None)

    window = LauncherWindow(CONFIG_PATH)
    try:
        assert window._auto_check_id is not None, "a remembered workbook should schedule one"
        window._cancel_auto_check()
        assert window._auto_check_id is None
        # cancelling twice is harmless
        window._cancel_auto_check()
    finally:
        window.root.destroy()


def test_the_logic_imports_without_tkinter(monkeypatch) -> None:
    """A headless machine must still be able to use the functions. tkinter is
    imported inside the window class for exactly this reason."""
    import importlib
    import sys

    monkeypatch.setitem(sys.modules, "tkinter", None)
    module = importlib.reload(importlib.import_module("mobo_kit.launcher"))
    assert callable(module.generate_next_round)
