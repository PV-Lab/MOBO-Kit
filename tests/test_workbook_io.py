from __future__ import annotations

import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from mobo_kit.campaign import (
    load_campaign_config,
    measurement_entry_columns,
    model_source_columns,
    objective_names,
)
from mobo_kit.workbook_io import (
    candidate_workbook_path,
    CandidateSheetError,
    detect_round,
    read_campaign_workbook,
    sheet_name_for_round,
    write_candidate_sheet,
)

CONFIG_PATH = "configs/campaign_d2d_perovskite.yaml"
SOURCE = "local_inputs/Summary Table.xlsx"

pytestmark = pytest.mark.skipif(
    not __import__("pathlib").Path(SOURCE).exists(),
    reason="requires the ignored private campaign workbook",
)


@pytest.fixture(scope="module")
def config() -> dict:
    return load_campaign_config(CONFIG_PATH)


@pytest.fixture
def workbook(tmp_path):
    destination = tmp_path / "Summary Table.xlsx"
    shutil.copy2(SOURCE, destination)
    return destination


def _conditions(config: dict, n: int = 5) -> pd.DataFrame:
    names = [item["name"] for item in config["inputs"]]
    rows = [
        [float(item["start"]) + i * float(item["step"]) for item in config["inputs"]]
        for i in range(n)
    ]
    return pd.DataFrame(rows, columns=names)


def test_reads_inputs_and_computes_one_value_per_objective(workbook, config) -> None:
    contents = read_campaign_workbook(workbook, config)
    assert contents.n_rows == 15
    assert list(contents.inputs.columns) == [i["name"] for i in config["inputs"]]
    assert list(contents.model_values.columns) == list(objective_names(config))
    # thickness must arrive in nanometres, not as its score
    assert contents.model_values["thickness"].max() > 100.0
    assert contents.model_values.notna().all().all()


def test_computed_values_agree_with_the_workbook_within_tolerance(
    workbook, config
) -> None:
    """The recomputation is not a different quantity: on the R0 rows it reproduces
    the stored cells to floating-point noise, and thickness only to half a
    nanometre because `Thickness (avg)` is ROUND(mean(T1..T4))."""
    contents = read_campaign_workbook(workbook, config)
    stored = contents.workbook_values
    assert (
        (contents.model_values["uniformity"] - stored["Uniformity score"]).abs().max()
        < 1e-12
    )
    assert (
        (contents.model_values["optoelectronic"] - stored["Optoelectronic score"])
        .abs()
        .max()
        < 1e-12
    )
    thickness_gap = (
        (contents.model_values["thickness"] - stored["Thickness (avg)"]).abs().max()
    )
    assert thickness_gap <= 0.5
    # and it is genuinely unrounded, or the gap would be zero
    assert thickness_gap > 0.0


def test_the_r0_rows_produce_no_errors_and_flag_the_disagreeing_films(
    workbook, config
) -> None:
    contents = read_campaign_workbook(workbook, config)
    assert contents.errors == ()
    disagreeing = {
        finding.sample_id
        for finding in contents.warnings
        if finding.code == "readings_disagree"
    }
    # samples 8, 12 and 15 hold thickness readings that split into two clusters
    assert disagreeing == {8, 12, 15}
    excluded = {
        finding.sample_id
        for finding in contents.findings
        if finding.code == "reading_excluded"
    }
    assert excluded == {4, 14}


def test_thickness_records_how_many_readings_each_row_used(workbook, config) -> None:
    """Two readings and four readings do not carry the same weight; Phase 4 needs
    the count to turn a spread into an observation variance."""
    contents = read_campaign_workbook(workbook, config)
    counts = contents.inputs_used["thickness"]
    assert counts.min() == 2 and counts.max() == 4
    assert counts.value_counts().to_dict() == {2: 9, 3: 3, 4: 3}


def test_stops_at_the_first_blank_sample_number(workbook, config) -> None:
    """Rows below the data block are notes, not observations."""
    contents = read_campaign_workbook(workbook, config)
    assert contents.sample_ids == tuple(range(1, 16))


def test_candidate_sheet_asks_for_raw_measurements_not_derived_scores(
    workbook, config
) -> None:
    """The objectives are computed now, so the sheet must collect what they are
    computed from. Offering a `Thickness (avg)` cell would invite someone to fill
    in a value that nothing reads."""
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    sheet = load_workbook(candidate_workbook_path(workbook, "R1"))[
        sheet_name_for_round("R1")
    ]
    headers = [c.value for c in sheet[1]]
    required, optional = measurement_entry_columns(config)
    for column in (*required, *optional):
        assert column in headers
    assert "T1" in headers and "T4" in headers and "T anom" in headers
    for derived in model_source_columns(config):
        assert derived not in headers


def test_writing_leaves_the_source_byte_identical(workbook, config) -> None:
    """openpyxl discards cached formula values on save, and Uniformity score is a
    formula column. Writing beside the workbook makes that impossible rather than
    merely unlikely."""
    import hashlib

    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    assert out != workbook
    assert out.exists()
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before


def test_formula_columns_survive_because_the_source_is_not_rewritten(
    workbook, config
) -> None:
    """The regression this design exists to prevent."""
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    contents = read_campaign_workbook(workbook, config)
    # Uniformity score is =L*N*O; a rewritten workbook reads it as NaN, which
    # would now surface as a cross_check_empty warning rather than as bad training
    # data -- but the invariant worth holding is still that it survives
    assert contents.workbook_values["Uniformity score"].notna().all()
    assert contents.workbook_values["Uniformity score"].max() > 0.0
    assert [f for f in contents.warnings if f.code == "cross_check_empty"] == []


def test_three_replicate_rows_per_condition(workbook, config) -> None:
    write_candidate_sheet(
        workbook, config, _conditions(config, 5), round_name="R1", replicates=3
    )
    sheet = load_workbook(candidate_workbook_path(workbook, "R1"))[
        sheet_name_for_round("R1")
    ]
    rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(rows) == 15
    assert len({r[1] for r in rows}) == 5


def test_refuses_to_overwrite_an_existing_sheet(workbook, config) -> None:
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    with pytest.raises(CandidateSheetError, match="already exists"):
        write_candidate_sheet(
            workbook, config, _conditions(config), round_name="R1", make_backup=False
        )


def test_round_detection_walks_r1_then_r2(workbook, config) -> None:
    assert detect_round(workbook, config).next_round == "R1"
    write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    state = detect_round(workbook, config)
    assert state.next_round is None
    assert "no results have been entered" in state.reason


def test_partially_scored_sheet_is_refused_with_a_readable_message(
    workbook, config
) -> None:
    """Fail closed: guessing at a half-filled sheet is how a round gets built on
    data the experimentalist had not finished entering."""
    out = write_candidate_sheet(workbook, config, _conditions(config), round_name="R1")
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [c.value for c in sheet[1]]
    required, _ = measurement_entry_columns(config)
    for column in (*required, "T1"):
        sheet.cell(row=2, column=headers.index(column) + 1).value = 1.0
    book.save(out)

    state = detect_round(workbook, config)
    assert state.next_round is None
    assert state.scored_rows == 1 and state.total_rows == 15
    assert "partly filled in" in state.reason
    assert "1 of 15" in state.reason


def test_a_row_with_two_thickness_readings_counts_as_complete(workbook, config) -> None:
    """Nine of the fifteen R0 rows have only T1 and T2. Demanding all four would
    hold a finished round hostage to measurements nobody intended to take."""
    out = write_candidate_sheet(workbook, config, _conditions(config, 1), round_name="R1")
    book = load_workbook(out)
    sheet = book[sheet_name_for_round("R1")]
    headers = [c.value for c in sheet[1]]
    required, _ = measurement_entry_columns(config)
    for row in (2, 3, 4):
        for column in (*required, "T1", "T2"):
            sheet.cell(row=row, column=headers.index(column) + 1).value = 1.0
    book.save(out)

    # accepting R1 as complete is what advances the campaign to R2; a stricter
    # rule would leave it stuck reporting "partly filled in" forever
    state = detect_round(workbook, config)
    assert state.next_round == "R2"


def test_missing_source_sheet_is_a_plain_sentence(tmp_path, config) -> None:
    from openpyxl import Workbook

    path = tmp_path / "wrong.xlsx"
    book = Workbook()
    book.active.title = "Renamed"
    book.save(path)
    # The message must name the sheet the CONFIG asked for, the key that decides
    # it, and what the workbook actually has. "Rename it back" was the old advice
    # and stopped being right once the sheet became configuration: the likelier
    # cause is now a workbook belonging to a different campaign.
    with pytest.raises(CandidateSheetError, match="campaign.source_sheet"):
        read_campaign_workbook(path, config)
