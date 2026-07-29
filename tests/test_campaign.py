from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
import torch

from mobo_kit.campaign import (
    FIXED_SCALING_MODES,
    assert_scaling_is_campaign_fixed,
    write_worklist_csv,
    BatchValidityError,
    CampaignConfigError,
    build_objective_transform,
    expand_replicates,
    load_campaign_config,
    model_source_columns,
    run_r0_lhs,
    validate_batch,
)
from mobo_kit.design import build_design_from_config

CONFIG_PATH = "configs/FA0.9CS0.1PbI3_260407_Config.yaml"


@pytest.fixture(scope="module")
def config() -> dict:
    return load_campaign_config(CONFIG_PATH)


def test_campaign_config_is_runnable(config: dict) -> None:
    """The canonical config must no longer be a baseline-only stub."""
    assert config["campaign"]["status"] == "active"
    assert len(config["inputs"]) == 10
    assert config["objectives"]["contract_version"]
    assert len(config["objectives"]["specs"]) == 3
    assert config["constraints"] == []


def test_thickness_trains_on_nanometres_not_the_score(config: dict) -> None:
    """The whole point of the objective rework: source column != utility."""
    sources = model_source_columns(config)
    assert sources == (
        "Uniformity score",
        "Optoelectronic score",
        "Thickness (avg)",
    )
    transform = build_objective_transform(config)
    thickness = transform.specs[2]
    assert thickness.transform == "gaussian_target"
    assert thickness.target == pytest.approx(650.0)
    # workbook writes exp(-((T-650)/250)^2); this convention carries the 1/2
    assert thickness.sigma == pytest.approx(250.0 / np.sqrt(2.0))


def test_transform_reproduces_the_workbook_thickness_score(config: dict) -> None:
    """Thickness is a log-link objective: the GP emits log(nm), so the transform
    exponentiates before applying the 650 nm Gaussian. Feeding it raw nm would
    silently score exp(687) instead of 687."""
    transform = build_objective_transform(config)
    assert transform.specs[2].model_link == "log"
    nm = np.array([687.0, 1303.0])
    model_output = torch.tensor(
        [[0.0, -8.0, np.log(nm[0])], [0.0, -8.0, np.log(nm[1])]], dtype=torch.double
    )
    got = transform(model_output)[:, 2].numpy()
    expected = np.exp(-(((nm - 650.0) / 250.0) ** 2))
    np.testing.assert_allclose(got, expected, atol=1e-12)


def test_reference_point_is_declared_in_utility_space(config: dict) -> None:
    """A raw-scale reference silently weighted optoelectronic 4x; utility space
    puts every axis on a comparable scale."""
    point = config["reference_point_utility"]
    assert len(point) == 3
    assert all(abs(float(v)) < 1.0 for v in point)


def test_objectives_without_a_source_column_are_rejected() -> None:
    bad = {
        "objectives": {
            "contract_version": "x",
            "specs": [{"name": "a", "goal": "maximize", "transform": "identity"}],
        }
    }
    with pytest.raises(CampaignConfigError, match="model_source_column"):
        build_objective_transform(bad)


def test_empty_objective_list_cannot_propose() -> None:
    with pytest.raises(CampaignConfigError, match="specs"):
        build_objective_transform(
            {"objectives": {"contract_version": "x", "specs": []}}
        )


# --------------------------------------------------------------------------- #
# validity gate
# --------------------------------------------------------------------------- #


def _valid_batch(config: dict) -> pd.DataFrame:
    design = build_design_from_config(dict(config))
    rows = [[float(design.var_array[j][i * 2]) for j in range(10)] for i in range(3)]
    return pd.DataFrame(rows, columns=list(design.names))


def test_validate_batch_accepts_a_clean_batch(config: dict) -> None:
    design = build_design_from_config(dict(config))
    report = validate_batch(_valid_batch(config), design, expected_count=3)
    assert report["unique"] and report["on_grid"] and report["in_bounds"]
    assert report["actual_count"] == 3


@pytest.mark.parametrize(
    "mutate, match",
    [
        (lambda d: d.iloc[:2], "Expected exactly 3"),
        (lambda d: pd.concat([d.iloc[:2], d.iloc[[0]]]), "unique"),
        (lambda d: d.assign(anti_time=12.0), "grid"),
        (lambda d: d.assign(speed_1=99999.0), "grid"),
        (lambda d: d.assign(speed_1=float("nan")), "non-finite"),
    ],
)
def test_validate_batch_refuses_real_defects(config: dict, mutate, match) -> None:
    design = build_design_from_config(dict(config))
    with pytest.raises(BatchValidityError, match=match):
        validate_batch(mutate(_valid_batch(config)), design, expected_count=3)


def test_validate_batch_enforces_minimum_spacing(config: dict) -> None:
    design = build_design_from_config(dict(config))
    batch = _valid_batch(config)
    with pytest.raises(BatchValidityError, match="pairwise distance"):
        validate_batch(batch, design, expected_count=3, min_pairwise_distance=10.0)


def test_validity_report_carries_no_approval_flags(config: dict) -> None:
    """The debug/production tiers are gone. Approval is a human decision recorded
    outside the code, not something a validity check can compute."""
    design = build_design_from_config(dict(config))
    report = validate_batch(_valid_batch(config), design, expected_count=3)
    for banned in (
        "debug_only",
        "approved_for_experiment",
        "approved_for_production",
        "experimental_approval_false",
        "production_approval_false",
    ):
        assert banned not in report


# --------------------------------------------------------------------------- #
# replicates
# --------------------------------------------------------------------------- #


def test_expand_replicates_groups_three_films_per_condition(config: dict) -> None:
    batch = _valid_batch(config)
    films = expand_replicates(batch, replicates=3, round_name="R1")
    assert len(films) == 9
    assert films["replicate_group"].nunique() == 3
    assert set(films["replicate_index"]) == {1, 2, 3}
    assert set(films["round"]) == {"R1"}
    for _, group in films.groupby("replicate_group"):
        inputs = group[list(batch.columns)].drop_duplicates()
        assert len(inputs) == 1, "replicates must share identical inputs"


def test_expand_replicates_rejects_zero(config: dict) -> None:
    with pytest.raises(ValueError, match="at least 1"):
        expand_replicates(_valid_batch(config), replicates=0, round_name="R1")


# --------------------------------------------------------------------------- #
# R0
# --------------------------------------------------------------------------- #


def test_run_r0_lhs_produces_a_valid_on_grid_worklist(config: dict) -> None:
    result = run_r0_lhs(config, n=8, seed=7)
    assert result.round_name == "R0"
    assert result.n_conditions == 8
    assert list(result.conditions.columns) == [i["name"] for i in config["inputs"]]
    assert result.diagnostics["validity"]["on_grid"]
    assert (
        len(result.replicates) == 8 * config["rounds"]["r1"]["replicates_per_condition"]
    )


def test_run_r0_lhs_is_deterministic_for_a_seed(config: dict) -> None:
    a = run_r0_lhs(config, n=6, seed=11).conditions
    b = run_r0_lhs(config, n=6, seed=11).conditions
    pd.testing.assert_frame_equal(a, b)


# --------------------------------------------------------------------------- #
# encoding
# --------------------------------------------------------------------------- #

# Non-ASCII that breaks under a locale default codec such as GBK or cp1252.
NON_ASCII = "sigma \u2014 \u00b5m \u00b1 5% \u2013 caf\u00e9 \u4e2d\u6587"


def test_config_round_trips_non_ascii(tmp_path) -> None:
    """Loaders must force UTF-8.

    The default codec is locale dependent -- on a Chinese-locale Windows install
    it is GBK -- so a bare open() fails on the first non-ASCII byte, and it fails
    on a different machine from the one the file was written on.
    """
    import yaml

    path = tmp_path / "cfg.yaml"
    payload = {
        "campaign": {"name": NON_ASCII},
        "objectives": {
            "contract_version": NON_ASCII,
            "specs": [
                {
                    "name": "a",
                    "goal": "maximize",
                    "transform": "identity",
                    "model_source_column": NON_ASCII,
                }
            ],
        },
    }
    path.write_text(yaml.safe_dump(payload, allow_unicode=True), encoding="utf-8")

    loaded = load_campaign_config(path)
    assert loaded["campaign"]["name"] == NON_ASCII
    assert model_source_columns(loaded) == (NON_ASCII,)
    assert build_objective_transform(loaded).version == NON_ASCII


def test_workbook_reader_round_trips_non_ascii(tmp_path) -> None:
    """A non-ASCII column header or cell must survive the workbook boundary."""
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "wb.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["Sample number", NON_ASCII])
    sheet.append([1, NON_ASCII])
    book.save(path)

    reread = load_workbook(path, data_only=True)["Sheet1"]
    rows = list(reread.iter_rows(values_only=True))
    assert rows[0][1] == NON_ASCII
    assert rows[1][1] == NON_ASCII


def test_csv_round_trips_non_ascii(tmp_path) -> None:
    """pandas defaults to UTF-8 for both directions; pin it with a test so a
    future explicit encoding= cannot silently regress it."""
    path = tmp_path / "out.csv"
    frame = pd.DataFrame({"label": [NON_ASCII], "value": [1.0]})
    frame.to_csv(path, index=False)
    pd.testing.assert_frame_equal(pd.read_csv(path), frame)


# --------------------------------------------------------------------------- #
# campaign-fixed scaling
# --------------------------------------------------------------------------- #


def test_campaign_declares_fixed_scaling(config: dict) -> None:
    assert config["objectives"]["scaling_mode"] in FIXED_SCALING_MODES
    assert_scaling_is_campaign_fixed(config)


def test_data_derived_scaling_is_refused(config: dict) -> None:
    """If the scale tracks the data, hypervolume stops being comparable between
    rounds. The temptation arrives the moment R1 measurements land."""
    import copy

    bad = copy.deepcopy(dict(config))
    bad["objectives"]["scaling_mode"] = "observed_min_max"
    with pytest.raises(CampaignConfigError, match="incomparable"):
        assert_scaling_is_campaign_fixed(bad)
    with pytest.raises(CampaignConfigError, match="incomparable"):
        build_objective_transform(bad)


def test_affine_objective_without_declared_anchors_is_refused(config: dict) -> None:
    """Two layers refuse this: ObjectiveSpec at construction, and
    assert_scaling_is_campaign_fixed for anything that reaches it. Whichever
    fires first, an affine objective can never end up with implicit anchors."""
    import copy

    bad = copy.deepcopy(dict(config))
    del bad["objectives"]["specs"][0]["upper_anchor"]
    with pytest.raises(ValueError, match="upper_anchor"):
        build_objective_transform(bad)


def test_inverted_anchors_are_refused(config: dict) -> None:
    import copy

    bad = copy.deepcopy(dict(config))
    spec = bad["objectives"]["specs"][0]
    spec["lower_anchor"], spec["upper_anchor"] = 1.0, 0.0
    with pytest.raises(ValueError, match="anchor"):
        build_objective_transform(bad)


def test_worklist_csv_is_excel_safe(tmp_path) -> None:
    """Excel does not detect plain UTF-8 and falls back to the system ANSI
    codepage, mangling non-ASCII cells on the reader's machine rather than the
    writer's. utf-8-sig writes the BOM; other readers strip it transparently."""
    path = tmp_path / "worklist.csv"
    frame = pd.DataFrame({"label": [NON_ASCII], "speed_1": [1000.0]})
    write_worklist_csv(frame, path)

    assert path.read_bytes().startswith(b"\xef\xbb\xbf"), "missing UTF-8 BOM"
    pd.testing.assert_frame_equal(pd.read_csv(path), frame)
    pd.testing.assert_frame_equal(pd.read_csv(path, encoding="utf-8-sig"), frame)
