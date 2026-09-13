"""The batch review artifact.

Two kinds of test here, deliberately separated.

*Config parsing* runs against the live campaign YAML, because what the campaign
declares -- the low-speed probe, the anneal_temp note -- is part of what shipped
and should break if someone deletes it.

*Artifact building* runs against a small purpose-built config with three inputs.
It needs a real GP fit, and `model_validation`'s signal-collapse guard is strict
for good reason: it refuses a fit whose residual carries no signal. Manufacturing
ten well-conditioned observations in a 10-input space just to test a table is
fighting the wrong battle, and a test that trips a legitimate guard teaches the
next person to weaken the guard.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from mobo_kit.batch_review import (
    NOT_APPROVED,
    SD_MATERIALITY_RATIO,
    ProbeSpec,
    build_batch_review,
    classify_probe_objective,
    probe_specs_from_config,
    review_notes_from_config,
    write_review_sheet,
)
from mobo_kit.campaign import load_campaign_config
from mobo_kit.design import build_design_from_config
from mobo_kit.lhs import lhs_dataframe_optimized
from mobo_kit.scores import ScoreFinding, ScoreSeverity

LIVE_CONFIG_PATH = "configs/campaign_d2d_perovskite.yaml"


@pytest.fixture(scope="module")
def live_config() -> dict:
    return load_campaign_config(LIVE_CONFIG_PATH)


def _config() -> dict:
    """Three inputs, the same three objective shapes as the campaign.

    Keeps the log-link thickness objective and the identity-link linear mean on
    `anneal_temp`, because those are the two paths the review has to decode
    correctly.
    """
    return {
        "inputs": [
            {"name": "speed_1", "unit": "rpm", "start": 1000, "stop": 6000, "step": 500},
            {"name": "precur_conc", "unit": "M", "start": 1.0, "stop": 2.0, "step": 0.05},
            {"name": "anneal_temp", "unit": "C", "start": 100, "stop": 185, "step": 5},
        ],
        "objectives": {
            "contract_version": "review-test-v1",
            "scaling_mode": "fixed_affine",
            "specs": [
                {
                    "name": "uniformity",
                    "model_source_column": "U",
                    "transform": "affine",
                    "goal": "maximize",
                    "lower_anchor": 0.0,
                    "upper_anchor": 1.0,
                },
                {
                    "name": "optoelectronic",
                    "model_source_column": "O",
                    "transform": "affine",
                    "goal": "maximize",
                    "lower_anchor": -10.0,
                    "upper_anchor": -6.0,
                    "mean_function": {
                        "response": "identity",
                        "features": [{"column": "anneal_temp", "transform": "identity"}],
                    },
                },
                {
                    "name": "thickness",
                    "model_source_column": "T",
                    "transform": "gaussian_target",
                    "goal": "target",
                    "target": 650.0,
                    "sigma": 176.7766952966369,
                    "mean_function": {
                        "response": "log",
                        "features": [
                            {"column": "speed_1", "transform": "log"},
                            {"column": "precur_conc", "transform": "log"},
                        ],
                    },
                },
            ],
        },
        "reference_point_utility": [-0.01, -0.01, -0.01],
        "rounds": {
            "r1": {
                "method": "ucb_hvi",
                "batch_size": 3,
                "replicates_per_condition": 3,
                "beta": 4.0,
                "candidate_pool_size": 512,
                "posterior_samples": 64,
                "moment_method": "monte_carlo",
            }
        },
        "local_penalization": {"radius": 0.25, "min_batch_distance": 0.15},
        "model": {"variant": "dim_scaled_prior"},
        "reproducibility": {"seed": 7},
        "constraints": [],
        "review": {
            "probes": [
                {
                    "name": "low-speed corner",
                    "column": "speed_1",
                    "value": 1000,
                    "note": "The two observations here contradict each other.",
                }
            ],
            "notes": ["anneal_temp sits at a range edge by construction."],
        },
    }


@pytest.fixture(scope="module")
def config() -> dict:
    return _config()


@pytest.fixture(scope="module")
def design(config):
    return build_design_from_config(dict(config))


def _rows(config: dict, n: int, *, seed: int) -> np.ndarray:
    design = build_design_from_config(dict(config))
    frame = lhs_dataframe_optimized(design, n, seed=seed, snap_to_grids=True)
    return frame.to_numpy(dtype=float)


def _observations(config: dict, n: int = 12, *, seed: int = 1) -> tuple[np.ndarray, np.ndarray]:
    """Trends the mean functions can find, plus residual structure the GPs can.

    Residual structure matters: data following a mean function exactly leaves a
    pure-noise residual, which the collapse guard rejects -- correctly, since a GP
    with nothing to model has no business scoring candidates.
    """
    X = _rows(config, n, seed=seed)
    speed, concentration, temperature = X[:, 0], X[:, 1], X[:, 2]
    scaled_temperature = (temperature - 100.0) / 85.0
    scaled_concentration = (concentration - 1.0) / 1.0

    thickness = np.exp(
        9.6
        - 0.38 * np.log(speed)
        + 0.5 * np.log(concentration)
        + 0.25 * np.sin(3.0 * scaled_temperature)
    )
    optoelectronic = (
        -6.2 - 0.012 * temperature + 0.35 * np.cos(2.5 * scaled_concentration)
    )
    uniformity = np.clip(
        0.5 + 0.35 * np.sin(2.0 * scaled_concentration + 0.7 * scaled_temperature),
        0.02,
        0.98,
    )
    return X, np.column_stack([uniformity, optoelectronic, thickness])


@pytest.fixture(scope="module")
def review(config):
    X, Y = _observations(config)
    conditions = pd.DataFrame(
        _rows(config, 3, seed=99), columns=[i["name"] for i in config["inputs"]]
    )
    return build_batch_review(
        config,
        X,
        Y,
        conditions,
        round_name="R1",
        findings=(
            ScoreFinding(
                severity=ScoreSeverity.WARNING,
                code="readings_disagree",
                objective="thickness",
                row_position=11,
                sample_id=12,
                message="2 readings span 891: ['1600', '709'].",
            ),
            ScoreFinding(
                severity=ScoreSeverity.NOTE,
                code="reading_excluded",
                objective="thickness",
                row_position=3,
                sample_id=4,
                message="'T anom' holds 1618, judged anomalous.",
            ),
        ),
        context={"Round": "R1", "Seed": 7},
    )


# --------------------------------------------------------------------------- #
# the judgment, on its own
# --------------------------------------------------------------------------- #


def test_worse_with_the_same_uncertainty_reads_as_known_and_bad() -> None:
    """The numbers from the campaign's own R0 fit: thickness utility 0.223 against
    0.786, sd ratio 1.02. A bare `>` on sd calls that 'more uncertain' and prints
    the benign verdict, which is how this nearly shipped wrong."""
    assert classify_probe_objective(0.223, 0.786, 1.02) == "known_and_bad"


def test_worse_but_materially_more_uncertain_reads_as_a_tradeoff() -> None:
    assert classify_probe_objective(0.223, 0.786, 2.0) == "uncertain_tradeoff"


def test_not_worse_means_the_model_has_no_objection() -> None:
    assert classify_probe_objective(0.9, 0.8, 1.0) is None
    assert classify_probe_objective(0.8, 0.8, 1.0) is None


def test_the_threshold_is_a_ratio_not_an_inequality() -> None:
    assert classify_probe_objective(0.1, 0.5, SD_MATERIALITY_RATIO - 0.01) == "known_and_bad"
    assert (
        classify_probe_objective(0.1, 0.5, SD_MATERIALITY_RATIO + 0.01)
        == "uncertain_tradeoff"
    )


def test_a_zero_batch_sd_does_not_divide_by_zero() -> None:
    assert classify_probe_objective(0.1, 0.5, float("inf")) == "uncertain_tradeoff"


# --------------------------------------------------------------------------- #
# what the live campaign declares
# --------------------------------------------------------------------------- #


def test_the_live_config_declares_the_low_speed_probe(live_config) -> None:
    probes = probe_specs_from_config(live_config)
    assert [p.column for p in probes] == ["speed_1"]
    assert probes[0].value == 1000.0
    assert "contradict each other" in probes[0].note


def test_the_live_config_declares_the_anneal_temp_note(live_config) -> None:
    notes = review_notes_from_config(live_config)
    assert any("anneal_temp" in note and "mean function" in note for note in notes)
    assert any("exploration-only" in note for note in notes)


def test_no_review_block_means_no_probes_and_no_notes() -> None:
    assert probe_specs_from_config({}) == ()
    assert review_notes_from_config({}) == ()


def test_a_single_probe_mapping_is_accepted() -> None:
    probes = probe_specs_from_config(
        {"review": {"probes": {"name": "p", "column": "speed_1", "value": 1000}}}
    )
    assert probes == (ProbeSpec("p", "speed_1", 1000.0),)


def test_a_probe_naming_an_undeclared_input_is_refused_before_any_fitting(
    config,
) -> None:
    """A typo in a config column name should cost nothing, not three GP fits."""
    broken = dict(config)
    broken["review"] = {"probes": [{"name": "bad", "column": "not_an_input", "value": 1}]}
    names = [i["name"] for i in config["inputs"]]
    with pytest.raises(ValueError, match="not a declared input"):
        build_batch_review(
            broken,
            np.zeros((0, 3)),
            np.zeros((0, 3)),
            pd.DataFrame(_rows(config, 1, seed=3), columns=names),
            round_name="R1",
        )


# --------------------------------------------------------------------------- #
# the candidate table
# --------------------------------------------------------------------------- #


def test_every_candidate_gets_a_utility_and_an_sd_per_objective(review) -> None:
    for name in ("uniformity", "optoelectronic", "thickness"):
        assert f"{name}_utility" in review.candidates
        assert f"{name}_sd" in review.candidates
        assert review.candidates[f"{name}_sd"].gt(0).all()
    assert len(review.candidates) == 3


def test_the_physical_prediction_is_reported_in_the_measurement_s_units(review) -> None:
    """A utility of 0.87 means nothing at the coater; nanometres do."""
    assert review.candidates["thickness_predicted"].between(50.0, 5000.0).all()
    for column in ("thickness_lo68", "thickness_hi68"):
        # numeric, so a spreadsheet can sort, plot and compare it
        assert pd.api.types.is_float_dtype(review.candidates[column])


def test_a_log_link_objective_reports_a_median_not_a_mean(review) -> None:
    """exp of a mean of logs is the median. The interval brackets it
    multiplicatively; an additive one would be symmetric in nanometres, which is
    what transforming a mean rather than decoding a posterior would produce."""
    low = review.candidates["thickness_lo68"]
    high = review.candidates["thickness_hi68"]
    middle = review.candidates["thickness_predicted"]
    assert (low < middle).all() and (middle < high).all()
    assert np.allclose(middle / low, high / middle, rtol=1e-12)
    assert np.allclose(np.sqrt(low * high), middle, rtol=1e-12)


def test_an_identity_link_objective_gets_a_symmetric_interval(review) -> None:
    low = review.candidates["optoelectronic_lo68"]
    high = review.candidates["optoelectronic_hi68"]
    middle = review.candidates["optoelectronic_predicted"]
    assert np.allclose((low + high) / 2.0, middle, rtol=1e-12)


def test_distance_to_the_nearest_observed_point_is_zero_when_reproposed(config) -> None:
    X, Y = _observations(config)
    names = [i["name"] for i in config["inputs"]]
    built = build_batch_review(
        config, X, Y, pd.DataFrame(X[:2], columns=names), round_name="R1"
    )
    assert built.candidates["distance_to_nearest"].max() < 1e-9
    assert built.candidates["nearest_observed_row"].tolist() == [1, 2]


def test_range_edges_are_counted_and_named(config, design) -> None:
    """Which coordinates are pinned matters more than how many: 'anneal_temp=min'
    on every row is a mean function speaking, and a count alone hides that."""
    names = list(design.names)
    X, Y = _observations(config)
    row = list(X[0])
    row[names.index("speed_1")] = float(design.lowers[names.index("speed_1")])
    row[names.index("anneal_temp")] = float(design.uppers[names.index("anneal_temp")])
    built = build_batch_review(
        config, X, Y, pd.DataFrame([row], columns=names), round_name="R1"
    )
    assert built.candidates["n_at_range_edge"].iloc[0] == 2
    which = built.candidates["which_at_range_edge"].iloc[0]
    assert "speed_1=min" in which and "anneal_temp=max" in which


def test_a_candidate_at_no_range_edge_says_so_rather_than_leaving_a_blank(
    review,
) -> None:
    which = review.candidates["which_at_range_edge"]
    assert which.notna().all()
    assert (which.str.len() > 0).all()


# --------------------------------------------------------------------------- #
# probes and text
# --------------------------------------------------------------------------- #


def test_the_probe_moves_every_candidate_to_the_probed_value(review) -> None:
    kinds = review.probes["kind"].tolist()
    assert sum("moved to speed_1=1000" in kind for kind in kinds) == 3


def test_the_probe_reports_observations_already_in_the_region(config) -> None:
    X, Y = _observations(config)
    X[0, 0] = 1000.0
    names = [i["name"] for i in config["inputs"]]
    built = build_batch_review(
        config, X, Y, pd.DataFrame(_rows(config, 2, seed=42), columns=names), round_name="R1"
    )
    kinds = built.probes["kind"].tolist()
    assert any("observed row 1" in kind for kind in kinds)
    assert built.probes["measured"].notna().any()


def test_the_verdict_names_the_mean_function_when_the_probed_column_is_in_it(
    review,
) -> None:
    """speed_1 is a feature of the thickness mean, so a probe there asks a fitted
    trend to extrapolate to its range edge -- a different claim from a GP
    interpolating between neighbours, and the artifact should say which it is."""
    text = " ".join(review.probe_verdicts)
    if "KNOWN AND BAD" in text:
        assert "thickness carries speed_1 in its mean function" in text
        assert "fitted global trend" in text


def test_the_verdict_prints_the_sd_ratio_and_its_threshold(review) -> None:
    """The ratio is printed whichever branch fires, so a reader who wants a
    different threshold can apply their own."""
    import re

    text = " ".join(review.probe_verdicts)
    assert re.search(r"sd \d+\.\d+ vs \d+\.\d+ \(x\d+\.\d+\)", text)
    assert f"{SD_MATERIALITY_RATIO:g}x counts as materially more uncertain" in text


def test_the_verdict_states_how_many_observations_sit_in_the_region(review) -> None:
    assert "observation(s) already sit there" in " ".join(review.probe_verdicts)


def test_the_text_artifact_stands_alone(review) -> None:
    text = review.to_text()
    for section in (
        "BATCH REVIEW - R1",
        "PROPOSED CONDITIONS",
        "PROBES",
        "PROBE 'low-speed corner'",
        "NOTES",
        "CARRIED FROM THE MEASURED DATA",
    ):
        assert section in text
    assert "1600" in text  # the carried finding, verbatim
    assert text.rstrip().endswith("outside this file.")
    assert NOT_APPROVED.split(".")[0] in text


def test_findings_are_carried_worst_first(review) -> None:
    text = review.to_text()
    assert text.index("[warning] sample 12") < text.index("[note] sample 4")


# --------------------------------------------------------------------------- #
# when the mean function explains the data
# --------------------------------------------------------------------------- #


@pytest.fixture
def collapsed_residual(monkeypatch):
    """Force the condition rather than hope data produces it.

    Whether a given dataset lands on a collapsed residual is knife-edge -- measured
    across residual magnitudes from 0 to 0.3, it fires at 0, 1e-4, 0.01 and 0.03 but
    not at 0.001 or 0.1, because it depends on where the MLL optimiser lands. A test
    that depended on that would be a flake. The guard's own decision is tested
    directly in test_model_validation.py; what these tests check is that its warning
    reaches the people who need it.

    Only objectives with a `StructuredMean` are collapsed, which is exactly the case
    being emulated: the mean function explains the data, so the residual GP has
    nothing left. Uniformity has no mean function and stays healthy -- collapsing it
    would be a true collapse and must still hard-fail.
    """
    import mobo_kit.model_validation as validation_module
    from mobo_kit.structured_mean import StructuredMean

    real_fit = validation_module.fit_gpytorch_mll

    def collapse_structured_only(mll):
        real_fit(mll)
        if isinstance(getattr(mll.model, "mean_module", None), StructuredMean):
            import torch

            mll.model.covar_module.outputscale = torch.tensor(1e-12, dtype=torch.double)
            mll.model.likelihood.noise = torch.tensor(0.9, dtype=torch.double)
        return mll

    monkeypatch.setattr(validation_module, "fit_gpytorch_mll", collapse_structured_only)


def test_a_round_still_proposes_when_the_mean_function_explains_the_data(
    config, collapsed_residual
) -> None:
    """The behaviour that matters: refusing here would dead-end the campaign
    exactly when the physics model started working, with no way out -- better data
    cannot be collected without first proposing conditions."""
    from mobo_kit.campaign import run_r1_ucb

    X, Y = _observations(config)
    result = run_r1_ucb(config, X, Y, n=2)
    assert result.n_conditions == 2

    warnings = result.diagnostics["model_fit_warnings"]
    assert warnings, "the collapsed residual GP should have warned"
    joined = " ".join(warnings)
    assert "exploration term has degenerated" in joined
    assert "UNDERSTATED" in joined
    # the two objectives with a mean function, and not the one without
    assert any(w.startswith("thickness") for w in warnings)
    assert not any(w.startswith("uniformity") for w in warnings)


def test_the_round_diagnostics_carry_no_library_deprecation_noise(
    config, collapsed_residual
) -> None:
    """`record.warnings` also collects every Python warning raised while fitting --
    on this stack, ~18 numpy-2.0 deprecation notices per fit. Putting those in front
    of someone reviewing a batch is how people learn to ignore warnings."""
    from mobo_kit.campaign import run_r1_ucb

    warnings = run_r1_ucb(config, *_observations(config), n=2).diagnostics[
        "model_fit_warnings"
    ]
    assert warnings
    assert not any("numpy" in w.lower() or "__array__" in w for w in warnings)


def test_the_unfiltered_warnings_are_kept_but_not_surfaced(config) -> None:
    """Filtered out of the human channel, retained for debugging. A BoTorch or
    scipy convergence warning the filter dropped is exactly what someone needs when
    a fit looks strange weeks later."""
    from mobo_kit.campaign import run_r1_ucb

    diagnostics = run_r1_ucb(config, *_observations(config), n=2).diagnostics
    raw = diagnostics["fit_warnings_raw"]
    assert raw, "the fits do raise library warnings on this stack"
    assert any("numpy" in entry.lower() or "__array__" in entry for entry in raw)
    # each entry says which objective and which stage it came from
    assert all("|" in entry for entry in raw)
    # and the surfaced channel is still clean
    assert not diagnostics["model_fit_warnings"]


def _collapsed_review(config):
    X, Y = _observations(config)
    names = [i["name"] for i in config["inputs"]]
    return build_batch_review(
        config,
        X,
        Y,
        pd.DataFrame(_rows(config, 2, seed=77), columns=names),
        round_name="R1",
    )


def test_the_review_carries_the_warning_above_the_numbers(
    config, collapsed_residual
) -> None:
    built = _collapsed_review(config)
    assert built.model_warnings

    text = built.to_text()
    assert "READ THIS BEFORE THE NUMBERS" in text
    # before, not after: it changes how every number below should be read
    assert text.index("READ THIS BEFORE THE NUMBERS") < text.index("PROPOSED CONDITIONS")
    assert "understated" in text.lower()
    assert "no uncertainty" in text


def test_the_warning_reaches_the_review_sheet_too(
    tmp_path, config, collapsed_residual
) -> None:
    built = _collapsed_review(config)
    path = tmp_path / "candidates.xlsx"
    _candidate_book(path)
    write_review_sheet(path, built)
    body = _sheet_text(path)
    assert "READ THIS BEFORE THE NUMBERS" in body
    assert "UNDERSTATED" in body


def test_a_healthy_fit_carries_no_warning(review) -> None:
    """The warning has to mean something, which means it must not always fire."""
    assert review.model_warnings == ()
    assert "READ THIS BEFORE THE NUMBERS" not in review.to_text()


def test_the_review_uses_the_measured_noise_the_round_was_given(config) -> None:
    """Under replicate_pooled the proposing GPs carry measured noise. A review that
    refits without it describes a different model from the one that chose the
    batch -- on the first R2 the uniformity utilities came out up to 0.027 high. So
    the review takes the same variance, and with it reproduces the round's model."""
    from mobo_kit.batch_review import _physical_predictions
    from mobo_kit.campaign import build_objective_transform, fit_campaign_models

    X, Y = _observations(config)
    names = [i["name"] for i in config["inputs"]]
    conditions = pd.DataFrame(_rows(config, 2, seed=42), columns=names)
    # a measured variance in each objective's MODEL space: log for a log link
    links = [spec.model_link for spec in build_objective_transform(config).specs]
    target = np.column_stack(
        [np.log(Y[:, i]) if link == "log" else Y[:, i] for i, link in enumerate(links)]
    )
    Yvar = np.tile(0.05 * target.var(axis=0), (len(Y), 1))

    measured = build_batch_review(
        config, X, Y, conditions, round_name="R2", observed_Yvar=Yvar
    )
    fitted = build_batch_review(config, X, Y, conditions, round_name="R2")
    model, _ = fit_campaign_models(config, X, Y, Yvar=Yvar)
    expected = _physical_predictions(config, model, conditions[names].to_numpy(float))

    for objective, values in expected.items():
        assert np.allclose(
            measured.candidates[f"{objective}_predicted"].to_numpy(float), values[:, 0]
        )
    assert any(
        not np.allclose(
            measured.candidates[f"{objective}_utility"].to_numpy(float),
            fitted.candidates[f"{objective}_utility"].to_numpy(float),
            rtol=1e-6,
        )
        for objective in expected
    )


# --------------------------------------------------------------------------- #
# the sheet
# --------------------------------------------------------------------------- #


def _candidate_book(path) -> None:
    book = Workbook()
    book.active.title = "R1_Candidates"
    book["R1_Candidates"]["A1"] = "candidate_id"
    book.save(path)


def _sheet_text(path, sheet: str = "Review") -> str:
    return "\n".join(
        str(value)
        for row in load_workbook(path)[sheet].iter_rows(values_only=True)
        for value in row
        if value is not None
    )


def test_the_review_is_written_as_its_own_sheet(tmp_path, review) -> None:
    path = tmp_path / "candidates.xlsx"
    _candidate_book(path)
    write_review_sheet(path, review)

    written = load_workbook(path)
    assert written.sheetnames == ["R1_Candidates", "Review"]
    # the worklist is untouched
    assert written["R1_Candidates"]["A1"].value == "candidate_id"

    body = _sheet_text(path)
    assert "BATCH REVIEW - R1" in body
    assert "PROPOSED CONDITIONS" in body
    assert "APPROVAL" in body
    assert "Nothing here is approved" in body
    assert "1600" in body


def test_rewriting_the_review_replaces_it_rather_than_duplicating(
    tmp_path, review
) -> None:
    path = tmp_path / "candidates.xlsx"
    _candidate_book(path)
    write_review_sheet(path, review)
    write_review_sheet(path, review)
    assert load_workbook(path).sheetnames == ["R1_Candidates", "Review"]


def test_non_finite_cells_are_written_as_blanks_not_the_text_nan(
    tmp_path, review
) -> None:
    """The probe table carries NaN in the 'selected' columns of observed rows.
    Excel showing the literal text 'nan' would read as a measurement."""
    path = tmp_path / "candidates.xlsx"
    _candidate_book(path)
    write_review_sheet(path, review)
    assert "nan" not in _sheet_text(path).lower().replace("anneal", "")
