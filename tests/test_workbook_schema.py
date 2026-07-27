from __future__ import annotations

import hashlib
import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mobo_kit.d2d_campaign import load_d2d_debug_config, load_d2d_workbook_frame
from mobo_kit.workbook_schema import (
    WorkbookInputExceptionRule,
    audit_campaign_workbook,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PRIVATE_WORKBOOK = os.environ.get("MOBO_KIT_D2D_PRIVATE_WORKBOOK")
PRIVATE_CONFIG = os.environ.get("MOBO_KIT_D2D_PRIVATE_CONFIG")
SYNTHETIC_SAMPLE_IDS = tuple(range(1001, 1016))
V3_HEADERS = (
    "Sample number",
    "speed_1",
    "time_1",
    "speed_2",
    "time_2",
    "precur_conc",
    "precur_vol (uL)",
    "anneal_temp",
    "anneal_time",
    "anti_vol",
    "anti_time",
    "Coverage",
    "Uniformity",
    "1 - Uniformity",
    "Phase purity",
    "PL - Implied Voc (Max)",
    "Photoconductance (Max)",
    "Log10 (Photoconductance (Max) x PL - Implied Voc (Max))",
    "T1",
    "T2",
    "T3",
    "T4",
    "T anom",
    "Thickness (avg)",
    "Normalized thickness (sigma = 250)",
    "Uniformity score",
    "Optoelectronic score",
    "Thickness score",
    "Stability score?",
    "Total combination - addition",
    "Total combination - multiplied",
    "Uniformity score absolute difference",
    "Optoelectronic score absolute difference",
    "Thickness absolute difference",
    "Total score absolute difference",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _create_sanitized_summary_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    headers: list[str | None] = [
        "Sample number",
        "speed_1",
        "time_1",
        "speed_2",
        "time_2",
        "precur_conc",
        "precur_vol",
        "anneal_temp",
        "anneal_time",
        "anti_vol",
        "anti_time",
        "Anneal Temp",
        "Raw measurement A",
        "Raw measurement B",
        "Raw measurement C",
        "Derived score A",
        "Uniformity score",
        "Derived score B",
        "Derived score C",
        "Uniformity score",
        "Thickness score",
        "Optoelectronic score",
        "Total combination - multiplication",
        "Total combination - addition",
        None,
        "BO objective 1",
        "BO objective 2",
        "BO objective 3",
        "Notes",
    ]
    assert len(headers) == 29

    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=column, value=header)

    for row, sample_id in enumerate(SYNTHETIC_SAMPLE_IDS, start=2):
        worksheet.cell(row=row, column=1, value=sample_id)

    # Ensure worksheet.max_row/max_column are misleading. The auditor must
    # ignore this style-only cell and report the range of nonblank values.
    worksheet["AZ100"].number_format = "0.00"
    workbook.save(path)
    workbook.close()


def _create_sanitized_v2_workbook(
    path: Path, *, note_cells: tuple[str, str] = ("P18", "R18")
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    headers: list[str | None] = [
        "Sample number",
        "speed_1",
        "time_1",
        "speed_2",
        "time_2",
        "precur_conc",
        "precur_vol (uL)",
        "anneal_temp",
        "anneal_time",
        "anti_vol",
        "anti_time",
        "Coverage",
        "Uniformity",
        "1 - Uniformity",
        "Phase purity",
        "PL - Implied Voc (Max)",
        "PL - Implied Voc (Max) Normalized",
        "Photoconductance (Max)",
        "Photoconductance (Max) Normalized",
        "Uniformity score",
        "Optoelectronic score",
        "Thickness (avg)",
        "Normalized thickness (sigma = 250)",
        "Total combination - addition",
        None,
        "Uniformity score absolute difference",
        "Optoelectronic score absolute difference",
        "Thickness absolute difference",
        "Total score absolute difference",
    ]
    assert len(headers) == 29
    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=column, value=header)

    grids = [
        list(range(1000, 6001, 500)),
        list(range(5, 51, 5)),
        list(range(0, 5001, 500)),
        list(range(10, 61, 5)),
        [round(1.0 + index * 0.05, 2) for index in range(21)],
        list(range(40, 201, 10)),
        list(range(100, 186, 5)),
        list(range(10, 61, 5)),
        list(range(100, 201, 5)),
        list(range(9, 26, 2)),
    ]
    for sample_index in range(15):
        row = sample_index + 2
        worksheet.cell(row=row, column=1, value=SYNTHETIC_SAMPLE_IDS[sample_index])
        for dimension, grid in enumerate(grids):
            value = grid[(sample_index * (dimension + 1) + dimension) % len(grid)]
            worksheet.cell(row=row, column=dimension + 2, value=value)
    note = "*Likely normalize to a max theoretical, to be discussed"
    for cell in note_cells:
        worksheet[cell] = note
    worksheet["AS47"].number_format = "0.00"
    workbook.save(path)
    workbook.close()


def _create_sanitized_v3_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    assert len(V3_HEADERS) == 35
    for column, header in enumerate(V3_HEADERS, start=1):
        worksheet.cell(row=1, column=column, value=header)

    grids = [
        list(range(1000, 6001, 500)),
        list(range(5, 51, 5)),
        list(range(0, 5001, 500)),
        list(range(10, 61, 5)),
        [round(1.0 + index * 0.05, 2) for index in range(21)],
        list(range(40, 201, 10)),
        list(range(100, 186, 5)),
        list(range(10, 61, 5)),
        list(range(100, 201, 5)),
        list(range(9, 26, 2)),
    ]
    for sample_index in range(15):
        row = sample_index + 2
        worksheet.cell(row=row, column=1, value=SYNTHETIC_SAMPLE_IDS[sample_index])
        for dimension, grid in enumerate(grids):
            value = grid[(sample_index * (dimension + 1) + dimension) % len(grid)]
            worksheet.cell(row=row, column=dimension + 2, value=value)
        worksheet.cell(row=row, column=12, value=0.9)
        worksheet.cell(row=row, column=13, value=0.2)
        worksheet.cell(row=row, column=14, value=0.8)
        worksheet.cell(row=row, column=15, value=0.95)
        worksheet.cell(row=row, column=16, value=1.1 + sample_index * 0.01)
        worksheet.cell(row=row, column=17, value=100 + sample_index)
        worksheet.cell(row=row, column=18, value=f"=LOG10(P{row}*Q{row})")
        for column, thickness in zip(range(19, 23), (620, 640, 660, 680)):
            worksheet.cell(row=row, column=column, value=thickness + sample_index)
        worksheet.cell(row=row, column=23, value=999)
        worksheet.cell(row=row, column=24, value=f"=AVERAGE(S{row}:V{row})")
        worksheet.cell(
            row=row,
            column=25,
            value=f"=EXP(-((AVERAGE(S{row}:V{row})-650)/250)^2)",
        )
        worksheet.cell(row=row, column=26, value=0.65 + sample_index * 0.01)
        worksheet.cell(row=row, column=27, value=2.0 + sample_index * 0.01)
        worksheet.cell(row=row, column=28, value=0.9 - sample_index * 0.01)

    for row in range(17, 21):
        worksheet.cell(row=row, column=1, value=f"Synthetic note row {row}")
    worksheet["AI20"] = "Synthetic ignored-summary note"
    # Whitespace-only values, including NBSP, must not extend the content range
    # or be reported as notes.
    worksheet["AJ21"] = " \u00a0\t "
    worksheet["S17"] = "\u00a0"
    worksheet["AZ100"].number_format = "0.00"
    workbook.save(path)
    workbook.close()


def test_audit_campaign_workbook_preserves_raw_schema_and_file(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sanitized_summary_table.xlsx"
    _create_sanitized_summary_workbook(workbook_path)

    before_hash = _sha256(workbook_path)
    before_mtime_ns = workbook_path.stat().st_mtime_ns

    audit = audit_campaign_workbook(workbook_path)

    assert audit.workbook_path == workbook_path
    assert audit.sheet_names == ["Sheet1"]
    assert audit.active_sheet == "Sheet1"
    assert audit.used_range == "A1:AC16"
    assert audit.column_count == 29
    assert len(audit.raw_headers) == 29
    assert audit.sample_row_count == 15
    assert audit.profile == "d2d_summary_step1_historical"

    assert audit.raw_headers[16] == "Uniformity score"  # Q
    assert audit.raw_headers[19] == "Uniformity score"  # T
    assert audit.duplicate_headers == {"Uniformity score": [17, 20]}
    assert audit.raw_headers[24] is None  # Y
    assert audit.blank_header_columns == [25]

    duplicate_warning = next(
        warning for warning in audit.warnings if "Duplicate header" in warning
    )
    assert "Uniformity score" in duplicate_warning
    assert "[17, 20]" in duplicate_warning

    anneal_warning = next(
        warning
        for warning in audit.warnings
        if "anneal_temp" in warning and "Anneal Temp" in warning
    )
    assert "normalize" in anneal_warning

    blank_warning = next(
        warning for warning in audit.warnings if "Blank header" in warning
    )
    assert "column 25 (Y)" in blank_warning
    assert "Total combination - addition" in blank_warning

    assert _sha256(workbook_path) == before_hash
    assert workbook_path.stat().st_mtime_ns == before_mtime_ns


def test_updated_v2_profile_aliases_grid_rows_notes_and_file_integrity(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "sanitized_summary_v2.xlsx"
    _create_sanitized_v2_workbook(workbook_path)
    before_hash = _sha256(workbook_path)
    before_mtime_ns = workbook_path.stat().st_mtime_ns

    audit = audit_campaign_workbook(workbook_path)

    assert audit.profile == "d2d_summary_v2"
    assert audit.sheet_names == ["Sheet1"]
    assert audit.active_sheet == "Sheet1"
    assert audit.used_range == "A1:AC18"
    assert audit.column_count == 29
    assert audit.sample_row_count == 15
    assert audit.sample_rows == list(range(2, 17))
    assert audit.raw_headers[6] == "precur_vol (uL)"
    assert audit.canonical_input_mapping["precur_vol (uL)"] == "precur_vol"
    assert audit.canonical_input_positions == {
        "speed_1": 2,
        "time_1": 3,
        "speed_2": 4,
        "time_2": 5,
        "precur_conc": 6,
        "precur_vol": 7,
        "anneal_temp": 8,
        "anneal_time": 9,
        "anti_vol": 10,
        "anti_time": 11,
    }
    assert "Anneal Temp" not in audit.raw_headers
    assert audit.header_positions["Uniformity score"] == [20]
    assert audit.header_positions["Optoelectronic score"] == [21]
    assert audit.header_positions["Thickness (avg)"] == [22]
    assert audit.header_positions["Normalized thickness (sigma = 250)"] == [23]
    assert audit.blank_header_columns == [25]
    assert [note.cell for note in audit.notes] == ["P18", "R18"]
    assert audit.input_rows_valid is True
    assert audit.input_validation_errors == []
    assert audit.objective_mapping_approved is False
    assert audit.formula_cell_count == 0
    assert not any("P18/R18" in warning for warning in audit.warnings)
    assert _sha256(workbook_path) == before_hash
    assert workbook_path.stat().st_mtime_ns == before_mtime_ns


def test_v3_profile_exact_contract_and_file_integrity(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "sanitized_summary_v3.xlsx"
    _create_sanitized_v3_workbook(workbook_path)
    before_hash = _sha256(workbook_path)
    before_mtime_ns = workbook_path.stat().st_mtime_ns

    audit = audit_campaign_workbook(workbook_path)

    assert audit.profile == "d2d_summary_v3_scores"
    assert audit.sheet_names == ["Sheet1"]
    assert audit.active_sheet == "Sheet1"
    assert audit.used_range == "A1:AI20"
    assert audit.column_count == 35
    assert tuple(audit.raw_headers) == V3_HEADERS
    assert audit.sample_row_count == 15
    assert audit.sample_rows == list(range(2, 17))
    assert {note.row for note in audit.notes} == {17, 18, 19, 20}
    assert all(note.cell != "S17" for note in audit.notes)

    assert audit.canonical_input_mapping["precur_vol (uL)"] == "precur_vol"
    assert audit.canonical_input_positions == {
        "speed_1": 2,
        "time_1": 3,
        "speed_2": 4,
        "time_2": 5,
        "precur_conc": 6,
        "precur_vol": 7,
        "anneal_temp": 8,
        "anneal_time": 9,
        "anti_vol": 10,
        "anti_time": 11,
    }
    assert audit.header_positions["Uniformity score"] == [26]
    assert audit.header_positions["Optoelectronic score"] == [27]
    assert audit.header_positions["Thickness score"] == [28]
    assert audit.canonical_objective_mapping == {
        "Uniformity score": "uniformity_score",
        "Optoelectronic score": "optoelectronic_score",
        "Thickness score": "thickness_score",
    }
    assert audit.canonical_objective_positions == {
        "uniformity_score": 26,
        "optoelectronic_score": 27,
        "thickness_score": 28,
    }
    assert audit.ignored_model_positions == {
        "Stability score?": 29,
        "Total combination - addition": 30,
        "Total combination - multiplied": 31,
        "Uniformity score absolute difference": 32,
        "Optoelectronic score absolute difference": 33,
        "Thickness absolute difference": 34,
        "Total score absolute difference": 35,
    }
    assert audit.objective_mapping_resolved_for_debug is True
    assert audit.objective_mapping_approved is False
    assert audit.blank_header_columns == []
    assert audit.formula_cell_count == 45

    assert audit.input_rows_valid is True
    assert audit.input_validation_errors == []
    assert audit.input_exceptions == []
    assert any(
        "resolved for debug only" in warning
        and "not approved for production" in warning
        for warning in audit.warnings
    )

    assert _sha256(workbook_path) == before_hash
    assert workbook_path.stat().st_mtime_ns == before_mtime_ns


def test_v3_profile_accepts_explicit_algorithmic_synthetic_exception(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "synthetic_exception_v3.xlsx"
    _create_sanitized_v3_workbook(workbook_path)
    grid = list(range(1000, 6001, 500))
    sentinel = (grid[0] + grid[1]) / 2.0
    workbook = load_workbook(workbook_path)
    workbook["Sheet1"]["B2"] = sentinel
    workbook.save(workbook_path)
    workbook.close()

    rule = WorkbookInputExceptionRule(
        sample_id=SYNTHETIC_SAMPLE_IDS[0],
        input_name="speed_1",
        observed_value=sentinel,
        reason="algorithmic synthetic off-grid sentinel",
    )
    audit = audit_campaign_workbook(workbook_path, allowed_input_exceptions=(rule,))

    assert audit.input_rows_valid is True
    assert len(audit.input_exceptions) == 1
    assert audit.input_exceptions[0].sample_id == SYNTHETIC_SAMPLE_IDS[0]
    assert audit.input_exceptions[0].observed_value == sentinel
    assert any("Observed-only input exception" in warning for warning in audit.warnings)


def test_v3_profile_rejects_nonexception_grid_and_bounds_errors(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "invalid_summary_v3.xlsx"
    _create_sanitized_v3_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Sheet1"]
    speed_grid = list(range(1000, 6001, 500))
    worksheet["B3"] = (speed_grid[0] + speed_grid[1]) / 2.0
    worksheet["B4"] = 999
    worksheet["C5"] = "NaN"
    workbook.save(workbook_path)
    workbook.close()

    audit = audit_campaign_workbook(workbook_path)

    assert audit.profile == "d2d_summary_v3_scores"
    assert audit.input_rows_valid is False
    assert any(
        "Excel row 3" in error and "off-grid" in error
        for error in audit.input_validation_errors
    )
    assert any(
        "Excel row 4" in error and "out-of-bounds" in error
        for error in audit.input_validation_errors
    )
    assert any(
        "Row 5" in error and "non-finite" in error
        for error in audit.input_validation_errors
    )
    assert audit.input_exceptions == []


def test_v3_profile_requires_exact_content_range(tmp_path: Path) -> None:
    workbook_path = tmp_path / "short_summary_v3.xlsx"
    _create_sanitized_v3_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Sheet1"]
    worksheet["A20"] = None
    worksheet["AI20"] = None
    workbook.save(workbook_path)
    workbook.close()

    audit = audit_campaign_workbook(workbook_path)

    assert audit.profile == "d2d_summary_v3_scores"
    assert audit.used_range == "A1:AI19"
    assert audit.input_rows_valid is False
    assert any(
        "requires content range A1:AI20" in error
        for error in audit.input_validation_errors
    )


def test_v3_profile_rejects_missing_shifted_duplicate_and_ambiguous_headers(
    tmp_path: Path,
) -> None:
    def audit_with_headers(name: str, updates: dict[str, str | None]):
        workbook_path = tmp_path / f"{name}.xlsx"
        _create_sanitized_v3_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        worksheet = workbook["Sheet1"]
        for cell, value in updates.items():
            worksheet[cell] = value
        workbook.save(workbook_path)
        workbook.close()
        return audit_campaign_workbook(workbook_path)

    missing = audit_with_headers("missing", {"Z1": None})
    assert missing.profile == "unknown"
    assert missing.blank_header_columns == [26]

    shifted = audit_with_headers(
        "shifted",
        {"AA1": "Thickness score", "AB1": "Optoelectronic score"},
    )
    assert shifted.profile == "unknown"
    assert shifted.header_positions["Optoelectronic score"] == [28]
    assert shifted.header_positions["Thickness score"] == [27]

    duplicate = audit_with_headers("duplicate", {"AA1": "Uniformity score"})
    assert duplicate.profile == "unknown"
    assert duplicate.duplicate_headers == {"Uniformity score": [26, 27]}
    assert any("Duplicate header" in warning for warning in duplicate.warnings)

    ambiguous = audit_with_headers("ambiguous", {"AC1": "Thickness-score"})
    assert ambiguous.profile == "unknown"
    assert any(
        "Ambiguous related headers" in warning
        and "Thickness score" in warning
        and "Thickness-score" in warning
        for warning in ambiguous.warnings
    )


def test_v2_profile_reports_off_grid_and_duplicate_recipe_errors(tmp_path: Path):
    workbook_path = tmp_path / "invalid_summary_v2.xlsx"
    _create_sanitized_v2_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Sheet1"]
    worksheet["B2"] = 1001
    for column in range(2, 12):
        worksheet.cell(
            row=3, column=column, value=worksheet.cell(row=4, column=column).value
        )
    workbook.save(workbook_path)
    workbook.close()
    audit = audit_campaign_workbook(workbook_path)
    assert audit.profile == "d2d_summary_v2"
    assert audit.input_rows_valid is False
    assert any("off-grid" in error for error in audit.input_validation_errors)
    assert any(
        "duplicate input grid tuples" in error
        for error in audit.input_validation_errors
    )


def test_v2_note_location_anomaly_is_reported_without_changing_profile(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "sanitized_summary_v2_note_anomaly.xlsx"
    _create_sanitized_v2_workbook(workbook_path, note_cells=("Q18", "S18"))

    audit = audit_campaign_workbook(workbook_path)

    assert audit.profile == "d2d_summary_v2"
    assert [note.cell for note in audit.notes] == ["Q18", "S18"]
    assert any("P18/R18" in warning for warning in audit.warnings)


def test_v2_profile_requires_exact_headers_sheet_rows_and_sample_ids(
    tmp_path: Path,
) -> None:
    wrong_header_path = tmp_path / "wrong_header.xlsx"
    _create_sanitized_v2_workbook(wrong_header_path)
    workbook = load_workbook(wrong_header_path)
    workbook["Sheet1"]["L1"] = "Coverage renamed"
    workbook.save(wrong_header_path)
    workbook.close()
    assert audit_campaign_workbook(wrong_header_path).profile == "unknown"

    wrong_sheet_path = tmp_path / "wrong_sheet.xlsx"
    _create_sanitized_v2_workbook(wrong_sheet_path)
    workbook = load_workbook(wrong_sheet_path)
    workbook["Sheet1"].title = "Not Sheet1"
    workbook.save(wrong_sheet_path)
    workbook.close()
    assert audit_campaign_workbook(wrong_sheet_path).profile == "unknown"

    incomplete_path = tmp_path / "incomplete_rows.xlsx"
    _create_sanitized_v2_workbook(incomplete_path)
    workbook = load_workbook(incomplete_path)
    worksheet = workbook["Sheet1"]
    worksheet["A16"] = None
    worksheet["Z17"] = "not blank"
    workbook.save(incomplete_path)
    workbook.close()
    incomplete = audit_campaign_workbook(incomplete_path)
    assert incomplete.profile == "d2d_summary_v2"
    assert incomplete.input_rows_valid is False
    assert any("rows exactly" in error for error in incomplete.input_validation_errors)
    assert any(
        "unique numeric sample identifiers" in error
        for error in incomplete.input_validation_errors
    )
    assert any("row 17" in error for error in incomplete.input_validation_errors)


def test_sample_count_uses_only_nonblank_sample_identifiers(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sample_ids.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Sample ID"
    worksheet["B1"] = "Measurement"
    worksheet["A2"] = "S-001"
    worksheet["B3"] = "row without a sample identifier"
    worksheet["A4"] = "S-002"
    workbook.save(workbook_path)
    workbook.close()

    audit = audit_campaign_workbook(workbook_path)

    assert audit.used_range == "A1:B4"
    assert audit.sample_row_count == 2


@pytest.mark.local_input
@pytest.mark.skipif(
    not PRIVATE_WORKBOOK or not PRIVATE_CONFIG,
    reason=(
        "set MOBO_KIT_D2D_PRIVATE_WORKBOOK and MOBO_KIT_D2D_PRIVATE_CONFIG "
        "to opt into the ignored private integration test"
    ),
)
def test_private_workbook_schema_if_explicitly_enabled() -> None:
    private_workbook = Path(str(PRIVATE_WORKBOOK)).expanduser().resolve()
    private_config = load_d2d_debug_config(
        Path(str(PRIVATE_CONFIG)).expanduser().resolve()
    )
    before_hash = _sha256(private_workbook)
    before_mtime_ns = private_workbook.stat().st_mtime_ns
    _, audit = load_d2d_workbook_frame(
        private_workbook,
        expected_profile=private_config.workbook_profile,
        expected_sample_ids=private_config.expected_sample_ids,
        allowed_input_exceptions=private_config.off_grid_exceptions,
    )

    assert audit.sheet_names == ["Sheet1"]
    assert audit.active_sheet == "Sheet1"
    assert audit.profile == "d2d_summary_v3_scores"
    assert audit.used_range == "A1:AI20"
    assert audit.column_count == 35
    assert audit.sample_row_count == 15
    assert audit.sample_rows == list(range(2, 17))
    assert tuple(audit.raw_headers) == V3_HEADERS
    assert audit.header_positions.get("Uniformity score") == [26]
    assert audit.header_positions.get("Optoelectronic score") == [27]
    assert audit.header_positions.get("Thickness score") == [28]
    assert audit.canonical_input_mapping.get("precur_vol (uL)") == "precur_vol"
    assert audit.canonical_objective_positions == {
        "uniformity_score": 26,
        "optoelectronic_score": 27,
        "thickness_score": 28,
    }
    assert audit.ignored_model_positions == {
        "Stability score?": 29,
        "Total combination - addition": 30,
        "Total combination - multiplied": 31,
        "Uniformity score absolute difference": 32,
        "Optoelectronic score absolute difference": 33,
        "Thickness absolute difference": 34,
        "Total score absolute difference": 35,
    }
    assert audit.blank_header_columns == []
    assert {note.row for note in audit.notes} == {17, 18, 19, 20}
    assert audit.input_rows_valid is True
    assert audit.input_validation_errors == []
    assert len(audit.input_exceptions) == len(private_config.off_grid_exceptions)
    assert audit.objective_mapping_resolved_for_debug is True
    assert audit.objective_mapping_approved is False
    assert _sha256(private_workbook) == before_hash
    assert private_workbook.stat().st_mtime_ns == before_mtime_ns
