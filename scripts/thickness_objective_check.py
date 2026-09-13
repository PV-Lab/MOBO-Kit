"""Acceptance check for plan rev. 2 section 1.2.

Compares two ways of getting a thickness *score* prediction out of a GP:

  A. train on the score directly            (what the repo did before v2; it now does B)
  B. train on raw nanometres, then push the posterior through the 650 nm
     Gaussian analytically via ObjectiveTransform.expected_transform

Both are scored against the true score by exact leave-one-out, so the comparison
is like for like.  Also reported is the naive variant of B that transforms only
the posterior mean, to show what ignoring the variance costs.

    python scripts/thickness_objective_check.py
"""

from __future__ import annotations

import argparse
import math
import warnings
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook
from scipy.stats import spearmanr

from mobo_kit.model_validation import (
    DIM_SCALED_PRIOR,
    LEGACY_NO_PRIOR,
    fit_model_variant,
)
from mobo_kit.objectives import ObjectiveSpec, ObjectiveTransform

DESIGN = (
    ("speed_1", 1000.0, 6000.0),
    ("time_1", 5.0, 50.0),
    ("speed_2", 0.0, 5000.0),
    ("time_2", 10.0, 60.0),
    ("precur_conc", 1.0, 2.0),
    ("precur_vol", 40.0, 200.0),
    ("anneal_temp", 100.0, 185.0),
    ("anneal_time", 10.0, 60.0),
    ("anti_vol", 100.0, 200.0),
    ("anti_time", 9.0, 25.0),
)
TARGET_NM = 650.0
# workbook: =EXP(-(((X-650)/250)^2)), which is exp(-0.5*((X-650)/s)^2) with s = 250/sqrt(2)
SIGMA_NM = 250.0 / math.sqrt(2.0)

THICKNESS_UTILITY = ObjectiveTransform(
    [
        ObjectiveSpec(
            "thickness", "target", "gaussian_target", target=TARGET_NM, sigma=SIGMA_NM
        )
    ],
    version="D2D-thickness-nm-v1",
)


def load(path: Path):
    ws = load_workbook(path, data_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    lo = np.array([d[1] for d in DESIGN])
    hi = np.array([d[2] for d in DESIGN])
    X = (np.array([[r[j] for j in range(1, 11)] for r in rows], float) - lo) / (hi - lo)
    T_nm = np.array([r[23] for r in rows], float)  # X: Thickness (avg)
    score = np.array([r[27] for r in rows], float)  # AB: Thickness score
    ids = tuple(int(r[0]) for r in rows)
    return ids, X, T_nm, score


def loo_posterior(X, y, ids, variant, seed=73):
    """Exact LOO posterior mean and variance for a single objective."""
    n = len(y)
    mean = np.empty(n)
    var = np.empty(n)
    for i in range(n):
        keep = [j for j in range(n) if j != i]
        torch.manual_seed(seed)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            rec = fit_model_variant(
                torch.tensor(X[keep], dtype=torch.double),
                torch.tensor(y[keep], dtype=torch.double).unsqueeze(-1),
                sample_ids=tuple(ids[j] for j in keep),
                objective_names=("y",),
                variant=variant,
                seed=seed,
            )
            rec.model.eval()
            with torch.no_grad():
                post = rec.model.posterior(
                    torch.tensor(X[i : i + 1], dtype=torch.double)
                )
                mean[i] = float(post.mean.reshape(-1)[0])
                var[i] = float(post.variance.reshape(-1)[0])
    return mean, var


def r2(y, p):
    return 1.0 - np.sum((y - p) ** 2) / np.sum((y - y.mean()) ** 2)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", default="local_inputs/Summary Table.xlsx")
    ap.add_argument(
        "--variant",
        default="dim_scaled_prior",
        choices=["dim_scaled_prior", "legacy_matern_no_prior"],
    )
    args = ap.parse_args()

    variant = (
        DIM_SCALED_PRIOR if args.variant == "dim_scaled_prior" else LEGACY_NO_PRIOR
    )
    ids, X, T_nm, score = load(Path(args.workbook))
    n = len(ids)
    null = 1.0 - (n / (n - 1)) ** 2

    # sanity: the workbook score must be reproducible from raw nm
    recomputed = (
        THICKNESS_UTILITY(torch.tensor(T_nm, dtype=torch.double).unsqueeze(-1))
        .numpy()
        .reshape(-1)
    )
    print(
        f"score == exp(-((T-650)/250)^2) from raw nm ?  "
        f"max|diff| = {np.abs(recomputed - score).max():.2e}"
    )
    print(f"model variant: {variant.name}")
    print(f"N = {n}, LOO-mean null R2 = {null:+.4f}\n")

    # --- A: train on the score directly -----------------------------------
    a_mean, _ = loo_posterior(X, score, ids, variant)

    # --- B: train on nanometres, transform the posterior -------------------
    b_mu, b_var = loo_posterior(X, T_nm, ids, variant)
    mu_t = torch.tensor(b_mu, dtype=torch.double).unsqueeze(-1)
    var_t = torch.tensor(b_var, dtype=torch.double).unsqueeze(-1)
    b_expected = THICKNESS_UTILITY.expected_transform(mu_t, var_t).numpy().reshape(-1)
    b_meanonly = THICKNESS_UTILITY(mu_t).numpy().reshape(-1)

    print("=== predicting the THICKNESS SCORE, exact leave-one-out ===")
    print(f"{'approach':>46} {'LOO R2':>9} {'Spearman':>10}")
    for label, pred in (
        ("A  train on score directly", a_mean),
        ("B  train on nm -> E[score] (analytic)", b_expected),
        ("B' train on nm -> score(mean) only", b_meanonly),
    ):
        print(
            f"{label:>46} {r2(score, pred):>+9.4f} "
            f"{spearmanr(score, pred).statistic:>+10.4f}"
        )
    print(f"{'null (LOO mean)':>46} {null:>+9.4f} {-1.0:>+10.4f}")

    print("\n=== the underlying raw-nm model ===")
    print(
        f"  LOO R2 = {r2(T_nm, b_mu):+.4f}   Spearman = "
        f"{spearmanr(T_nm, b_mu).statistic:+.4f}"
    )
    print(
        f"  posterior sd over the 15 folds: min {np.sqrt(b_var).min():.1f} nm, "
        f"median {np.median(np.sqrt(b_var)):.1f} nm, max {np.sqrt(b_var).max():.1f} nm"
    )

    print("\n=== what the uncertainty penalty is doing ===")
    print(
        f"{'sample':>7} {'true nm':>9} {'pred nm':>9} {'sd nm':>8} "
        f"{'true score':>11} {'E[score]':>10} {'score(mean)':>12}"
    )
    for i, sid in enumerate(ids):
        print(
            f"{sid:>7} {T_nm[i]:>9.0f} {b_mu[i]:>9.0f} {math.sqrt(b_var[i]):>8.0f} "
            f"{score[i]:>11.4f} {b_expected[i]:>10.4f} {b_meanonly[i]:>12.4f}"
        )


if __name__ == "__main__":
    main()
