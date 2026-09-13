"""Screen candidate objectives built from RAW measurements instead of stored scores.

WHY THIS EXISTS. Every contract so far has handed the GP a composite score --
clamped, capped, normalised, averaged -- and every one of those composites has come
back unlearnable. The extended C1&C2 sheet showed the mechanism on the one axis
where both forms exist: leave-one-recipe-out R2 is **-0.2151** on the stored
thickness score and **+0.4082** on the same films' raw nanometres. The Gaussian
squash `EXP(-((T-650)/250)^2)` is non-monotone, so 500 nm and 800 nm map to the
same score and the GP is asked to learn a fold that is not there.

This script asks whether the same is true of uniformity and optoelectronic: give
the model Coverage, 1-Uniformity, phase purity, Voc, photoconductance and
photosensitivity AS MEASURED, and see which of them it can predict. Composites
are then built in UTILITY space, after the GP, where a monotone squash costs
nothing.

    python scripts/raw_component_screen.py --list
    python scripts/raw_component_screen.py --candidates coverage,phase_purity,log_photocond
    python scripts/raw_component_screen.py --spec '[{"name":"x","expr":"np.log(g_light)"}]'
    python scripts/raw_component_screen.py --candidates phase_purity --permutations 600

THE STATISTICS, AND THE TRAP. N = 15.

**-0.1480 IS NOT A SIGNIFICANCE THRESHOLD, and this script used to imply it was.**
`1-(N/(N-1))^2` is the score of the leave-one-out MEAN predictor -- predict every
held-out film with the average of the other fourteen. A fitted GP does not do
that. Measured here on 2026-09-04, 300 permutations of `phase_purity` with the
campaign's own model: the fitted GP's null has median **-0.4210** and 95th
percentile **+0.2890**, and **28.7% of pure-noise shuffles score above -0.1480**.
Beating it is a one-in-four event under no signal at all. The honest
single-candidate bar is the 95th percentile of the candidate's OWN empirical
null, which `--calibrate` measures; it runs about 0.3 to 0.4 R2 units above the
number this project quoted for a year.

A mean function LOWERS that null rather than raising it -- an OLS trend fitted on
14 rows of shuffled y is a noise fit, and extrapolating it to the held-out row
adds error. Median goes -0.4075 (no mean) -> -0.4368 (one feature) -> -0.5384
(two). So a mean-function candidate is not flattered by its null; it was simply
being scored against a bar five times too low, like everything else.

The bootstrap resolution sd is **+-0.236** -- wider than most effects anyone will
find here. Screening K candidates and reporting the best R2 is selection on the
outcome: at K = 30 several candidates clear any fixed bar by chance alone. So

  * every run prints how many candidates were screened, in the summary, always;
  * `--permutations` runs the rank permutation test, which is the project's
    adjudicator because rank is what the acquisition consumes;
  * `--family-size K` Bonferroni-adjusts that p for a screen of K candidates.
    Report the ADJUSTED p when the candidate was chosen by looking at this data.

A candidate that beats the null but whose adjusted p is not significant is a
hypothesis for the next batch of films, not a finding.
"""

from __future__ import annotations

import argparse
import ast
import json
import sys
import warnings
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
import openpyxl
import torch
from scipy.stats import spearmanr

from mobo_kit.campaign import load_campaign_config, normalise_inputs
from mobo_kit.loocv import RESOLUTION_SD_AT_15, null_loo_r2, resolution_sd
from mobo_kit.model_validation import DIM_SCALED_PRIOR, fit_model_variant
from mobo_kit.structured_mean import (
    MeanFeature,
    StructuredMeanSpec,
    build_structured_mean,
)

warnings.filterwarnings("ignore")

DEFAULT_WORKBOOK = "local_inputs/Final Summary Table.xlsx"
DEFAULT_CONFIG = "configs/campaign_d2d_perovskite_final.yaml"

#: The raw measurement namespace, by workbook column. Names are what a candidate
#: expression may refer to; the letters are where they live on the R0 sheet.
#: Deliberately includes BOTH the raw and the normalised form of everything, so a
#: candidate can be written either way and the difference measured rather than
#: assumed.
COLUMNS: dict[str, str] = {
    # design
    "speed_1": "B", "time_1": "C", "speed_2": "D", "time_2": "E",
    "precur_conc": "F", "precur_vol": "G", "anneal_temp": "H",
    "anneal_time": "I", "anti_vol": "J", "anti_time": "K",
    # uniformity family
    "coverage": "L",
    "uniformity_raw": "M",
    "uniformity_clamped": "N",
    "one_minus_unif": "O",
    "phase_purity": "P",
    # optoelectronic family
    "voc_raw": "Q",
    "voc_clamped": "R",
    "voc_norm": "S",
    "g_light": "T",
    "g_dark": "U",
    "g_dark_floor": "V",
    "photocond_raw": "W",
    "photocond": "X",
    "photocond_norm": "Y",
    "photosens_ratio": "Z",
    "photosens_capped": "AA",
    "photosens_norm": "AB",
    # thickness family
    "thickness_nm": "AH",
    "thickness_norm": "AI",
    # the stored composites, for reference only
    "score_uniformity": "AJ",
    "score_opto": "AK",
    "score_thickness": "AL",
}

#: The baseline screen. Every RAW component on its own, then the stored scores it
#: is being compared against, then the handful of composites that can be argued
#: for from the chemistry rather than fitted from the data.
BUILT_IN: list[dict[str, str]] = [
    # --- uniformity family, raw ---
    {"name": "coverage", "expr": "coverage", "family": "uniformity"},
    {"name": "one_minus_unif", "expr": "one_minus_unif", "family": "uniformity"},
    {"name": "uniformity_raw", "expr": "uniformity_raw", "family": "uniformity"},
    {"name": "log_uniformity_raw", "expr": "np.log(uniformity_raw)", "family": "uniformity"},
    {"name": "phase_purity", "expr": "phase_purity", "family": "uniformity"},
    {"name": "logit_phase_purity", "expr": "np.log(phase_purity / (1 - phase_purity))",
     "family": "uniformity"},
    # --- optoelectronic family, raw ---
    {"name": "voc_raw", "expr": "voc_raw", "family": "optoelectronic"},
    {"name": "g_light", "expr": "g_light", "family": "optoelectronic"},
    {"name": "log_g_light", "expr": "np.log(g_light)", "family": "optoelectronic"},
    {"name": "photocond", "expr": "photocond", "family": "optoelectronic"},
    {"name": "log_photocond", "expr": "np.log(photocond)", "family": "optoelectronic"},
    {"name": "photosens_ratio", "expr": "photosens_ratio", "family": "optoelectronic"},
    {"name": "log_g_dark_floor", "expr": "np.log(g_dark_floor)", "family": "optoelectronic"},
    # --- thickness family, the known-good control ---
    {"name": "thickness_nm", "expr": "thickness_nm", "family": "thickness"},
    {"name": "log_thickness_nm", "expr": "np.log(thickness_nm)", "family": "thickness"},
    # --- the stored composites, as the thing to beat ---
    {"name": "STORED_score_uniformity", "expr": "score_uniformity", "family": "stored"},
    {"name": "STORED_score_opto", "expr": "score_opto", "family": "stored"},
    {"name": "STORED_score_thickness", "expr": "score_thickness", "family": "stored"},
]

#: AST nodes a candidate expression may contain. No attribute access except the
#: `np.` namespace, no calls except to numpy, no comprehensions, no names outside
#: the measurement namespace. These expressions come from the analyst (or from an
#: agent proposing candidates), not from the workbook, but a screen that silently
#: evaluates arbitrary text is a bad instrument regardless of who is typing.
_ALLOWED_NODES = (
    ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
    ast.Call, ast.Attribute, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow,
    ast.USub, ast.UAdd, ast.Mod, ast.Tuple, ast.keyword,
)
_ALLOWED_NP = {
    "log", "log10", "log1p", "exp", "sqrt", "abs", "clip", "minimum", "maximum",
    "power", "square", "cbrt", "sign", "arctan", "tanh", "mean", "prod", "sum",
}


def _check_expression(expr: str) -> None:
    tree = ast.parse(expr, mode="eval")
    for node in ast.walk(tree):
        if not isinstance(node, _ALLOWED_NODES):
            raise ValueError(f"{type(node).__name__} is not allowed in a candidate expression")
        if isinstance(node, ast.Attribute):
            if not (isinstance(node.value, ast.Name) and node.value.id == "np"):
                raise ValueError("only the `np.` namespace may be attribute-accessed")
            if node.attr not in _ALLOWED_NP:
                raise ValueError(f"np.{node.attr} is not on the allowed list")
        if isinstance(node, ast.Name) and node.id not in COLUMNS and node.id != "np":
            raise ValueError(
                f"unknown name {node.id!r}; the measurement namespace is "
                + ", ".join(sorted(COLUMNS))
            )


def read_measurements(workbook: Path, sheet: str, n_rows: int | None = None) -> dict[str, np.ndarray]:
    """Every named column, as float, stopping at the first blank sample number."""
    ws = openpyxl.load_workbook(workbook, data_only=True)[sheet]
    last = 1
    for row in range(2, ws.max_row + 1):
        if ws[f"A{row}"].value is None:
            break
        last = row
    if n_rows is not None:
        last = min(last, 1 + n_rows)
    out: dict[str, np.ndarray] = {}
    for name, letter in COLUMNS.items():
        values = [ws[f"{letter}{row}"].value for row in range(2, last + 1)]
        out[name] = np.array(
            [np.nan if v is None else float(v) for v in values], dtype=float
        )
    return out


def evaluate(expr: str, space: Mapping[str, np.ndarray]) -> np.ndarray:
    _check_expression(expr)
    value = eval(  # noqa: S307 - namespace is whitelisted by _check_expression
        compile(ast.parse(expr, mode="eval"), "<candidate>", "eval"),
        {"__builtins__": {}, "np": np},
        dict(space),
    )
    return np.asarray(value, dtype=float)


def _mean_spec(item: Mapping[str, Any], response: str) -> StructuredMeanSpec | None:
    """Build a candidate's structured mean, if it declares one.

    A feature may be a bare column name (identity transform) or
    ``{"column": ..., "transform": "log"}`` -- the campaign config's own shape.
    Without the second form this screen COULD NOT EXPRESS the mean function the
    live campaign ran until its withdrawal on 2026-09-06, ``log(speed_1) +
    log(precur_conc)`` on a log response, so every "beats the incumbent" comparison
    it made was against a
    different model. Found 2026-09-04 by an adversarial verifier; the incumbent
    measures +0.7423, matching the config's own recorded +0.7422, against the
    +0.7633 the screen had been calling it.
    """
    features = item.get("mean_features")
    if not features:
        return None
    return StructuredMeanSpec(
        response=response,
        features=tuple(
            MeanFeature(feature)
            if isinstance(feature, str)
            else MeanFeature(
                str(feature["column"]), str(feature.get("transform", "identity"))
            )
            for feature in features
        ),
    )


def loo_r2(
    config: Mapping[str, Any],
    X_phys: np.ndarray,
    y: np.ndarray,
    *,
    seed: int = 73,
    mean_spec: StructuredMeanSpec | None = None,
) -> dict[str, Any]:
    """Exact leave-one-out under the campaign's own model variant.

    Refits the structured mean INSIDE every fold when one is given. Fitting it
    once on everything leaks the held-out value into the mean function, which is
    the single easiest way to manufacture a result on 15 rows.
    """
    X_norm = normalise_inputs(config, np.asarray(X_phys, float))
    y = np.asarray(y, float)
    n = len(y)
    design_names = [item["name"] for item in config["inputs"]]
    lowers = np.array([float(item["start"]) for item in config["inputs"]])
    uppers = np.array([float(item["stop"]) for item in config["inputs"]])
    mu = np.empty(n)
    collapsed = 0
    previous = torch.get_num_threads()
    torch.set_num_threads(1)
    try:
        for held in range(n):
            keep = [i for i in range(n) if i != held]
            target = y[keep]
            module = None
            if mean_spec is not None:
                module, target = build_structured_mean(
                    np.asarray(X_phys, float)[keep], y[keep], mean_spec,
                    design_names, lowers, uppers,
                )
            torch.manual_seed(seed)
            try:
                record = fit_model_variant(
                    torch.tensor(X_norm[keep], dtype=torch.double),
                    torch.tensor(target, dtype=torch.double).unsqueeze(-1),
                    sample_ids=tuple(range(len(keep))),
                    objective_names=("y",),
                    variant=DIM_SCALED_PRIOR,
                    seed=seed,
                    mean_module=module,
                )
            except Exception:
                collapsed += 1
                mu[held] = target.mean()
                continue
            gp = record.model.models[0]
            gp.eval()
            with torch.no_grad():
                value = float(
                    gp.posterior(
                        torch.tensor(X_norm[held: held + 1], dtype=torch.double)
                    ).mean.reshape(-1)[0]
                )
            mu[held] = np.exp(value) if (mean_spec and mean_spec.response == "log") else value
    finally:
        torch.set_num_threads(previous)
    if collapsed == n:
        # THE FALLBACK'S OWN SCORE IS THE NUMBER THIS PROJECT USED AS ITS NULL.
        # Every fold falling back to its training mean makes `mu` the
        # leave-one-out mean predictor exactly, whose R2 is 1-(n/(n-1))^2 -- so a
        # totally broken run would report -0.1480 and rho -1.0000 and look like an
        # ordinary no-signal result. Refuse instead. Found 2026-09-04 by an
        # adversarial verifier that reproduced it with a 1e-8 perturbation of the
        # design matrix, which pushes `normalise_inputs` a few parts in a billion
        # outside [0, 1] and makes every fit raise.
        raise RuntimeError(
            f"All {n} folds failed to fit, so every prediction is its fold's "
            "training mean. That degenerate predictor scores exactly "
            f"{1.0 - (n / (n - 1)) ** 2:+.4f} with Spearman -1.0000, which is "
            "indistinguishable from an ordinary no-signal result. Check the "
            "candidate for non-finite values, a constant response, or inputs "
            "outside their declared grids."
        )
    residual = ((y - mu) ** 2).sum()
    total = ((y - y.mean()) ** 2).sum()
    return {
        "r2": float(1.0 - residual / total),
        "spearman": float(spearmanr(y, mu).statistic),
        "collapsed_folds": collapsed,
        "predicted": mu.tolist(),
    }


def permutation_p(
    config, X_phys, y, observed_rho, *, permutations: int, seed: int = 73, mean_spec=None
) -> dict[str, Any]:
    """Rank permutation null: shuffle y, redo the whole fold loop, count exceedances.

    Rank rather than R2 because rank is what the acquisition consumes -- it never
    sees R2 -- and because a rank null is unaffected by the heavy tails that make
    R2 unstable at N = 15.
    """
    rng = np.random.default_rng(seed)
    y = np.asarray(y, float)
    null = []
    for index in range(permutations):
        shuffled = rng.permutation(y)
        null.append(loo_r2(config, X_phys, shuffled, seed=seed, mean_spec=mean_spec)["spearman"])
        if (index + 1) % 50 == 0:
            print(f"    permutation {index + 1}/{permutations}", file=sys.stderr, flush=True)
    null = np.asarray(null, float)
    exceed = int((null >= observed_rho).sum())
    p = (exceed + 1) / (permutations + 1)
    return {
        "permutations": permutations,
        "null_mean": float(null.mean()),
        "null_sd": float(null.std(ddof=1)),
        "exceedances": exceed,
        "p": float(p),
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--config", default=DEFAULT_CONFIG)
    parser.add_argument("--sheet", default="R0")
    parser.add_argument("--seed", type=int, default=73)
    parser.add_argument("--list", action="store_true", help="print the measurement namespace and exit")
    parser.add_argument(
        "--candidates",
        default=None,
        help="comma-separated names from the built-in screen; default is all of them",
    )
    parser.add_argument(
        "--spec",
        default=None,
        help=(
            'JSON list of {"name","expr"[,"family"][,"mean_features"]} to screen '
            "INSTEAD of the built-ins. `expr` is a numpy expression over the "
            "measurement namespace (see --list). `mean_features` is a list of "
            'design column names to fit a linear mean on, e.g. ["precur_conc"].'
        ),
    )
    parser.add_argument("--mean-response", default="identity", choices=("identity", "log"))
    parser.add_argument("--permutations", type=int, default=0)
    parser.add_argument(
        "--calibrate",
        type=int,
        default=0,
        metavar="DRAWS",
        help=(
            "measure the EMPIRICAL null for each candidate instead of screening "
            "it: permute y this many times, rerun the whole fold loop, and report "
            "the null's median and 95th percentile. The 95th percentile is the "
            "honest single-candidate bar. Use it before quoting any R2. "
            "SLOW and single-threaded: every draw is a full fold loop, about "
            "3 s per draw per candidate at N=15, so 500 draws is ~25 min. "
            "Parallelise across processes if you need more."
        ),
    )
    parser.add_argument(
        "--family-size",
        type=int,
        default=None,
        help="K for the Bonferroni adjustment; defaults to the number screened",
    )
    parser.add_argument("--out", default=None, help="write the full result table as JSON")
    args = parser.parse_args(argv)

    if args.list:
        print("measurement namespace (name -> R0 column):")
        for name, letter in COLUMNS.items():
            print(f"  {name:24} {letter}")
        print("\nbuilt-in candidates:")
        for item in BUILT_IN:
            print(f"  {item['name']:26} = {item['expr']}")
        return 0

    config = load_campaign_config(args.config)
    space = read_measurements(Path(args.workbook), args.sheet)
    X = np.column_stack([space[item["name"]] for item in config["inputs"]])
    n = len(X)
    null = null_loo_r2(n)

    if args.spec:
        candidates = json.loads(args.spec)
    else:
        candidates = list(BUILT_IN)
        if args.candidates:
            wanted = {name.strip() for name in args.candidates.split(",")}
            unknown = wanted - {item["name"] for item in candidates}
            if unknown:
                raise SystemExit(f"unknown candidate(s): {sorted(unknown)}")
            candidates = [item for item in candidates if item["name"] in wanted]

    family_size = args.family_size if args.family_size is not None else len(candidates)

    print(f"workbook   {args.workbook}  sheet {args.sheet}  N = {n}")
    print(f"null LOO R2 {null:+.4f}   resolution sd +-{resolution_sd(n):.3f} "
          f"(bootstrapped {RESOLUTION_SD_AT_15} at N=15)")
    print(f"screening  {len(candidates)} candidates   Bonferroni family size K = {family_size}")
    print()
    header = f"{'candidate':30} {'LOO R2':>9} {'rho':>7} {'beats null':>11}  expression"
    print(header)
    print("-" * (len(header) + 10))

    if args.calibrate:
        print(
            f"CALIBRATION: {args.calibrate} permutations per candidate. The "
            "theoretical -0.1480 is the score of the leave-one-out MEAN "
            "predictor, which is NOT what a fitted GP does; the empirical null "
            "below is."
        )
        print()
        header = (
            f"{'candidate':30} {'observed':>9} {'null med':>9} {'null p95':>9} "
            f"{'p':>7}  {'>-0.1480':>9}"
        )
        print(header)
        print("-" * len(header))
        rng = np.random.default_rng(args.seed)
        for item in candidates:
            y = evaluate(item["expr"], space)
            mean_spec = _mean_spec(item, args.mean_response)
            observed = loo_r2(config, X, y, seed=args.seed, mean_spec=mean_spec)
            draws = np.array(
                [
                    loo_r2(
                        config, X, rng.permutation(y), seed=args.seed,
                        mean_spec=mean_spec,
                    )["r2"]
                    for _ in range(args.calibrate)
                ]
            )
            p_value = float((int((draws >= observed["r2"]).sum()) + 1) / (args.calibrate + 1))
            print(
                f"{item['name']:30} {observed['r2']:>+9.4f} "
                f"{np.median(draws):>+9.4f} {np.percentile(draws, 95):>+9.4f} "
                f"{p_value:>7.4f}  {(draws > null).mean():>8.1%}"
            )
        print()
        print(
            "The last column is how often PURE NOISE beats -0.1480. If it is not "
            "near zero, -0.1480 is not a significance threshold for this model."
        )
        return 0

    results = []
    for item in candidates:
        y = evaluate(item["expr"], space)
        if not np.isfinite(y).all():
            print(f"{item['name']:30} {'SKIPPED':>9}   non-finite values")
            continue
        mean_spec = _mean_spec(item, args.mean_response)
        outcome = loo_r2(config, X, y, seed=args.seed, mean_spec=mean_spec)
        beats = outcome["r2"] > null
        row = {**item, "n": n, "null_loo_r2": null, **outcome}
        if args.permutations and beats:
            print(f"  permuting {item['name']} ...", file=sys.stderr, flush=True)
            perm = permutation_p(
                config, X, y, outcome["spearman"],
                permutations=args.permutations, seed=args.seed, mean_spec=mean_spec,
            )
            perm["p_bonferroni"] = float(min(1.0, perm["p"] * family_size))
            row["permutation"] = perm
        results.append(row)
        flag = "YES" if beats else "no"
        note = "  COLLAPSED" if outcome["collapsed_folds"] else ""
        print(
            f"{item['name']:30} {outcome['r2']:>+9.4f} {outcome['spearman']:>+7.3f} "
            f"{flag:>11}  {item['expr']}{note}"
        )
        if "permutation" in row:
            perm = row["permutation"]
            print(
                f"{'':30} {'permutation':>9}: p {perm['p']:.4f}  "
                f"Bonferroni x{family_size} = {perm['p_bonferroni']:.4f}  "
                f"(null rho {perm['null_mean']:+.3f} sd {perm['null_sd']:.3f})"
            )

    beat = [r for r in results if r["r2"] > null]
    print()
    print(f"{len(beat)} of {len(results)} screened candidates beat the null.")
    if beat:
        best = max(beat, key=lambda r: r["r2"])
        print(f"best: {best['name']}  R2 {best['r2']:+.4f}  rho {best['spearman']:+.3f}")
    print(
        "REMINDER: these candidates were chosen by looking at this data. At N=15 "
        f"the resolution sd is +-{resolution_sd(n):.3f}, so a screen of "
        f"{family_size} will produce apparent winners by chance. Quote the "
        "Bonferroni-adjusted permutation p, not the R2."
    )
    if args.out:
        Path(args.out).write_text(json.dumps(results, indent=1), encoding="utf-8")
        print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
