"""Confirm the declarative structured means reproduce the hand-rolled results.

Targets, exact leave-one-out, null = -0.1480:

    thickness (nm)   plain GP +0.183  ->  structured +0.384
    optoelectronic   plain GP  ?      ->  structured +0.244 (anneal_temp alone)

    python scripts/validate_structured_means.py
"""

from __future__ import annotations

import argparse
import warnings
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook
from scipy.stats import spearmanr

from mobo_kit.campaign import load_campaign_config
from mobo_kit.model_validation import DIM_SCALED_PRIOR, fit_model_variant
from mobo_kit.structured_mean import (
    MeanFeature,
    StructuredMeanSpec,
    apply_structured_mean,
    fit_structured_mean,
)

warnings.filterwarnings("ignore")
torch.set_num_threads(1)

THICKNESS = StructuredMeanSpec(
    response="log",
    features=(MeanFeature("speed_1", "log"), MeanFeature("precur_conc", "log")),
)
OPTO = StructuredMeanSpec(
    response="identity", features=(MeanFeature("anneal_temp", "identity"),)
)


def _gp_residual_posterior(X_norm, resid, test_norm, seed=73):
    rec = fit_model_variant(
        torch.tensor(X_norm, dtype=torch.double),
        torch.tensor(resid, dtype=torch.double).unsqueeze(-1),
        sample_ids=tuple(range(len(X_norm))),
        objective_names=("residual",),
        variant=DIM_SCALED_PRIOR,
        seed=seed,
    )
    rec.model.eval()
    with torch.no_grad():
        post = rec.model.posterior(torch.tensor(test_norm, dtype=torch.double))
        return (
            float(post.mean.reshape(-1)[0]),
            float(post.variance.reshape(-1)[0]),
        )


def loo(X_phys, X_norm, y, names, spec):
    """Exact LOO. Mean coefficients refit on the 14 training rows each fold."""
    n = len(y)
    mean = np.empty(n)
    var = np.empty(n)
    for i in range(n):
        keep = [j for j in range(n) if j != i]
        if spec is None:
            r_mu, r_var = _gp_residual_posterior(
                X_norm[keep], np.asarray(y, float)[keep], X_norm[i : i + 1]
            )
            mean[i], var[i] = r_mu, r_var
        else:
            coef, resid = fit_structured_mean(
                X_phys[keep], np.asarray(y, float)[keep], spec, names
            )
            r_mu, r_var = _gp_residual_posterior(X_norm[keep], resid, X_norm[i : i + 1])
            post = apply_structured_mean(
                coef,
                X_phys[i : i + 1],
                spec,
                names,
                np.array([r_mu]),
                np.array([r_var]),
            )
            if post.link == "log":
                # report on the original scale for comparability
                mean[i] = float(np.exp(post.mean[0] + post.variance[0] / 2.0))
                var[i] = float(
                    (np.exp(post.variance[0]) - 1.0)
                    * np.exp(2 * post.mean[0] + post.variance[0])
                )
            else:
                mean[i], var[i] = float(post.mean[0]), float(post.variance[0])
    return mean, var


def r2(y, p):
    y = np.asarray(y, float)
    return 1.0 - np.sum((y - p) ** 2) / np.sum((y - y.mean()) ** 2)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", default="local_inputs/Summary Table.xlsx")
    ap.add_argument("--config", default="configs/campaign_d2d_perovskite.yaml")
    args = ap.parse_args()

    config = load_campaign_config(args.config)
    names = [item["name"] for item in config["inputs"]]
    lo = np.array([item["start"] for item in config["inputs"]], float)
    hi = np.array([item["stop"] for item in config["inputs"]], float)

    ws = load_workbook(Path(args.workbook), data_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    X_phys = np.array([[float(r[j]) for j in range(1, 11)] for r in rows])
    X_norm = (X_phys - lo) / (hi - lo)
    thickness_nm = np.array([float(r[23]) for r in rows])
    optoelectronic = np.array([float(r[26]) for r in rows])

    n = len(rows)
    null = 1.0 - (n / (n - 1)) ** 2
    print(f"N = {n}, null LOO R2 = {null:+.4f}\n")
    print(
        f"{'objective':>16} {'mean function':>34} {'LOO R2':>9} {'Spearman':>10} {'target':>9}"
    )

    cases = [
        ("thickness nm", "none (plain GP)", thickness_nm, None, "+0.183"),
        (
            "thickness nm",
            "log T ~ log(speed_1)+log(conc)",
            thickness_nm,
            THICKNESS,
            "+0.384",
        ),
        ("optoelectronic", "none (plain GP)", optoelectronic, None, "-"),
        ("optoelectronic", "linear anneal_temp", optoelectronic, OPTO, "+0.244"),
    ]
    for label, mean_label, y, spec, target in cases:
        mu, _ = loo(X_phys, X_norm, y, names, spec)
        print(
            f"{label:>16} {mean_label:>34} {r2(y, mu):>+9.4f} "
            f"{spearmanr(y, mu).statistic:>+10.4f} {target:>9}"
        )


if __name__ == "__main__":
    main()
