"""Two evaluations that gate R1 generation.

1. Permutation test on the route-B Spearman: permute the thickness measurements,
   redo the full leave-one-out fit plus analytic transform, and build the null
   distribution.  Turns "rank improved" into a p-value.

2. Structured mean function.  Spin-coating physics says T ~ speed^-0.5, and
   log T ~ log(speed_1) + log(precur_conc) reaches LOO R2 +0.449 while the plain
   10-input GP reaches +0.183.  Test whether giving the GP a linear mean on those
   two log inputs closes the gap.  Selection of the two inputs is from physics,
   fixed before fitting; the linear coefficients are refitted inside every fold.

    python scripts/thickness_permutation_and_mean.py --permutations 200
"""

from __future__ import annotations

import argparse
import math
import warnings
from pathlib import Path

import numpy as np
import torch
from botorch.fit import fit_gpytorch_mll
from botorch.models import SingleTaskGP
from botorch.models.transforms.outcome import Standardize
from botorch.models.utils.gpytorch_modules import (
    get_covar_module_with_dim_scaled_prior,
    get_gaussian_likelihood_with_lognormal_prior,
)
from gpytorch.kernels import ScaleKernel
from gpytorch.means import LinearMean
from gpytorch.mlls import ExactMarginalLogLikelihood
from openpyxl import load_workbook
from scipy.stats import spearmanr
from sklearn.linear_model import LinearRegression

from mobo_kit.objectives import ObjectiveSpec, ObjectiveTransform

warnings.filterwarnings("ignore")
torch.set_default_dtype(torch.double)

DESIGN = [
    ("speed_1", 1000, 6000),
    ("time_1", 5, 50),
    ("speed_2", 0, 5000),
    ("time_2", 10, 60),
    ("precur_conc", 1, 2),
    ("precur_vol", 40, 200),
    ("anneal_temp", 100, 185),
    ("anneal_time", 10, 60),
    ("anti_vol", 100, 200),
    ("anti_time", 9, 25),
]
TARGET_NM, SIGMA_NM = 650.0, 250.0 / math.sqrt(2.0)
UTILITY = ObjectiveTransform(
    [
        ObjectiveSpec(
            "thickness", "target", "gaussian_target", target=TARGET_NM, sigma=SIGMA_NM
        )
    ],
    version="D2D-thickness-nm-v1",
)
SPEED_1, PRECUR_CONC = 0, 4


def load(path: Path):
    ws = load_workbook(path, data_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    lo = np.array([d[1] for d in DESIGN], float)
    hi = np.array([d[2] for d in DESIGN], float)
    Xp = np.array([[r[j] for j in range(1, 11)] for r in rows], float)
    return (
        (Xp - lo) / (hi - lo),
        Xp,
        np.array([r[23] for r in rows], float),
        np.array([r[27] for r in rows], float),
    )


def _gp(X, y, *, mean_features=None):
    """SingleTaskGP under the dim_scaled_prior contract, optional linear mean."""
    base = get_covar_module_with_dim_scaled_prior(
        ard_num_dims=X.shape[1], use_rbf_kernel=False
    )
    model = SingleTaskGP(
        X,
        y,
        covar_module=ScaleKernel(base),
        likelihood=get_gaussian_likelihood_with_lognormal_prior(),
        outcome_transform=Standardize(m=1),
    )
    if mean_features is not None:
        model.mean_module = LinearMean(input_size=mean_features, bias=True)
    fit_gpytorch_mll(ExactMarginalLogLikelihood(model.likelihood, model))
    return model


def loo_nm(X, y, *, structured_mean=False, Xphys=None, seed=73):
    """LOO posterior over raw nm. With structured_mean, a physics linear trend on
    log(speed_1) and log(precur_conc) is removed first and added back after."""
    n = len(y)
    mu = np.empty(n)
    var = np.empty(n)
    for i in range(n):
        keep = [j for j in range(n) if j != i]
        torch.manual_seed(seed)
        if structured_mean:
            F = np.c_[np.log(Xphys[:, SPEED_1]), np.log(Xphys[:, PRECUR_CONC])]
            lin = LinearRegression().fit(F[keep], np.log(y[keep]))
            resid = np.log(y[keep]) - lin.predict(F[keep])
            m = _gp(torch.tensor(X[keep]), torch.tensor(resid).unsqueeze(-1))
            m.eval()
            with torch.no_grad():
                p = m.posterior(torch.tensor(X[i : i + 1]))
                r_mu = float(p.mean.reshape(-1)[0])
                r_var = float(p.variance.reshape(-1)[0])
            log_mu = float(lin.predict(F[i : i + 1])[0]) + r_mu
            # lognormal moments back to nm
            mu[i] = math.exp(log_mu + r_var / 2.0)
            var[i] = (math.exp(r_var) - 1.0) * math.exp(2 * log_mu + r_var)
        else:
            m = _gp(torch.tensor(X[keep]), torch.tensor(y[keep]).unsqueeze(-1))
            m.eval()
            with torch.no_grad():
                p = m.posterior(torch.tensor(X[i : i + 1]))
                mu[i] = float(p.mean.reshape(-1)[0])
                var[i] = float(p.variance.reshape(-1)[0])
    return mu, var


def r2(y, p):
    return 1.0 - np.sum((y - p) ** 2) / np.sum((y - y.mean()) ** 2)


def expected_score(mu, var):
    return (
        UTILITY.expected_transform(
            torch.tensor(mu).unsqueeze(-1), torch.tensor(var).unsqueeze(-1)
        )
        .numpy()
        .reshape(-1)
    )


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", default="local_inputs/Summary Table.xlsx")
    ap.add_argument("--permutations", type=int, default=200)
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument(
        "--structured", action="store_true", help="permute the structured-mean pipeline"
    )
    args = ap.parse_args()

    X, Xphys, T_nm, score = load(Path(args.workbook))
    n = len(T_nm)
    null_r2 = 1.0 - (n / (n - 1)) ** 2

    print("=== 2. structured mean on log(speed_1), log(precur_conc) ===")
    plain_mu, plain_var = loo_nm(X, T_nm)
    struct_mu, struct_var = loo_nm(X, T_nm, structured_mean=True, Xphys=Xphys)
    print(f"{'raw-nm model':>34} {'LOO R2':>9} {'Spearman':>10}")
    for label, mu in (
        ("plain GP, 10 inputs", plain_mu),
        ("GP + physics linear mean", struct_mu),
    ):
        print(
            f"{label:>34} {r2(T_nm, mu):>+9.4f} "
            f"{spearmanr(T_nm, mu).statistic:>+10.4f}"
        )
    print(f"{'2-input log-log reference':>34} {'+0.4494':>9} {'+0.7143':>10}")

    print(f"\n{'resulting score prediction':>34} {'LOO R2':>9} {'Spearman':>10}")
    plain_s = expected_score(plain_mu, plain_var)
    struct_s = expected_score(struct_mu, struct_var)
    for label, s in (
        ("plain GP -> E[score]", plain_s),
        ("structured mean -> E[score]", struct_s),
    ):
        print(
            f"{label:>34} {r2(score, s):>+9.4f} "
            f"{spearmanr(score, s).statistic:>+10.4f}"
        )
    print(f"{'null':>34} {null_r2:>+9.4f} {-1.0:>+10.4f}")

    print("\n=== 3. how much of the structured-mean result rests on sample 1? ===")
    print("    sample 1 is the off-grid literature control and the one")
    print("    extrapolation point, so LOO metrics are sensitive to it")
    keep1 = np.arange(1, n)
    mu_x, var_x = loo_nm(
        X[keep1], T_nm[keep1], structured_mean=True, Xphys=Xphys[keep1]
    )
    s_x = expected_score(mu_x, var_x)
    n_x = len(keep1)
    print(f"{'':>34} {'LOO R2':>9} {'Spearman':>10}")
    print(
        f"{'raw nm, all 15':>34} {r2(T_nm, struct_mu):>+9.4f} "
        f"{spearmanr(T_nm, struct_mu).statistic:>+10.4f}"
    )
    print(
        f"{'raw nm, sample 1 excluded (N=14)':>34} {r2(T_nm[keep1], mu_x):>+9.4f} "
        f"{spearmanr(T_nm[keep1], mu_x).statistic:>+10.4f}"
    )
    print(
        f"{'score, sample 1 excluded':>34} {r2(score[keep1], s_x):>+9.4f} "
        f"{spearmanr(score[keep1], s_x).statistic:>+10.4f}"
    )
    print(f"{'null at N=14':>34} {1.0 - (n_x/(n_x-1))**2:>+9.4f} {-1.0:>+10.4f}")

    print(f"\n=== 1. permutation test, {args.permutations} shuffles ===")
    print("    permuting the nm measurements and redoing LOO + transform")
    print(
        f"    structured mean: {args.structured} "
        "(when true the linear mean is refit inside every null fold too)"
    )
    rng = np.random.default_rng(args.seed)
    observed = struct_s if args.structured else plain_s
    obs_rho = spearmanr(score, observed).statistic
    obs_r2 = r2(score, observed)
    null_rho, null_r2s = [], []
    for k in range(args.permutations):
        perm = rng.permutation(n)
        T_p, s_p = T_nm[perm], score[perm]
        mu_p, var_p = loo_nm(X, T_p, structured_mean=args.structured, Xphys=Xphys)
        e_p = expected_score(mu_p, var_p)
        null_rho.append(spearmanr(s_p, e_p).statistic)
        null_r2s.append(r2(s_p, e_p))
        if (k + 1) % 25 == 0:
            print(
                f"    {k+1}/{args.permutations} done, running null mean rho = "
                f"{np.mean(null_rho):+.4f}"
            )
    null_rho = np.array(null_rho)
    null_r2s = np.array(null_r2s)
    print(f"\n  observed Spearman = {obs_rho:+.4f}")
    print(
        f"  null Spearman     : mean {null_rho.mean():+.4f}, sd {null_rho.std():.4f}, "
        f"95th pct {np.percentile(null_rho, 95):+.4f}"
    )
    print(f"  p(null >= observed) = {np.mean(null_rho >= obs_rho):.4f}")
    print(f"\n  observed R2 = {obs_r2:+.4f}")
    print(f"  null R2     : mean {null_r2s.mean():+.4f}, sd {null_r2s.std():.4f}")
    print(f"  p(null >= observed) = {np.mean(null_r2s >= obs_r2):.4f}")
    print("\n  Note the null mean for Spearman is NOT zero: the leave-one-out")
    print("  shrinkage artifact drags it negative, which is exactly why a")
    print("  positive observed value is meaningful.")


if __name__ == "__main__":
    main()
