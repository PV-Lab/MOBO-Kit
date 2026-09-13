"""Render a round's figures from a terminal, exactly as the launcher does.

    python scripts/generate_round_report.py --workbook "local_inputs/Final Summary Table.xlsx"
    python scripts/generate_round_report.py --workbook <path> --data-only

Two modes, matching the two buttons:

* **default** re-derives the proposal from the config and the measured rows rather
  than reading it back from the worklist, so what the figures describe is the
  model's answer at this seed. **When a worklist for that round already exists,
  the two are compared by batch hash and the result is printed as MATCH or
  DRIFT.** They should match; if they do not, the config, the data or the seed has
  moved since the sheet was written, and the figures describe the model rather
  than the films anyone is about to run.
* **--data-only** renders everything that depends on measurements alone. This is
  the mode to use the moment a round's results are entered.

Nothing here writes to the source workbook, and nothing here approves anything.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from mobo_kit.campaign import load_campaign_config, run_r1_ucb, run_r2_qlognehvi
from mobo_kit.batch_review import build_batch_review
from mobo_kit.launcher import gather_observations, inspect_campaign
from mobo_kit.round_report import generate_round_report
from mobo_kit.workbook_io import read_campaign_workbook


def _worklist_drift(workbook, config, round_name: str, proposal) -> str:
    """Does the re-derived proposal still match the worklist on disk?

    The figures describe a proposal computed here and now. The films someone runs
    come from a sheet written earlier. Those are the same batch only if the
    config, the data and the seed have not moved -- and if they have, the figures
    are about a different experiment than the one on the bench, which is exactly
    the sort of quiet divergence that is worth a line of output.

    Compared by ``batch_hash``, so ordering is not mistaken for a difference.
    """
    from mobo_kit.candidate_diagnostics import batch_hash
    from mobo_kit.workbook_io import candidate_workbook_path, sheet_name_for_round

    path = candidate_workbook_path(workbook, round_name)
    if not path.exists():
        return f"no {path.name} on disk yet, so there is nothing to compare"
    try:
        from openpyxl import load_workbook

        sheet = load_workbook(path, data_only=True)[sheet_name_for_round(round_name)]
        header = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
        names = [item["name"] for item in config["inputs"]]
        columns = [header.index(name) for name in names]
        seen: list[list[float]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            values = [float(row[c]) for c in columns]
            if values not in seen:
                seen.append(values)
    except Exception as exc:  # noqa: BLE001 - a check must not break the report
        return f"could not read {path.name} to compare ({type(exc).__name__}: {exc})"

    on_disk = batch_hash(seen)
    derived = batch_hash(proposal.conditions.to_numpy(float))
    if on_disk == derived:
        return f"MATCH - the re-derived batch is {path.name}'s ({derived})"
    return (
        f"DRIFT - re-derived {derived} against {on_disk} in {path.name}. The "
        "config, the data or the seed has moved since that sheet was written, so "
        "these figures describe a different batch than the one on the bench."
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--config", default="configs/campaign_d2d_perovskite_final.yaml")
    parser.add_argument("--outdir", default=None)
    parser.add_argument("--seed", type=int, default=None)
    parser.add_argument(
        "--data-only",
        action="store_true",
        help="figures from the measurements alone; no batch is proposed",
    )
    parser.add_argument(
        "--shap-instances",
        type=int,
        default=15,
        help=(
            "rows to attribute. NOT the main runtime knob -- the leave-one-out "
            "refits are about two thirds of the cost and are not optional. "
            "Recorded in the manifest either way."
        ),
    )
    args = parser.parse_args(argv)

    config = load_campaign_config(args.config)
    workbook = Path(args.workbook)

    proposal = None
    review = None
    observations = None
    if not args.data_only:
        status = inspect_campaign(workbook, config)
        if not status.can_generate:
            print(f"No round is due: {status.reason}")
            print("Rendering the data-only report instead.")
        else:
            round_name = str(status.next_round)
            print(f"Proposing {round_name} to describe it...")
            X, Y, Yvar, _ = gather_observations(
                workbook, config, for_round=round_name
            )
            runner = run_r1_ucb if round_name == "R1" else run_r2_qlognehvi
            proposal = runner(config, X, Y, seed=args.seed, observed_Yvar=Yvar)
            contents = read_campaign_workbook(workbook, config)
            review = build_batch_review(
                config,
                X,
                Y,
                proposal.conditions,
                round_name=round_name,
                seed=proposal.diagnostics.get("seed"),
                findings=contents.findings,
                observed_Yvar=Yvar,
            )
            observations = (X, Y, Yvar)
            drift = _worklist_drift(workbook, config, round_name, proposal)
            print(f"  worklist check: {drift}")

    manifest = generate_round_report(
        workbook,
        config,
        proposal=proposal,
        review=review,
        outdir=args.outdir,
        seed=args.seed,
        shap_max_instances=args.shap_instances,
        progress=lambda message: print(f"  {message}", flush=True),
        observations=observations,
    )
    print()
    print(manifest.summary())
    print()
    print(f"{manifest.runtime_seconds:.1f} s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
